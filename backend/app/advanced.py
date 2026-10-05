from __future__ import annotations

import io
import re
import unicodedata
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

from app.excel import gerar_modelo_abas

COLUNAS = ["DESCRIÇÃO", "DATA", "VALOR", "DÉBITO", "CRÉDITO", "HISTÓRICO"]


def normalizar_texto(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto).casefold().strip()


def texto_celula_seguro(valor) -> str:
    if valor is None:
        return ""
    try:
        if pd.isna(valor):
            return ""
    except Exception:
        pass
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def limpar_caracteres_ilegais(valor) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(valor or ""))


def limpar_valor_monetario(valor) -> float:
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)) and not pd.isna(valor):
        return float(valor)
    texto = str(valor).strip().replace("R$", "").replace(" ", "")
    if not texto:
        return 0.0
    negativo = texto.startswith("(") and texto.endswith(")")
    texto = texto.strip("()")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        numero = float(re.sub(r"[^0-9+\-.]", "", texto))
    except ValueError:
        return 0.0
    return -abs(numero) if negativo else numero


def identificar_estorno_de_baixa(*campos):
    texto = normalizar_texto(" ".join(texto_celula_seguro(c) for c in campos))
    tokens = re.findall(r"[a-z0-9]+", texto)
    pos_estorno = [i for i, token in enumerate(tokens) if token.startswith(("estorn", "revers"))]
    pos_baixa = [i for i, token in enumerate(tokens) if token.startswith("baix")]
    return any(abs(i - j) <= 6 for i in pos_estorno for j in pos_baixa)


def _template_bytes() -> bytes:
    import base64
    resource = Path(__file__).resolve().parents[1] / "resources" / "modelo_dominio.b64"
    return base64.b64decode(resource.read_text(encoding="utf-8").strip())


def append_dataframe_sheets(workbook: bytes, extras: dict[str, pd.DataFrame]) -> bytes:
    wb = load_workbook(io.BytesIO(workbook))
    for raw_name, df in extras.items():
        name = str(raw_name)[:31] or "Relatório"
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name)
        frame = df.copy() if isinstance(df, pd.DataFrame) else pd.DataFrame(df)
        if frame.empty:
            ws.append(["Nenhum registro"])
            continue
        ws.append([str(c) for c in frame.columns])
        for row in frame.itertuples(index=False, name=None):
            values = []
            for value in row:
                if isinstance(value, pd.Timestamp):
                    value = value.to_pydatetime()
                elif value is None:
                    value = ""
                else:
                    try:
                        if pd.isna(value):
                            value = ""
                    except Exception:
                        pass
                values.append(value)
            ws.append(values)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def processar_mapa_autokraft(file_bytes: bytes, filename: str = ""):
    xls = pd.ExcelFile(io.BytesIO(file_bytes))
    abas_diarias = [
        aba for aba in xls.sheet_names
        if re.fullmatch(r"\d{2}[.-]\d{2}", str(aba).strip())
    ]
    if not abas_diarias:
        raise ValueError("Nenhuma aba diária no formato DD-MM ou DD.MM foi encontrada.")

    ano_nome = re.search(r"(?<!\d)(20\d{2})(?!\d)", str(filename))
    ano_referencia = int(ano_nome.group(1)) if ano_nome else datetime.now().year
    registros = {"Itaú": [], "Daycoval": []}

    for nome_aba in abas_diarias:
        df = pd.read_excel(xls, sheet_name=nome_aba, header=None, dtype=object)
        if df.empty or df.shape[1] < 6:
            continue
        data_raw = df.iloc[1, 2] if len(df.index) > 1 and df.shape[1] > 2 else None
        if isinstance(data_raw, (int, float)) and not pd.isna(data_raw):
            data_aba = pd.to_datetime(data_raw, unit="D", origin="1899-12-30", errors="coerce")
        else:
            data_aba = pd.to_datetime(data_raw, dayfirst=True, errors="coerce")
        if pd.isna(data_aba):
            partes = re.split(r"[.-]", str(nome_aba).strip())
            if len(partes) != 2:
                continue
            data_aba = pd.Timestamp(year=ano_referencia, month=int(partes[1]), day=int(partes[0]))

        banco_atual = None
        for _, linha in df.iterrows():
            bloco = normalizar_texto(texto_celula_seguro(linha.iloc[0]))
            if bloco == "itau":
                banco_atual = "Itaú"
            elif bloco == "daycoval":
                banco_atual = "Daycoval"

            hist_credito = texto_celula_seguro(linha.iloc[2])
            hist_debito = texto_celula_seguro(linha.iloc[4])
            txt_credito = normalizar_texto(hist_credito)
            txt_debito = normalizar_texto(hist_debito)
            if txt_credito.startswith("total de creditos") or txt_debito.startswith("total de debitos"):
                banco_atual = None
                continue
            if banco_atual is None:
                continue

            if hist_credito and not txt_credito.startswith("total"):
                valor = abs(limpar_valor_monetario(linha.iloc[3]))
                if valor:
                    registros[banco_atual].append({
                        "DESCRIÇÃO": f"BANCO {banco_atual.upper()}",
                        "DATA": data_aba,
                        "VALOR": valor,
                        "DÉBITO": "",
                        "CRÉDITO": "",
                        "HISTÓRICO": f"Recebido: {limpar_caracteres_ilegais(hist_credito).strip()}",
                    })
            if hist_debito and not txt_debito.startswith("total"):
                valor = abs(limpar_valor_monetario(linha.iloc[5]))
                if valor:
                    registros[banco_atual].append({
                        "DESCRIÇÃO": f"BANCO {banco_atual.upper()}",
                        "DATA": data_aba,
                        "VALOR": -valor,
                        "DÉBITO": "",
                        "CRÉDITO": "",
                        "HISTÓRICO": f"Pago: {limpar_caracteres_ilegais(hist_debito).strip()}",
                    })

    result = {}
    for bank, rows in registros.items():
        df = pd.DataFrame(rows, columns=COLUNAS)
        if not df.empty:
            df = df.sort_values("DATA", kind="stable").reset_index(drop=True)
        result[bank] = df
    if not any(not df.empty for df in result.values()):
        raise ValueError("Nenhum lançamento bancário válido foi lido.")
    return result


def processar_nova_geracao_banco(file_bytes, nome_aba, conta_esperada, descricao_banco):
    xls = pd.ExcelFile(io.BytesIO(file_bytes))
    conta_normalizada = re.sub(r"\D", "", conta_esperada)
    df = colunas = None
    for aba in xls.sheet_names:
        candidato = pd.read_excel(xls, sheet_name=aba, dtype=object)
        mapa = {normalizar_texto(str(col)).strip(): col for col in candidato.columns}
        obrig = ["conta", "data", "valor", "lacto", "historico", "doc"]
        if not all(nome in mapa for nome in obrig):
            continue
        contas = candidato[mapa["conta"]].apply(lambda v: re.sub(r"\D", "", texto_celula_seguro(v)))
        if contas.eq(conta_normalizada).any():
            df, colunas = candidato, mapa
            break
    if df is None:
        raise ValueError(f"A conta {conta_esperada} ({nome_aba}) não foi encontrada.")

    principais, retirados = [], []
    for _, linha in df.iterrows():
        conta = re.sub(r"\D", "", texto_celula_seguro(linha[colunas["conta"]]))
        if conta != conta_normalizada:
            continue
        data_raw = linha[colunas["data"]]
        data = (
            pd.to_datetime(data_raw, unit="D", origin="1899-12-30", errors="coerce")
            if isinstance(data_raw, (int, float)) and not pd.isna(data_raw)
            else pd.to_datetime(data_raw, dayfirst=True, errors="coerce")
        )
        if pd.isna(data):
            continue
        lacto_original = texto_celula_seguro(linha[colunas["lacto"]])
        lacto_norm = normalizar_texto(lacto_original)
        lacto = re.sub(r"\bPAGAR\b", "PAGO", lacto_original, flags=re.I)
        lacto = re.sub(r"\b(?:RECEBER|RECEBIMENTO)\b", "RECEBIDO", lacto, flags=re.I)
        raw = linha[colunas["valor"]]
        valor_original = float(raw) if isinstance(raw, (int, float)) and not pd.isna(raw) else limpar_valor_monetario(raw)
        if not valor_original:
            continue
        tipo_col = colunas.get("tipo")
        tipo = normalizar_texto(linha[tipo_col]) if tipo_col else ""
        if lacto_norm.startswith(("pagar", "pago")):
            valor = -abs(valor_original)
        elif lacto_norm.startswith(("receber", "recebido", "recebimento")):
            valor = abs(valor_original)
        elif "debito" in tipo:
            valor = -abs(valor_original)
        elif "credito" in tipo:
            valor = abs(valor_original)
        else:
            valor = valor_original

        hist_raw = linha[colunas["historico"]]
        hist_exato = "" if hist_raw is None or pd.isna(hist_raw) else limpar_caracteres_ilegais(str(hist_raw))
        hist = texto_celula_seguro(hist_raw)
        doc = texto_celula_seguro(linha[colunas["doc"]])
        hist_limpo = re.sub(r"^(?:Pago|Recebido)\s*:\s*", "", hist, flags=re.I).strip()
        historico = re.sub(r"\s+", " ", " ".join(p for p in [lacto, hist_limpo, doc] if p)).strip()
        registro = {
            "DESCRIÇÃO": descricao_banco,
            "DATA": data,
            "VALOR": valor,
            "DÉBITO": "",
            "CRÉDITO": "",
            "HISTÓRICO": historico,
        }
        if identificar_estorno_de_baixa(lacto_original, hist, doc):
            r = dict(registro)
            r["HISTÓRICO"] = hist_exato
            r["MOTIVO"] = "Estorno de baixa identificado"
            retirados.append(r)
        else:
            principais.append(registro)
    if not principais and not retirados:
        raise ValueError(f"Nenhum lançamento da conta {nome_aba} {conta_esperada} foi encontrado.")
    return pd.DataFrame(principais, columns=COLUNAS), pd.DataFrame(retirados)


def processar_planilha_accede_sig(file_bytes, banco_nome, empresa=""):
    from razync.companies import CONFIGURACOES_ACCEDE
    from razync.accede_1000 import aplicar_regras_accede_1000

    xls = pd.ExcelFile(io.BytesIO(file_bytes))
    registros = []

    def texto_exato(valor):
        if valor is None or pd.isna(valor):
            return ""
        if isinstance(valor, float) and valor.is_integer():
            return str(int(valor))
        return limpar_caracteres_ilegais(str(valor)).strip()

    for nome_aba in xls.sheet_names:
        bruto = pd.read_excel(xls, sheet_name=nome_aba, header=None, dtype=object)
        if bruto.empty:
            continue
        idx_header = None
        nomes_header = None
        for idx in range(min(len(bruto), 30)):
            nomes = [normalizar_texto(texto_celula_seguro(v)) for v in bruto.iloc[idx].tolist()]
            if all(nome in nomes for nome in ["data", "complemento", "entrada", "saida"]):
                idx_header, nomes_header = idx, nomes
                break
        if idx_header is None:
            continue

        def coluna(nome):
            return nomes_header.index(nome) if nome in nomes_header else None

        c_data, c_dc, c_comp, c_conf = coluna("data"), coluna("d/c"), coluna("complemento"), coluna("conf")
        c_ent, c_sai = coluna("entrada"), coluna("saida")
        linhas = bruto.iloc[idx_header + 1:].reset_index(drop=True)
        i = 0
        while i < len(linhas):
            principal = linhas.iloc[i]
            data = pd.to_datetime(principal.iloc[c_data], dayfirst=True, errors="coerce")
            if pd.isna(data):
                i += 1
                continue
            j = i + 1
            detalhes = []
            while j < len(linhas):
                prox = pd.to_datetime(linhas.iloc[j].iloc[c_data], dayfirst=True, errors="coerce")
                if not pd.isna(prox):
                    break
                if any(texto_celula_seguro(v) for v in linhas.iloc[j].tolist()):
                    detalhes.append(linhas.iloc[j])
                j += 1

            entrada = abs(limpar_valor_monetario(principal.iloc[c_ent])) if c_ent is not None else 0
            saida = abs(limpar_valor_monetario(principal.iloc[c_sai])) if c_sai is not None else 0
            sinal = 1 if entrada else (-1 if saida else 0)
            dc = texto_exato(principal.iloc[c_dc]) if c_dc is not None else ""
            complemento = texto_exato(principal.iloc[c_comp]) if c_comp is not None else ""
            conf = texto_exato(principal.iloc[c_conf]) if c_conf is not None else ""
            descricao = "BANCO ITAÚ" if normalizar_texto(banco_nome) == "itau" else "SICREDI"

            detalhes_validos = []
            for detalhe in detalhes:
                conf_doc = texto_exato(detalhe.iloc[1]) if len(detalhe) > 1 else ""
                valor_individual = abs(limpar_valor_monetario(detalhe.iloc[2])) if len(detalhe) > 2 else 0
                favorecido = texto_exato(detalhe.iloc[3]) if len(detalhe) > 3 else ""
                if valor_individual:
                    detalhes_validos.append((conf_doc, valor_individual, favorecido))

            if detalhes_validos:
                for conf_doc, valor_individual, favorecido in detalhes_validos:
                    historico = " ".join(p for p in [favorecido, conf_doc] if p).strip()
                    if not historico:
                        historico = complemento or conf or dc or "MOVIMENTO BANCARIO"
                    registros.append({
                        "DESCRIÇÃO": descricao, "DATA": data, "VALOR": round(valor_individual * (sinal or -1), 2),
                        "DÉBITO": "", "CRÉDITO": "", "HISTÓRICO": historico,
                    })
            else:
                valor = entrada if entrada else (-saida if saida else 0)
                if valor:
                    registros.append({
                        "DESCRIÇÃO": descricao, "DATA": data, "VALOR": round(valor, 2),
                        "DÉBITO": "", "CRÉDITO": "", "HISTÓRICO": complemento or conf or dc or "MOVIMENTO BANCARIO",
                    })
            i = j

    df = pd.DataFrame(registros, columns=COLUNAS)
    if df.empty:
        raise ValueError(f"Nenhum lançamento válido foi encontrado na planilha SIG do {banco_nome}.")
    df = df.sort_values("DATA", kind="stable").reset_index(drop=True)
    if empresa == "accede_automacao":
        conta = CONFIGURACOES_ACCEDE[empresa]["contas_bancarias"][normalizar_texto(banco_nome)]
        df = aplicar_regras_accede_1000(df, conta)
    return df


def processar_itau_generico(content: bytes, conta: str = "") -> pd.DataFrame:
    from razync.engekraft_969 import processar_extrato_itau_modelo
    return processar_extrato_itau_modelo(content, conta or "", ("ITAU", "ITAÚ"), "arquivo selecionado")


def processar_daycoval_generico(content: bytes) -> pd.DataFrame:
    reader = PdfReader(io.BytesIO(content))
    textos = [(p.extract_text() or "") for p in reader.pages]
    texto_total = "\n".join(textos)
    if "daycoval" not in normalizar_texto(texto_total) and "dayconnect" not in normalizar_texto(texto_total):
        raise ValueError("O PDF não parece ser um extrato Daycoval.")
    padrao_data = re.compile(r"^\s*(\d{2})\s*/\s*(\d{2})\s+(.*)$")
    padrao_moeda = re.compile(r"([+-]?\s*R\s*\$\s*[+-]?\s*\d[\d\s.]*,\s*\d\s*\d)", re.I)
    ano_match = re.search(r"\b(20\d{2})\b", texto_total)
    ano = int(ano_match.group(1)) if ano_match else datetime.now().year
    rows = []
    for linha in texto_total.splitlines():
        m = padrao_data.match(linha)
        if not m:
            continue
        corpo = re.sub(r"\s+", " ", m.group(3)).strip()
        moedas = list(padrao_moeda.finditer(corpo))
        if not moedas:
            continue
        token = re.sub(r"\s+", "", moedas[0].group(1)).upper().replace("R$", "")
        sinal = -1 if token.startswith("-") else 1
        try:
            valor = sinal * float(token.lstrip("+-").replace(".", "").replace(",", "."))
        except ValueError:
            continue
        hist = corpo[:moedas[0].start()].strip(" -|")
        if not hist or "saldo" in normalizar_texto(hist):
            continue
        try:
            data = pd.Timestamp(year=ano, month=int(m.group(2)), day=int(m.group(1)))
        except ValueError:
            continue
        rows.append({"DESCRIÇÃO":"BANCO DAYCOVAL","DATA":data,"VALOR":round(valor,2),"DÉBITO":"","CRÉDITO":"","HISTÓRICO":hist})
    if not rows:
        raise ValueError("Nenhum lançamento Daycoval foi reconhecido.")
    return pd.DataFrame(rows, columns=COLUNAS)


def _aplicar_conta(df: pd.DataFrame, conta: str) -> pd.DataFrame:
    out = df.copy()
    out["DÉBITO"] = out["VALOR"].apply(lambda v: str(conta) if float(v) > 0 else "")
    out["CRÉDITO"] = out["VALOR"].apply(lambda v: str(conta) if float(v) < 0 else "")
    return out


def workflow_modelo(company_code: int, roles: dict[str, list[tuple[str, bytes]]], options: dict) -> tuple[bytes, str]:
    if company_code in {3, 178, 343}:
        role = roles.get("mapa") or []
        if not role:
            raise ValueError("Envie o mapa bancário.")
        name, content = role[0]
        dados = processar_mapa_autokraft(content, name)
        contas = {
            3: {"Itaú":"508","Daycoval":"2283"},
            178: {"Itaú":"508","Daycoval":"505"},
            343: {"Itaú":"508","Daycoval":"506"},
        }[company_code]
        sheets = {f"{bank} {contas[bank]}": _aplicar_conta(df, contas[bank]) for bank, df in dados.items() if not df.empty}
        return gerar_modelo_abas(sheets), f"RAZYNC_{company_code}_MODELO_DOMINIO.xlsx"

    if company_code in {266, 1396}:
        role = roles.get("consolidada") or []
        if not role:
            raise ValueError("Envie a planilha consolidada.")
        _, content = role[0]
        configs = (
            [("Itaú","99549-5","BANCO ITAÚ"),("Bradesco","451990-6","BANCO BRADESCO"),("Fibra","673947-1","BANCO FIBRA")]
            if company_code == 266 else
            [("Itaú","98002-6","BANCO ITAÚ"),("Bradesco","3084-8","BANCO BRADESCO")]
        )
        sheets, retirados = {}, {}
        for nome, conta, desc in configs:
            try:
                principal, retirado = processar_nova_geracao_banco(content, nome, conta, desc)
            except ValueError:
                continue
            if not principal.empty:
                sheets[nome] = principal
            if not retirado.empty:
                retirados[f"Retirados {nome}"] = retirado
        if not sheets:
            raise ValueError("Nenhuma conta configurada da Nova Geração foi encontrada.")
        book = gerar_modelo_abas(sheets)
        return append_dataframe_sheets(book, retirados), f"RAZYNC_{company_code}_MODELO_DOMINIO.xlsx"

    if company_code == 285:
        jaguar = (roles.get("jaguar") or [])
        entradas = (roles.get("entradas") or [])
        if not jaguar or not entradas:
            raise ValueError("Envie a planilha Jaguar e a planilha de entradas detalhadas.")
        from razync.lcarlos import processar_planilhas_lcarlos
        modelo, conciliacao, resumo = processar_planilhas_lcarlos(jaguar[0][1], entradas[0][1])
        book = gerar_modelo_abas({"Santander 513": modelo})
        return append_dataframe_sheets(book, {"Conciliação": conciliacao, "Resumo": pd.DataFrame([resumo])}), "LCARLOS_285_MODELO_DOMINIO.xlsx"

    if company_code == 242:
        original = roles.get("original") or []
        if not original:
            raise ValueError("Envie o arquivo original/principal da empresa 242.")
        ano = int(options.get("ano") or datetime.now().year)
        from razync.eletro_forte import (
            processar_despesas, processar_fornecedores, processar_recebidos,
            gerar_modelo_dominio_eletro_forte, corrigir_datas_com_francesinhas,
        )
        despesas = fornecedores = recebidos = None
        if roles.get("despesas"):
            despesas = processar_despesas(roles["despesas"][0][1], ano)
        if roles.get("fornecedores"):
            fornecedores = processar_fornecedores(roles["fornecedores"][0][1], ano)
        if roles.get("recebidos"):
            recebidos = processar_recebidos(roles["recebidos"][0][1], ano)
        extras = {}
        if roles.get("francesinhas") and recebidos:
            from razync.eletro_forte_francesinhas import processar_zip_francesinhas
            francesinhas, avisos = processar_zip_francesinhas(roles["francesinhas"][0][1])
            recebidos, resumo, pendencias = corrigir_datas_com_francesinhas(recebidos, francesinhas)
            extras["Pendências francesinhas"] = pendencias
            extras["Resumo francesinhas"] = pd.DataFrame([resumo | {"avisos": " | ".join(avisos)}])
        book = gerar_modelo_dominio_eletro_forte(
            original[0][1], original[0][0], _template_bytes(), despesas, fornecedores, recebidos
        )
        if extras:
            book = append_dataframe_sheets(book, extras)
        return book, "ELETRO_FORTE_242_MODELO_DOMINIO.xlsx"

    if company_code == 968:
        sheets, extras = {}, {}
        comprovantes = pd.DataFrame()
        all_dates = []
        if roles.get("sispag"):
            from razync.radani import consolidar_comprovantes_sispag
            # período largo; análise só usa datas coincidentes com o extrato.
            arquivos = [(n,b) for n,b in roles["sispag"]]
            comprovantes = consolidar_comprovantes_sispag(arquivos, "2000-01-01", "2100-12-31")
        if roles.get("itau"):
            frames = [processar_itau_generico(b, "508") for _, b in roles["itau"]]
            ext = pd.concat(frames, ignore_index=True)
            from razync.radani import analisar_desmembramentos
            analise = analisar_desmembramentos(ext, "Itaú", comprovantes)
            sheets["Itaú 508"] = _aplicar_conta(analise.organizado, "508")
            extras["Revisões Itaú"] = analise.revisoes
            extras["Detalhes Itaú"] = analise.detalhamentos
        if roles.get("bradesco"):
            from razync.bradesco_radani import processar_extrato_bradesco_radani
            from razync.radani import analisar_desmembramentos
            frames=[]
            diag=[]
            for _, b in roles["bradesco"]:
                df, d = processar_extrato_bradesco_radani(b)
                frames.append(df)
                if isinstance(d, pd.DataFrame): diag.append(d)
                elif d: diag.append(pd.DataFrame([d]))
            ext=pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
            analise=analisar_desmembramentos(ext, "Bradesco", None)
            sheets["Bradesco 9"]=_aplicar_conta(analise.organizado,"9")
            extras["Revisões Bradesco"]=analise.revisoes
            if diag: extras["Diagnóstico Bradesco"]=pd.concat(diag, ignore_index=True)
        if not sheets:
            raise ValueError("Envie ao menos um extrato Itaú ou Bradesco.")
        return append_dataframe_sheets(gerar_modelo_abas(sheets), extras), "RADANI_968_MODELO_DOMINIO.xlsx"

    if company_code in {1000,1001}:
        empresa = "accede_automacao" if company_code == 1000 else "accede_equipamentos"
        sheets={}
        for bank in ("itau","sicredi"):
            if roles.get(bank):
                frames=[processar_planilha_accede_sig(b, bank, empresa) for _,b in roles[bank]]
                df=pd.concat(frames,ignore_index=True).sort_values("DATA",kind="stable")
                conta={"itau":"508","sicredi":"505"}[bank]
                if company_code==1001:
                    df=_aplicar_conta(df,conta)
                sheets[bank.title()] = df
        if not sheets:
            raise ValueError("Envie a planilha SIG do Itaú e/ou Sicredi.")
        return gerar_modelo_abas(sheets), f"ACCEDE_{company_code}_MODELO_DOMINIO.xlsx"

    if company_code == 1096:
        from razync.up_pack import processar_planilha_up_pack
        sheets={}
        contas={"santander":"513","sicredi":"510"}
        for bank in ("santander","sicredi"):
            if roles.get(bank):
                frames=[processar_planilha_up_pack(b,bank) for _,b in roles[bank]]
                df=pd.concat(frames,ignore_index=True).sort_values("DATA",kind="stable")
                sheets[bank.title()] = _aplicar_conta(df,contas[bank])
        if not sheets:
            raise ValueError("Envie a planilha SIG do Santander e/ou Sicredi.")
        return gerar_modelo_abas(sheets), "UP_PACK_1096_MODELO_DOMINIO.xlsx"

    if company_code == 1211:
        extrato=roles.get("extrato") or []
        boletos=roles.get("boletos") or []
        if not extrato or not boletos:
            raise ValueError("Envie o extrato Itaú e o relatório de boletos liquidados.")
        from razync.gz_1211 import processar_gz
        modelo, diag, nao_usados, resumo = processar_gz(extrato[0][1], boletos[0][1])
        book=gerar_modelo_abas({"Itaú":modelo})
        return append_dataframe_sheets(book,{"Diagnóstico":diag,"Boletos não usados":nao_usados,"Resumo":pd.DataFrame([resumo])}), "GZ_1211_MODELO_DOMINIO.xlsx"

    if company_code == 1402:
        planilha=roles.get("planilha") or []
        extrato=roles.get("extrato") or []
        if not planilha or not extrato:
            raise ValueError("Envie a planilha de caixa e o extrato BTG.")
        from razync.vgv_1402 import processar_vgv
        modelo, ext, sem_planilha, resumo=processar_vgv(planilha[0][1],extrato[0][1])
        book=gerar_modelo_abas({"BTG 510":modelo[COLUNAS]})
        return append_dataframe_sheets(book,{"Conferência":modelo,"Sem planilha":sem_planilha,"Resumo":pd.DataFrame([resumo])}), "VGV_1402_MODELO_DOMINIO.xlsx"

    if company_code == 1408:
        extrato=roles.get("extrato") or []
        if not extrato:
            raise ValueError("Envie o extrato Itaú.")
        frames=[processar_itau_generico(b,"512") for _,b in extrato]
        movimentos=pd.concat(frames,ignore_index=True).sort_values("DATA",kind="stable")
        recebidos=(roles.get("recebidos") or [])
        francesinhas=None
        extras={}
        if roles.get("francesinhas"):
            from razync.eletro_forte_francesinhas import processar_zip_francesinhas
            francesinhas, avisos=processar_zip_francesinhas(roles["francesinhas"][0][1],"512")
            extras["Avisos francesinhas"]=pd.DataFrame({"AVISO":avisos})
        from razync.eletro_forte_filial_1408 import montar_modelo_1408
        modelo,resumo=montar_modelo_1408(
            movimentos.to_dict("records"),
            recebidos[0][1] if recebidos else None,
            int(options.get("ano") or datetime.now().year),
            francesinhas,
        )
        book=gerar_modelo_abas({"Itaú 512":modelo})
        extras["Resumo"]=pd.DataFrame([resumo])
        return append_dataframe_sheets(book,extras), "ELETRO_FORTE_1408_MODELO_DOMINIO.xlsx"

    if company_code == 1529:
        from razync.nibo import processar_extrato_nibo_pdf
        sheets={}
        for bank, conta, desc in [("itau","508","BANCO ITAÚ"),("banco_brasil","8","BANCO DO BRASIL")]:
            if roles.get(bank):
                frames=[]
                for _,b in roles[bank]:
                    df=processar_extrato_nibo_pdf(b)
                    df=df[COLUNAS].copy()
                    df["DESCRIÇÃO"]=desc
                    frames.append(_aplicar_conta(df,conta))
                sheets["Itaú" if bank=="itau" else "Banco do Brasil"]=pd.concat(frames,ignore_index=True)
        if not sheets:
            raise ValueError("Envie um PDF Nibo do Itaú e/ou Banco do Brasil.")
        return gerar_modelo_abas(sheets), "DIAS_PEREIRA_1529_MODELO_DOMINIO.xlsx"

    raise NotImplementedError(f"Fluxo avançado ainda não mapeado para a empresa {company_code}.")


# Keep the original universal PDF readers, including OCR, signs and balance checks.
def processar_itau_generico(file_bytes: bytes, conta: str = '') -> pd.DataFrame:
    from app import engine
    rows = engine.processar_extrato_conferencia_empresa(file_bytes, 'extrato_itau.pdf', 'itau')
    frame = pd.DataFrame(rows, columns=COLUNAS)
    if frame.empty:
        raise ValueError('Nenhum lançamento Itaú foi encontrado.')
    return _aplicar_conta(frame, conta) if conta else frame


def processar_daycoval_generico(file_bytes: bytes) -> pd.DataFrame:
    from app import engine
    rows = engine.processar_extrato_conferencia_empresa(file_bytes, 'extrato_daycoval.pdf', 'daycoval')
    frame = pd.DataFrame(rows, columns=COLUNAS)
    if frame.empty:
        raise ValueError('Nenhum lançamento Daycoval foi encontrado.')
    return frame
