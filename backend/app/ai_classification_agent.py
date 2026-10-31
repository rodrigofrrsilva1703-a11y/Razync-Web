"""Read-only evidence investigation; approved reviews reuse the existing Excel writer."""
from collections import defaultdict, OrderedDict
import copy
import hashlib
import io
import json
import re
import secrets
import threading
import time

from openpyxl import load_workbook
import pandas as pd
from app import classification_service as base, engine
from app.classification_coverage import _normalized_header, _classified, _needs_replacement, coverage
from app.migration_services import accounts, slug

VERSION = 'base-investigator-2'
_lock = threading.RLock()
_cache = OrderedDict()
_TTL = 1200


def fingerprint(records):
    return hashlib.sha256(json.dumps(records, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def eligible(company, content, options):
    """Use the same company/profile eligibility as coverage, never alter the supplied workbook."""
    book = load_workbook(io.BytesIO(content), data_only=False)
    banks = accounts(company)
    result = []
    try:
        for ws in book:
            title = engine.normalizar_texto(ws.title).strip()
            if 'retir' in title or (slug(company).startswith('eletro_forte') and title == 'principal'):
                continue
            header, cols = _normalized_header(ws)
            if not header:
                continue
            for n, cells in enumerate(ws.iter_rows(min_row=header+1, values_only=True), header+1):
                hist = engine.texto_celula_seguro(cells[cols['historico']-1]).strip()
                if not hist or _classified(cells, cols, slug(company), options):
                    continue
                if any(ws.cell(n, cols[k]).data_type == 'f' for k in ['historico','debito','credito']):
                    continue
                value = engine.limpar_valor_monetario(cells[cols['valor']-1]) if cols.get('valor') else 0
                signature = engine.criar_assinatura_classificacao(hist)
                nature = 'pago' if value < 0 else 'recebido' if value > 0 else signature.split('|')[0]
                if nature not in {'pago','recebido'}:
                    continue
                target, bankcol = (cols['debito'], cols['credito']) if nature == 'pago' else (cols['credito'], cols['debito'])
                current = engine.texto_celula_seguro(cells[target-1]).strip()
                replacement, placeholders = _needs_replacement(cells, cols, slug(company), options)
                targetname = 'debito' if nature == 'pago' else 'credito'
                if current not in {'','0'} and not (replacement == targetname and current in placeholders):
                    continue
                bankaccount = engine.texto_celula_seguro(cells[bankcol-1]).strip()
                candidates = [key for key, account in banks.items() if str(account) == bankaccount]
                identified = engine.identificar_chave_banco_empresa(ws.title)
                if cols.get('descricao'):
                    identified = engine.identificar_chave_banco_empresa(cells[cols['descricao']-1]) or identified
                bank = candidates[0] if len(candidates)==1 else identified if identified in candidates or not candidates else ''
                if bank not in banks or (bankaccount not in {'','0'} and str(banks[bank]) != bankaccount):
                    continue
                date = cells[cols['data']-1] if cols.get('data') else ''
                result.append({'_aba':ws.title,'_linha':n,'_banco':bank,'_col_destino':target,'_col_banco':bankcol,
                    '_col_data':cols.get('data',0),'Histórico':hist,'Valor':value,'Data':str(date or ''),
                    'natureza':nature,'assinatura':signature,'Conta da contrapartida':''})
    finally:
        book.close()
    return result


def candidates_for(row, records, bankaccounts):
    grouped = {}
    signature = row['assinatura'].split('|',1)[-1]
    words = set(signature.split())
    for record in records:
        bank = record.get('banco')
        if bank not in bankaccounts:
            continue
        debit, credit = str(record.get('debito','')).strip(), str(record.get('credito','')).strip()
        nature = 'pago' if credit == str(bankaccounts[bank]) and debit != credit else 'recebido' if debit == str(bankaccounts[bank]) and debit != credit else ''
        if nature != row['natureza']:
            continue
        sig = str(record.get('assinatura',''))
        prefix, _, suffix = sig.partition('|')
        if prefix not in {nature,'outro'}:
            continue
        exact = suffix == signature
        tokens = set(suffix.split())
        overlap = len(words & tokens)/max(1,len(words | tokens))
        if not exact and (len(words & tokens)<2 or overlap<.65):
            continue
        account = debit if nature=='pago' else credit
        if not re.fullmatch(r'[1-9][0-9]{0,14}', account) or account in {str(v) for v in bankaccounts.values()}:
            continue
        # Provisional accounts cannot become a suggestion based solely on their old use.
        if account in {'14','16','166'} and row.get('_company')==242:
            continue
        key = (account, bank, exact)
        item = grouped.setdefault(key, {'conta':account,'banco':bank,'exata':exact,'mesmo_banco':bank==row['_banco'],
            'periodos':set(),'ocorrencias':0,'historicos':[]})
        item['periodos'].update(p for p in record.get('periodos',[]) if re.fullmatch(r'\d{4}-\d{2}',str(p)))
        item['ocorrencias'] += max(0,int(record.get('ocorrencias',0)))
        example = str(record.get('exemplo_historico',''))[:240]
        if example and example not in item['historicos'] and len(item['historicos'])<3:
            item['historicos'].append(example)
    values = []
    for item in grouped.values():
        item['periodos'] = sorted(item['periodos'])
        values.append(item)
    return sorted(values,key=lambda x:(not x['exata'],not x['mesmo_banco'],-len(x['periodos']),x['conta']))


def account_catalog(records, bankaccounts):
    catalog = {}
    for record in records:
        bank = record.get('banco')
        if bank not in bankaccounts:
            continue
        debit, credit = str(record.get('debito','')), str(record.get('credito',''))
        nature = 'pago' if credit == str(bankaccounts[bank]) else 'recebido' if debit == str(bankaccounts[bank]) else ''
        account = debit if nature == 'pago' else credit
        if not nature or not re.fullmatch(r'[1-9][0-9]{0,14}',account) or account in {str(a) for a in bankaccounts.values()}:
            continue
        item = catalog.setdefault((nature,account), {'natureza':nature,'conta':account,'bancos':set(),'periodos':set(),'ocorrencias':0,'exemplos':[]})
        item['bancos'].add(bank)
        item['periodos'].update(str(p) for p in record.get('periodos',[]))
        item['ocorrencias'] += int(record.get('ocorrencias',0))
        history = str(record.get('exemplo_historico','')).strip()
        example = {'historico':history[:600], 'banco':bank,
                   'periodos':sorted(str(p) for p in record.get('periodos',[])),
                   'ocorrencias':int(record.get('ocorrencias',0))}
        if history and example not in item['exemplos'] and len(item['exemplos'])<8:
            item['exemplos'].append(example)
    return [dict(item,bancos=sorted(item['bancos']),periodos=sorted(item['periodos'])) for item in catalog.values()]


def investigate(company, content, options=None, use_ai=True):
    company = int(company)
    options = base.resolve_classification_options(company, options)
    records = base.records(company)
    digest = fingerprint(records)
    filehash = hashlib.sha256(content).hexdigest()
    key = fingerprint([VERSION,company,filehash,digest,options,bool(use_ai)])
    with _lock:
        now = time.monotonic()
        for old in list(_cache):
            if now-_cache[old]['time']>_TTL:
                del _cache[old]
        if key in _cache:
            cached = copy.deepcopy(_cache[key]['result'])
            cached['cache'] = True
            cached['ia']['chamadas_nesta_execucao'] = 0
            return cached
    rows = eligible(company,content,options)
    groups = defaultdict(list)
    for row in rows:
        row['_company'] = company
        # Preserve material differences in histories; never group by supplier name alone.
        identity=(row['_banco'],row['natureza'],row['assinatura'],engine.normalizar_texto(row['Histórico']))
        groups[identity].append(row)
    suggestions=[]
    for number, members in enumerate(groups.values(),1):
        row = members[0]
        evidence = candidates_for(row,records,accounts(company))
        same = [e for e in evidence if e['exata'] and e['mesmo_banco']]
        exact = [e for e in evidence if e['exata']]
        options_found = {e['conta'] for e in exact or evidence}
        chosen = next(iter(options_found)) if len(options_found)==1 and exact else ''
        periods = set(p for e in same for p in e['periodos'])
        strength = 'FORTE' if chosen and len(periods)>=2 else 'MODERADA' if chosen else 'MANUAL'
        reason = ('Conta única em assinatura exata, no mesmo banco, em pelo menos dois períodos.' if strength=='FORTE'
                  else 'Conta histórica compatível, mas com poucos períodos ou encontrada em outro banco.' if chosen
                  else 'Existem contas conflitantes; confira as operações antes de escolher.' if len(options_found)>1
                  else 'Sem correspondência exata suficiente. Os candidatos similares exigem conferência manual.' if evidence
                  else 'Não há evidência histórica compatível para sugerir uma conta.')
        for member in members:
            suggestions.append({**member,'grupo':number,'conta_sugerida':chosen,'situacao':strength,
                'justificativa':reason,'evidencias':evidence[:12],'aprovada':False})
    result={'empresa':company,'versao':VERSION,'total_pendencias':len(rows),'grupos_investigados':len(groups),
        'sugestoes':suggestions,'cache':False,'catalogo_contas':account_catalog(records,accounts(company)),
        'limitacoes':['O plano de contas não foi fornecido: confirme se a conta histórica continua ativa.',
                       'Sem correspondência documental, nomes semelhantes não comprovam a identidade da operação.'],
        'ia':{'estado':'dispensada','chamadas_nesta_execucao':0}}
    if use_ai and any(s['situacao']=='MANUAL' and s['evidencias'] for s in suggestions):
        from app.classification_agent_ai import interpret
        result['ia'] = interpret(suggestions)
    result['token'] = secrets.token_urlsafe(32)
    with _lock:
        _cache[key]={'time':time.monotonic(),'result':copy.deepcopy(result),'company':company,'filehash':filehash,
                     'basehash':digest,'options':options}
        while len(_cache)>32:
            _cache.popitem(last=False)
    return result


def apply(company, content, filename, revisions, token, remember=False):
    with _lock:
        return _apply(company, content, filename, revisions, token, remember)


def _apply(company, content, filename, revisions, token, remember=False):
    with _lock:
        ticket = next((copy.deepcopy(v) for v in _cache.values() if secrets.compare_digest(v['result']['token'],token)),None)
    if not ticket or time.monotonic()-ticket['time']>_TTL:
        raise ValueError('Investigação expirada. Investigue as pendências novamente.')
    if ticket['company']!=int(company) or ticket['filehash']!=hashlib.sha256(content).hexdigest() or ticket['basehash']!=fingerprint(base.records(company)):
        raise ValueError('Arquivo, empresa ou base alterados. Investigue novamente antes de aplicar.')
    eligible_rows = {(r['_aba'],r['_linha']):r for r in eligible(company,content,ticket['options'])}
    validated=[]; seen=set()
    known = {'pago':set(), 'recebido':set()}
    bankaccounts = accounts(company)
    for record in base.records(company):
        bankaccount = str(bankaccounts.get(record.get('banco'), ''))
        if not bankaccount:
            continue
        debit, credit = str(record.get('debito','')), str(record.get('credito',''))
        if credit == bankaccount:
            known['pago'].add(debit)
        if debit == bankaccount:
            known['recebido'].add(credit)
    for revision in revisions:
        identity=(revision.get('_aba'),int(revision.get('_linha',0)))
        account=str(revision.get('Conta da contrapartida','')).strip()
        if identity not in eligible_rows or identity in seen:
            raise ValueError('Linha inválida, repetida ou já classificada.')
        if not re.fullmatch(r'[1-9][0-9]{0,14}',account) or account not in known[eligible_rows[identity]['natureza']] or account in {str(v) for v in accounts(company).values()}:
            raise ValueError('Escolha uma conta de contrapartida presente no histórico desta empresa.')
        if int(company)==242 and account in {'14','16','166'}:
            raise ValueError('A conta provisória não pode ser usada como classificação final.')
        seen.add(identity)
        validated.append({**eligible_rows[identity],'Conta da contrapartida':account})
    if not validated:
        raise ValueError('Aprove pelo menos uma linha para aplicar.')
    book, applied, learned = engine.aplicar_revisoes_inteligentes(content,pd.DataFrame(validated),filename,slug(company),accounts(company))
    summary={'aplicadas':applied,'aprendizados':0,'cobertura':coverage(content,book,slug(company),ticket['options'])}
    if remember:
        for record in learned:
            record['origem'] = 'sugestao_agente_aprovada'
        con=base._db()
        try:
            con.execute('BEGIN IMMEDIATE')
            base._save(con,company,learned)
            con.commit()
            summary['aprendizados']=len(learned)
        finally:
            con.close()
    return book,summary
