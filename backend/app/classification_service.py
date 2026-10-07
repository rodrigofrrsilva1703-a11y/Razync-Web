"""Persistent adapter for the complete original classification and review rules."""
from __future__ import annotations
import hashlib
import io
import json
import os
import sqlite3
import urllib.request
import urllib.parse
from pathlib import Path
import pandas as pd
from app import engine
from app.migration_services import accounts, slug

DB_PATH = Path(os.getenv('RAZYNC_DB_PATH', '/data/razync.db'))

# Perfis originais da Base Inteligente da empresa 242 no Streamlit.
# O frontend envia apenas a origem; as contas elegíveis são decididas aqui.
ELETRO_242_CLASSIFICATION_PROFILES = {
    'consolidada': {
        'modo_consolidado': True,
        'coluna_substituir': '',
        'valores_substituiveis': [],
    },
    'despesa': {
        'modo_consolidado': False,
        'coluna_substituir': 'debito',
        'valores_substituiveis': ['0', ''],
    },
    'fornecedor': {
        'modo_consolidado': False,
        'coluna_substituir': 'debito',
        'valores_substituiveis': ['166', '0', ''],
    },
    'recebido': {
        'modo_consolidado': False,
        'coluna_substituir': 'credito',
        'valores_substituiveis': ['166', '0', '14', '16', ''],
    },
    'francesinhas': {
        'modo_consolidado': False,
        'coluna_substituir': 'credito',
        'valores_substituiveis': [''],
    },
}

def resolve_classification_options(company, options=None):
    resolved = dict(options or {})
    if int(company) == 242:
        origem = str(resolved.get('origem_242') or '').strip().casefold()
        if origem:
            if origem not in ELETRO_242_CLASSIFICATION_PROFILES:
                raise ValueError('Modo de classificação inválido para a empresa 242.')
            resolved.update(ELETRO_242_CLASSIFICATION_PROFILES[origem])
    elif int(company) == 1408:
        # No Streamlit a filial tem apenas "Modelo Domínio consolidado".
        # Pagamentos e recebimentos substituem somente conta vazia/0.
        resolved.update({
            'modo_consolidado': True,
            'coluna_substituir': '',
            'valores_substituiveis': [],
        })
    return resolved

def _db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB_PATH, timeout=30)
    con.row_factory = sqlite3.Row
    con.executescript('''
      CREATE TABLE IF NOT EXISTS classification_records(
        company INTEGER NOT NULL, id TEXT NOT NULL, payload TEXT NOT NULL,
        PRIMARY KEY(company,id));
      CREATE TABLE IF NOT EXISTS classification_imports(
        company INTEGER NOT NULL, digest TEXT NOT NULL, PRIMARY KEY(company,digest));
    ''')
    return con

def _save(con, company, records):
    for record in records:
        record = dict(record)
        record['empresa'] = slug(company)
        identity = hashlib.sha256(f"{record['empresa']}|{record['banco']}|{record['assinatura']}|{record['debito']}|{record['credito']}".encode()).hexdigest()
        record['id'] = identity
        old = con.execute('SELECT payload FROM classification_records WHERE company=? AND id=?', (company, identity)).fetchone()
        if old:
            previous = json.loads(old['payload'])
            record['periodos'] = sorted(set(record.get('periodos', []) + previous.get('periodos', [])))
            record['ocorrencias'] = max(int(record.get('ocorrencias', 1)), int(previous.get('ocorrencias', 1)))
        con.execute('INSERT OR REPLACE INTO classification_records VALUES(?,?,?)', (company, identity, json.dumps(record, ensure_ascii=False)))

def records(company):
    con = _db()
    try:
        result = [json.loads(row['payload']) for row in con.execute('SELECT payload FROM classification_records WHERE company=?', (company,))]
        # Preserve all previously learned v0.3 records without rewriting the table.
        if con.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='patterns'").fetchone():
            known = {(r['banco'], r['assinatura'], r['debito'], r['credito']) for r in result}
            for row in con.execute('SELECT * FROM patterns WHERE company=?', (company,)):
                bank = row['bank']
                bank = 'bb' if company == 242 and bank == 'banco_brasil' else 'itau_512' if company == 1408 and bank == 'itau' else bank
                account = accounts(company).get(bank)
                if not account:
                    continue
                debit, credit = (row['counterpart'], account) if row['nature'] == 'pago' else (account, row['counterpart'])
                identity = (bank, row['signature'], debit, credit)
                if identity in known:
                    continue
                matching = next((r for r in result if (r['banco'], r['assinatura'], r['debito'], r['credito']) == identity), None)
                if matching:
                    matching['periodos'] = sorted(set(matching['periodos'] + [row['period']]))
                    matching['ocorrencias'] += row['occurrences']
                else:
                    result.append({'empresa': slug(company), 'banco': bank, 'assinatura': row['signature'],
                        'debito': debit, 'credito': credit, 'periodos': [row['period']],
                        'ocorrencias': row['occurrences'], 'exemplo_historico': row['example']})
        return result
    finally:
        con.close()

def status(company):
    rows = records(company)
    return {'patterns': len(rows), 'banks': len({r['banco'] for r in rows}),
            'periods': len({p for r in rows for p in r.get('periodos', [])})}

class SourceFile(io.BytesIO):
    def __init__(self, content, filename):
        super().__init__(content)
        self.name = filename

def learn(company, content, filename, data_inicial='', data_final=''):
    digest_source = content + f"|{data_inicial}|{data_final}".encode("utf-8") if (data_inicial or data_final) else content
    digest = hashlib.sha256(digest_source).hexdigest()
    con = _db()
    try:
        con.execute('BEGIN IMMEDIATE')
        if con.execute('SELECT 1 FROM classification_imports WHERE company=? AND digest=?', (company, digest)).fetchone():
            return 0
        if filename.lower().endswith('.json'):
            payload = json.loads(content)
            if not isinstance(payload, dict) or int(payload.get('company', -1)) != company:
                raise ValueError('A base importada pertence a outra empresa.')
            imported = payload.get('records', [])
            if not isinstance(imported, list):
                raise ValueError('Base de classificações inválida.')
        else:
            imported = engine.importar_arquivos_classificados([SourceFile(content, filename)], slug(company), accounts(company))
        if bool(data_inicial) != bool(data_final):
            raise ValueError('Informe as duas datas do período.')
        if data_inicial:
            inicio = pd.Timestamp(data_inicial).to_period('M')
            fim = pd.Timestamp(data_final).to_period('M')
            if fim < inicio:
                raise ValueError('A Data Final não pode ser anterior à Data Inicial.')
            meses = {str(periodo) for periodo in pd.period_range(inicio, fim, freq='M')}
            imported = [
                item for item in imported
                if any(str(periodo) in meses for periodo in (item.get('periodos') or []))
            ]
            if not imported:
                raise ValueError('Nenhum padrão revisado foi encontrado no período informado.')
        for item in imported:
            if not all(k in item for k in ('banco', 'assinatura', 'debito', 'credito', 'periodos')) or item['banco'] not in accounts(company):
                raise ValueError('Padrão inválido ou banco não configurado para esta empresa.')
        if not imported:
            raise ValueError('Nenhum padrão revisado foi encontrado. Envie um Modelo Domínio ou Razão com as contas preenchidas.')
        _save(con, company, imported)
        con.execute('INSERT INTO classification_imports VALUES(?,?)', (company, digest))
        con.commit()
        return sum(int(r.get('ocorrencias', 1)) for r in imported)
    finally:
        con.close()

def _limit_classified_workbook_to_period(original_content, classified_content, data_inicial='', data_final=''):
    if bool(data_inicial) != bool(data_final):
        raise ValueError('Informe as duas datas do período.')
    if not data_inicial:
        return classified_content, None
    inicio = pd.Timestamp(data_inicial).normalize()
    fim = pd.Timestamp(data_final).normalize()
    if fim < inicio:
        raise ValueError('A Data Final não pode ser anterior à Data Inicial.')

    from openpyxl import load_workbook
    original = load_workbook(io.BytesIO(original_content), data_only=False)
    result = load_workbook(io.BytesIO(classified_content), data_only=False)
    changed = 0

    for ws in result.worksheets:
        if ws.title not in original.sheetnames or 'retir' in engine.normalizar_texto(ws.title):
            continue
        source = original[ws.title]
        header = None
        columns = {}
        for row_number in range(1, min(source.max_row, 30) + 1):
            test = {
                engine.normalizar_texto(engine.texto_celula_seguro(source.cell(row_number, col).value)).strip(): col
                for col in range(1, source.max_column + 1)
            }
            if all(name in test for name in ('historico', 'debito', 'credito')):
                header, columns = row_number, test
                break
        if header is None:
            continue
        data_col = columns.get('data')
        if data_col is None:
            raise ValueError(f'A aba {ws.title} não possui coluna DATA para aplicar o período.')
        debit_col, credit_col = columns['debito'], columns['credito']

        for row_number in range(header + 1, source.max_row + 1):
            raw_date = source.cell(row_number, data_col).value
            parsed = pd.to_datetime(raw_date, dayfirst=True, errors='coerce')
            inside = not pd.isna(parsed) and inicio <= parsed.normalize() <= fim
            source_debit = source.cell(row_number, debit_col).value
            source_credit = source.cell(row_number, credit_col).value
            if not inside:
                ws.cell(row_number, debit_col).value = source_debit
                ws.cell(row_number, credit_col).value = source_credit
            elif (
                ws.cell(row_number, debit_col).value != source_debit
                or ws.cell(row_number, credit_col).value != source_credit
            ):
                changed += 1

    output = io.BytesIO()
    result.save(output)
    return output.getvalue(), changed


def classify(company, content, filename, options=None):
    options = resolve_classification_options(company, options)
    workbook, summary = engine.classificar_planilha_final(
        content, filename, records(company), accounts(company),
        empresa_classificacao=slug(company), coluna_substituir=options.get('coluna_substituir', ''),
        valores_substituiveis=options.get('valores_substituiveis', []),
        modo_consolidado_eletro_forte=bool(options.get('modo_consolidado', company in {242, 1408}))
    )
    workbook, changed = _limit_classified_workbook_to_period(
        content, workbook, options.get('data_inicial', ''), options.get('data_final', '')
    )
    if changed is not None:
        summary = dict(summary)
        summary['automaticos'] = changed
        summary['periodo_aplicado'] = {
            'data_inicial': options.get('data_inicial', ''),
            'data_final': options.get('data_final', ''),
        }
    return workbook, summary

def pending(company, content, data_inicial='', data_final=''):
    frame = engine.extrair_pendencias_revisao_inteligente(content, accounts(company))
    if bool(data_inicial) != bool(data_final):
        raise ValueError('Informe as duas datas do período.')
    if data_inicial and not frame.empty:
        inicio = pd.Timestamp(data_inicial).normalize()
        fim = pd.Timestamp(data_final).normalize()
        if fim < inicio:
            raise ValueError('A Data Final não pode ser anterior à Data Inicial.')
        datas = pd.to_datetime(frame['Data'], dayfirst=True, errors='coerce')
        frame = frame.loc[datas.between(inicio, fim, inclusive='both')].copy()
    return json.loads(frame.to_json(orient='records', date_format='iso', force_ascii=False))

def review(company, content, filename, revisions, remember=False):
    expected = pending(company, content)
    identities = {(r['_aba'], int(r['_linha'])): r for r in expected}
    validated = []
    for requested in revisions:
        identity = (requested.get('_aba'), int(requested.get('_linha', 0)))
        if identity not in identities:
            raise ValueError('O lançamento já está classificado ou não pertence à revisão deste arquivo.')
        row = dict(identities[identity])
        row['Conta da contrapartida'] = str(requested.get('Conta da contrapartida', '')).strip()
        validated.append(row)
    workbook, applied, new_records = engine.aplicar_revisoes_inteligentes(content, pd.DataFrame(validated), filename, slug(company), accounts(company))
    if remember:
        con = _db()
        try:
            _save(con, company, new_records)
            con.commit()
        finally:
            con.close()
    return workbook, {'aplicadas': applied, 'aprendizados': len(new_records) if remember else 0}

def import_original_online(company):
    """Explicit read-only copy; never write to the original database."""
    url = os.getenv('SUPABASE_URL', '').split('/rest/v1')[0].rstrip('/')
    key = os.getenv('SUPABASE_SERVICE_KEY', '')
    if not url.startswith('https://') or not key:
        raise ValueError('Configure SUPABASE_URL e SUPABASE_SERVICE_KEY somente no Railway para importar a base original.')
    query = urllib.parse.urlencode({'empresa': 'eq.' + slug(company), 'select': '*'})
    imported, offset = [], 0
    while True:
        request = urllib.request.Request(f'{url}/rest/v1/classificacoes_bancarias?{query}&limit=1000&offset={offset}', headers={'apikey': key, 'Authorization': f'Bearer {key}'})
        with urllib.request.urlopen(request, timeout=30) as response:
            batch = json.load(response)
        imported.extend(batch)
        if len(batch) < 1000:
            break
        offset += 1000
    return learn(company, json.dumps({'company': company, 'records': imported}).encode(), 'base-original.json')

def clear(company):
    con = _db()
    try:
        for table in ('classification_records', 'classification_imports'):
            con.execute(f'DELETE FROM {table} WHERE company=?', (company,))
        if con.execute("SELECT 1 FROM sqlite_master WHERE name='patterns'").fetchone():
            con.execute('DELETE FROM patterns WHERE company=?', (company,))
        con.commit()
    finally:
        con.close()
