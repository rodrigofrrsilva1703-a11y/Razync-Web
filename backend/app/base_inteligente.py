from __future__ import annotations

import io, os, re, sqlite3, unicodedata
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

from app.registry import CAPABILITIES
from razync.accede_1000 import identificar_conta_folha_accede_1000

DB_PATH = Path(os.getenv("RAZYNC_DB_PATH", "/data/razync.db"))


def _db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    con=sqlite3.connect(DB_PATH)
    con.execute("""
    CREATE TABLE IF NOT EXISTS patterns(
      company INTEGER NOT NULL,
      bank TEXT NOT NULL,
      signature TEXT NOT NULL,
      nature TEXT NOT NULL,
      counterpart TEXT NOT NULL,
      period TEXT NOT NULL,
      example TEXT,
      occurrences INTEGER NOT NULL DEFAULT 1,
      PRIMARY KEY(company,bank,signature,nature,counterpart,period)
    )
    """)
    return con


def norm(v):
    t=unicodedata.normalize("NFKD",str(v or ""))
    t="".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+"," ",t).casefold().strip()


def safe(v):
    if v is None: return ""
    try:
        if pd.isna(v): return ""
    except Exception: pass
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return str(v).strip()


def signature(history):
    texto=norm(history)
    texto=re.sub(r"\b(?:pagar|pagamento)\b","pago",texto)
    texto=re.sub(r"\b(?:receber|recebimento)\b","recebido",texto)
    nature="pago" if re.search(r"\bpago\b",texto) else "recebido" if re.search(r"\brecebido\b",texto) else "outro"
    m=re.search(r"empresa\s*:\s*(.*?)(?:\s+obs\s*:|$)",texto)
    empresa=m.group(1) if m else texto
    empresa=re.sub(r"[^a-z0-9]+"," ",empresa)
    tokens=[x for x in empresa.split() if not(any(c.isdigit() for c in x) and len(x)>=3)]
    empresa=" ".join(tokens).strip()
    obs=""
    om=re.search(r"\s+obs\s*:\s*(.*)$",texto)
    if om:
        bruto=re.sub(r"[^a-z0-9]+"," ",om.group(1))
        obs=" ".join(x for x in bruto.split() if x.isalpha() and x not in {"doc","documento"})
    return "|".join([x for x in [nature,empresa,obs] if x]) if empresa else ""


def bank_from(value, company=None):
    t=norm(value); d=re.sub(r"\D","",safe(value))
    if "btg" in t or "pactual" in t: return "btg"
    if "daycoval" in t: return "daycoval"
    if "fibra" in t: return "fibra"
    if "sicredi" in t: return "sicredi"
    if "santander" in t: return "santander"
    if "bradesco" in t: return "bradesco"
    if "inter" in t: return "inter"
    if "caixa" in t: return "caixa"
    if "banco do brasil" in t or t=="bb": return "banco_brasil"
    if "itau" in t:
        if company==242:
            if "509" in d or "509" in t: return "itau_509"
            if "508" in d or "508" in t: return "itau_508"
        return "itau"
    return ""


def bank_account(company, bank):
    banks=CAPABILITIES.get(int(company),{}).get("banks",{})
    if bank in banks: return str(banks[bank] or "")
    if bank=="itau" and company==242: return ""
    return ""


def _find_header(raw):
    for i in range(min(30,len(raw))):
        names=[norm(x) for x in raw.iloc[i].tolist()]
        if all(x in names for x in ["historico","debito","credito"]):
            return i,names
    return None,None


def learn(company:int, content:bytes, filename:str):
    xls=pd.ExcelFile(io.BytesIO(content))
    learned=0
    con=_db()
    try:
        for sheet in xls.sheet_names:
            raw=pd.read_excel(xls,sheet_name=sheet,header=None,dtype=object)
            idx,names=_find_header(raw)
            if idx is None: continue
            df=raw.iloc[idx+1:].copy(); df.columns=[safe(x) for x in raw.iloc[idx].tolist()]
            mp={norm(c):c for c in df.columns}
            ch,cd,cc=mp.get("historico"),mp.get("debito"),mp.get("credito")
            cdata,cdesc=mp.get("data"),mp.get("descricao")
            if not ch or not cd or not cc: continue
            for _,row in df.iterrows():
                hist=safe(row[ch]); debit=safe(row[cd]); credit=safe(row[cc])
                if not hist or not debit or not credit: continue
                bank=bank_from(row[cdesc] if cdesc else sheet,company) or bank_from(sheet,company) or bank_from(filename,company)
                if not bank: continue
                acct=bank_account(company,bank)
                if not acct: continue
                if credit==acct and debit:
                    nature,counter="pago",debit
                elif debit==acct and credit:
                    nature,counter="recebido",credit
                else:
                    continue
                sig=signature(hist)
                if not sig: continue
                dt=pd.to_datetime(row[cdata],dayfirst=True,errors="coerce") if cdata else pd.NaT
                period=dt.strftime("%Y-%m") if not pd.isna(dt) else norm(filename)[:30]
                con.execute("""INSERT INTO patterns(company,bank,signature,nature,counterpart,period,example,occurrences)
                  VALUES(?,?,?,?,?,?,?,1)
                  ON CONFLICT(company,bank,signature,nature,counterpart,period)
                  DO UPDATE SET occurrences=occurrences+1, example=excluded.example""",
                  (company,bank,sig,nature,counter,period,hist[:500]))
                learned+=1
        con.commit()
    finally:
        con.close()
    return learned


def status(company:int):
    con=_db()
    try:
        row=con.execute("SELECT COUNT(*),COUNT(DISTINCT bank),COUNT(DISTINCT period) FROM patterns WHERE company=?",(company,)).fetchone()
        return {"patterns":row[0],"banks":row[1],"periods":row[2]}
    finally: con.close()


def classify(company:int,content:bytes,filename:str):
    wb=load_workbook(io.BytesIO(content))
    con=_db(); summary={"automaticos":0,"somente_banco":0,"fixas_1000":0,"ja_preenchidos":0}
    try:
      for ws in wb.worksheets:
        header=None; mp={}
        for r in range(1,min(ws.max_row,30)+1):
            test={norm(ws.cell(r,c).value):c for c in range(1,ws.max_column+1)}
            if all(x in test for x in ["historico","debito","credito"]):
                header,mp=r,test; break
        if not header: continue
        bank_sheet=bank_from(ws.title,company) or bank_from(filename,company)
        for r in range(header+1,ws.max_row+1):
            hist=safe(ws.cell(r,mp["historico"]).value)
            if not hist: continue
            debit=safe(ws.cell(r,mp["debito"]).value); credit=safe(ws.cell(r,mp["credito"]).value)
            if debit and credit:
                summary["ja_preenchidos"]+=1; continue
            desc=safe(ws.cell(r,mp.get("descricao",0)).value) if mp.get("descricao") else ""
            bank=bank_from(desc,company) or bank_sheet
            acct=bank_account(company,bank)
            if not acct: continue
            val=0.0
            if mp.get("valor"):
                try: val=float(ws.cell(r,mp["valor"]).value or 0)
                except Exception: val=0
            if val>0 and not debit:
                ws.cell(r,mp["debito"]).value=int(acct) if acct.isdigit() else acct
                debit=acct; summary["somente_banco"]+=1
            elif val<0 and not credit:
                ws.cell(r,mp["credito"]).value=int(acct) if acct.isdigit() else acct
                credit=acct; summary["somente_banco"]+=1

            if company==1000 and val<0 and not debit:
                fixed=identificar_conta_folha_accede_1000(hist)
                if fixed:
                    ws.cell(r,mp["debito"]).value=int(fixed) if fixed.isdigit() else fixed
                    debit=fixed; summary["fixas_1000"]+=1

            sig=signature(hist)
            nature="pago" if val<0 else "recebido" if val>0 else ""
            if not sig or not nature: continue
            rows=con.execute("""SELECT counterpart,COUNT(DISTINCT period) periods
              FROM patterns WHERE company=? AND bank=? AND signature=? AND nature=?
              GROUP BY counterpart""",(company,bank,sig,nature)).fetchall()
            eligible=[(cp,p) for cp,p in rows if p>=3]
            if len(eligible)!=1: continue
            cp=eligible[0][0]
            if nature=="pago" and not debit:
                ws.cell(r,mp["debito"]).value=int(cp) if cp.isdigit() else cp; summary["automaticos"]+=1
            elif nature=="recebido" and not credit:
                ws.cell(r,mp["credito"]).value=int(cp) if cp.isdigit() else cp; summary["automaticos"]+=1
    finally: con.close()
    out=io.BytesIO(); wb.save(out)
    return out.getvalue(),summary
