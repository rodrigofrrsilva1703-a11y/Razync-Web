"""Edit accounting model rows while preserving the original XLSX workbook."""
import io
import math
import re
import unicodedata
from copy import copy
from datetime import date, datetime
from openpyxl import Workbook, load_workbook


def norm(value):
    return ''.join(c for c in unicodedata.normalize('NFD', str(value or '')).upper() if not unicodedata.combining(c)).strip()


def load(content):
    if content.startswith(b'PK'):
        book = load_workbook(io.BytesIO(content))
    else:
        import pandas as pd
        book = Workbook()
        book.remove(book.active)
        with pd.ExcelFile(io.BytesIO(content)) as source:
            for name in source.sheet_names:
                sheet = book.create_sheet(name)
                frame = pd.read_excel(source, sheet_name=name, header=None, dtype=object)
                for row in frame.itertuples(index=False, name=None):
                    sheet.append([None if pd.isna(v) else v for v in row])
    return book


def layout(sheet):
    for row in sheet.iter_rows(max_row=min(sheet.max_row, 30)):
        names = [norm(c.value) for c in row]
        if all(name in names for name in ['DATA', 'VALOR', 'HISTORICO']):
            last = max(i for i, name in enumerate(names) if name) + 1
            if sheet.max_row - row[0].row > 20000 or last > 80:
                raise ValueError('O editor aceita até 20.000 linhas e 80 colunas por aba. Use o Excel para arquivos maiores.')
            return row[0].row, last
    return None


def display(value):
    if isinstance(value, (datetime, date)):
        return value.strftime('%d/%m/%Y')
    if value is None:
        return ''
    return value


def inspect(content):
    book = load(content)
    try:
        result = []
        for sheet in book:
            spec = layout(sheet)
            if not spec:
                continue
            header, width = spec
            if any(c.data_type == 'f' for row in sheet for c in row):
                raise ValueError('Esta aba contém fórmulas. Para preservar os cálculos, envie o Modelo Final com os valores calculados.')
            if sheet.protection.sheet or any(r.max_row > header for r in sheet.merged_cells.ranges):
                raise ValueError('A área de lançamentos está protegida ou contém células mescladas. Envie um Modelo Final sem essas restrições.')
            result.append({'name': sheet.title, 'header': header,
                'columns': [str(sheet.cell(header, c).value or '') for c in range(1, width + 1)],
                'rows': [[display(sheet.cell(r, c).value) for c in range(1, width + 1)] for r in range(header + 1, sheet.max_row + 1)]})
        if not result:
            raise ValueError('Nenhuma aba com DATA, VALOR e HISTÓRICO foi encontrada.')
        return {'sheets': result}
    finally:
        book.close()


def value_for(label, value):
    label = norm(label)
    if value is None or str(value).strip() == '':
        return None
    text = str(value).strip()
    if label == 'DATA':
        for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                pass
        raise ValueError('Data inválida. Use DD/MM/AAAA.')
    if label == 'VALOR':
        if isinstance(value, bool):
            raise ValueError('Valor inválido.')
        if isinstance(value, (int, float)):
            number = float(value)
        else:
            text = text.replace('R$', '').replace(' ', '')
            if ',' in text:
                text = text.replace('.', '').replace(',', '.')
            elif re.fullmatch(r'-?\d{1,3}(?:\.\d{3})+', text):
                text = text.replace('.', '')
            try:
                number = float(text)
            except ValueError as exc:
                raise ValueError('Valor inválido. Use, por exemplo, -1.234,56.') from exc
        if not math.isfinite(number):
            raise ValueError('Valor inválido.')
        return number
    # Treat histories and account codes as literal text, never spreadsheet formulas.
    return text


def apply(content, operations):
    if not isinstance(operations, list) or len(operations) > 50000:
        raise ValueError('Quantidade de alterações inválida.')
    metadata = {s['name']: s for s in inspect(content)['sheets']}
    book = load(content)
    try:
        for op in operations:
            if not isinstance(op, dict) or op.get('sheet') not in metadata:
                raise ValueError('Aba inválida.')
            sheet = book[op['sheet']]
            spec = metadata[sheet.title]
            row = op.get('row')
            if type(row) is not int or not spec['header'] < row <= sheet.max_row + 1:
                raise ValueError('Linha inválida. Os cabeçalhos não podem ser alterados.')
            kind = op.get('kind')
            if kind == 'insert':
                if sheet.max_row - spec['header'] >= 20000:
                    raise ValueError('Limite de 20.000 linhas por aba.')
                source = max(spec['header'] + 1, min(row, sheet.max_row))
                styles = [copy(sheet.cell(source, c)._style) for c in range(1, len(spec['columns']) + 1)]
                sheet.insert_rows(row)
                for c, style in enumerate(styles, 1):
                    sheet.cell(row, c)._style = style
            elif kind == 'delete':
                if row > sheet.max_row:
                    raise ValueError('Linha inexistente.')
                sheet.delete_rows(row)
            elif kind == 'set':
                col = op.get('col')
                if type(col) is not int or not 1 <= col <= len(spec['columns']):
                    raise ValueError('Coluna inválida.')
                cell = sheet.cell(row, col)
                cell.value = value_for(spec['columns'][col - 1], op.get('value'))
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                if norm(spec['columns'][col - 1]) == 'DATA':
                    cell.number_format = 'dd/mm/yyyy'
            else:
                raise ValueError('Operação inválida.')
        output = io.BytesIO()
        book.save(output)
        return output.getvalue()
    finally:
        book.close()
