from __future__ import annotations
import io
import json
import os
import hmac
from pathlib import Path
from pydantic import BaseModel
from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from fastapi.concurrency import run_in_threadpool
from app import classification_service as base
from app import migration_services as services

router = APIRouter(prefix='/api/v1')

@router.post('/model-editor/open')
async def open_model_editor(file: UploadFile = File(...)):
    from app.model_editor import inspect
    try:
        return await run_in_threadpool(inspect, await file.read())
    except Exception as exc:
        raise HTTPException(422, str(exc)) from exc

@router.post('/model-editor/apply')
async def apply_model_editor(file: UploadFile = File(...), operations_json: str = Form('[]')):
    from app.model_editor import apply
    try:
        content = await run_in_threadpool(apply, await file.read(), json.loads(operations_json))
        return download(content, 'MODELO_CORRIGIDO.xlsx')
    except Exception as exc:
        raise HTTPException(422, str(exc)) from exc


def workbook_preview(content):
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import from_excel
    import datetime
    book = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheets = []
    try:
        for sheet in book:
            if sheet.title.casefold() == 'principal':
                continue
            columns = None
            display, count = [], 0
            incoming, outgoing = 0, 0
            for row in sheet.iter_rows(values_only=True):
                if columns is None:
                    if not any(str(c or '').strip().upper() == 'HISTÓRICO' for c in row):
                        continue
                    columns = [str(c or '') for c in row]
                    index = next((i for i,c in enumerate(columns) if c.strip().upper() == 'VALOR'), None)
                    date_index = next((i for i,c in enumerate(columns) if c.strip().upper() == 'DATA'), None)
                    continue
                if not any(c is not None for c in row):
                    continue
                count += 1
                if index is not None and isinstance(row[index], (int,float)):
                    amount = float(row[index])
                    if amount > 0: incoming += amount
                    if amount < 0: outgoing += amount
                if len(display) < 500:
                    formatted = []
                    for i,c in enumerate(row):
                        if i == date_index and isinstance(c,(int,float)):
                            c = from_excel(c, book.epoch)
                        formatted.append(c.strftime('%d/%m/%Y') if isinstance(c,(datetime.datetime,datetime.date)) else c)
                    display.append(formatted)
            if columns is not None:
                sheets.append({'name':sheet.title,'count':count,'columns':columns,'rows':display,
                    'entradas':round(incoming,2),'saidas':round(-outgoing,2)})
    finally:
        book.close()
    return sheets

@router.post('/modelo-preview')
async def model_preview(file: UploadFile = File(...)):
    try:
        sheets = await run_in_threadpool(workbook_preview, await file.read())
        return {'sheets': sheets, 'diagnostics': []}
    except Exception as exc:
        raise HTTPException(422, str(exc)) from exc

@router.post('/workflow/{company_code}/preview')
async def workflow_preview(company_code: int, roles_json: str = Form(...), options_json: str = Form('{}'), files: list[UploadFile] = File(...)):
    try:
        names, options = json.loads(roles_json), json.loads(options_json)
        if not isinstance(names, list) or len(names) != len(files) or not isinstance(options, dict):
            raise ValueError('Configuração dos arquivos inválida.')
        roles = {}
        for name, upload in zip(names, files):
            roles.setdefault(str(name), []).append((upload.filename or 'arquivo.xlsx', await upload.read()))
        content, filename = await run_in_threadpool(services.workflow, company_code, roles, options)
        sheets = await run_in_threadpool(workbook_preview, content)
        tables = await run_in_threadpool(services.diagnostics, company_code, roles, options)
        diagnostics = [{'name':name,'columns':list(frame.columns),'rows':json.loads(frame.head(500).to_json(orient='values',date_format='iso',force_ascii=False)), 'count':len(frame)} for name,frame in tables.items()]
        return {'filename':filename,'sheets':sheets,'diagnostics':diagnostics}
    except Exception as exc:
        raise HTTPException(422, str(exc)) from exc

def download(content, filename, summary=None, mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'):
    headers = {'Content-Disposition': f'attachment; filename="{filename}"'}
    if summary is not None:
        headers['X-Razync-Summary'] = json.dumps(summary, ensure_ascii=True)
    return StreamingResponse(io.BytesIO(content), media_type=mime, headers=headers)

def require_admin(authorization: str = Header('')):
    credentials = [value for value in (os.getenv('CLASSIFICATION_ADMIN_PASSWORD'), os.getenv('RAZYNC_ACCESS_TOKEN')) if value]
    if not credentials:
        raise HTTPException(503, 'A senha administrativa não está configurada no servidor.')
    valid = False
    for expected in credentials:
        valid = hmac.compare_digest(authorization.encode('utf-8'), f'Bearer {expected}'.encode('utf-8')) or valid
    if not valid:
        raise HTTPException(401, 'Informe a chave de acesso administrativa.')

@router.post('/companies', dependencies=[Depends(require_admin)], status_code=201)
def create_company(payload: dict):
    from app.company_service import create
    try:
        row = create(payload.get('codigo'), payload.get('nome'), payload.get('regime'))
        row['capabilities'] = {'status': 'catalog_only', 'tools': [], 'banks': {}}
        return row
    except FileExistsError as exc:
        raise HTTPException(409, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@router.get('/migration-status')
def migration_status():
    manifest = Path(__file__).resolve().parents[1] / 'resources' / 'engine_manifest.json'
    data = json.loads(manifest.read_text(encoding='utf-8'))
    return {'reference': data['reference_repository'], 'extracted_functions': len(data['functions']),
            'validation': 'pending_real_files',
            'original_base_import_configured': bool(os.getenv('SUPABASE_URL') and os.getenv('SUPABASE_SERVICE_KEY'))}

@router.post('/tasks/importar-status', dependencies=[Depends(require_admin)])
async def import_task_statuses(file: UploadFile = File(...)):
    try:
        from app.tasks import import_company_statuses
        rows = json.loads(await file.read())
        if not isinstance(rows,list):
            raise ValueError('Snapshot de tarefas inválido.')
        return {'imported':await run_in_threadpool(import_company_statuses,rows)}
    except Exception as exc:
        raise HTTPException(422,str(exc)) from exc

@router.post('/workflow/{company_code}/reports')
async def report_package(company_code: int, roles_json: str = Form(...), options_json: str = Form('{}'), files: list[UploadFile] = File(...)):
    if services.CAPABILITIES.get(company_code, {}).get('workflow') != 'advanced':
        raise HTTPException(404, 'Empresa sem fluxo de arquivos complementares.')
    try:
        names, options = json.loads(roles_json), json.loads(options_json)
        if not isinstance(names, list) or len(names) != len(files) or not isinstance(options, dict):
            raise ValueError('Configuração dos arquivos inválida.')
        roles = {}
        for name, upload in zip(names, files):
            roles.setdefault(str(name), []).append((upload.filename or 'arquivo.xlsx', await upload.read()))
        reports = await run_in_threadpool(services.workflow_reports, company_code, roles, options)
        return download(services.package(reports), f'RAZYNC_{company_code}_RELATORIOS.zip', {'arquivos': list(reports)}, 'application/zip')
    except Exception as exc:
        raise HTTPException(422, str(exc)) from exc

@router.post('/francesinhas/{company_code}')
async def standalone_francesinhas(
    company_code: int,
    files: list[UploadFile] = File(...),
    data_inicial: str = Form(''),
    data_final: str = Form(''),
):
    if company_code not in {242, 1408}:
        raise HTTPException(404, 'Francesinhas disponíveis para Eletro Forte 242 e 1408.')
    try:
        uploads = [(u.filename or 'francesinhas.zip', await u.read()) for u in files]
        book, summary = await run_in_threadpool(
            services.francesinhas, company_code, uploads, data_inicial, data_final
        )
        return download(book, f'ELETRO_FORTE_{company_code}_FRANCESINHAS.xlsx', summary)
    except Exception as exc:
        raise HTTPException(422, str(exc)) from exc

@router.get('/base-inteligente/{company_code}/exportar', dependencies=[Depends(require_admin)])
def export_base(company_code: int):
    content = json.dumps({'version': 1, 'company': company_code, 'records': base.records(company_code)}, ensure_ascii=False).encode()
    return download(content, f'RAZYNC_{company_code}_BASE.json', mime='application/json')

@router.post('/base-inteligente/{company_code}/importar-original', dependencies=[Depends(require_admin)])
def import_base(company_code: int):
    try:
        return {'learned': base.import_original_online(company_code), **base.status(company_code)}
    except Exception as exc:
        raise HTTPException(422, str(exc)) from exc

class BaseSliceDeletion(BaseModel):
    bank: str
    period: str
    token: str


@router.get('/base-inteligente/{company_code}/recortes', dependencies=[Depends(require_admin)])
def base_slices(company_code: int):
    from app import base_maintenance
    return base_maintenance.slices(company_code)


@router.get('/base-inteligente/{company_code}/exclusao-previa', dependencies=[Depends(require_admin)])
def base_deletion_preview(company_code: int, bank: str, period: str):
    from app import base_maintenance
    try:
        return base_maintenance.preview(company_code, bank, period)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@router.post('/base-inteligente/{company_code}/excluir-recorte', dependencies=[Depends(require_admin)])
def base_delete_slice(company_code: int, payload: BaseSliceDeletion):
    from app import base_maintenance
    try:
        return base_maintenance.delete(company_code, payload.bank, payload.period, payload.token)
    except base_maintenance.StaleSelection as exc:
        raise HTTPException(409, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@router.delete('/base-inteligente/{company_code}', dependencies=[Depends(require_admin)])
def clear_base(company_code: int):
    base.clear(company_code)
    return {'ok': True}

@router.post('/base-inteligente/{company_code}/pendencias')
async def review_pending(
    company_code: int,
    file: UploadFile = File(...),
    data_inicial: str = Form(''),
    data_final: str = Form(''),
):
    try:
        return {'pendencias': base.pending(company_code, await file.read(), data_inicial, data_final)}
    except Exception as exc:
        raise HTTPException(422, str(exc)) from exc

@router.post('/base-inteligente/{company_code}/revisar')
async def apply_review(company_code: int, file: UploadFile = File(...), revisoes_json: str = Form(...), aprender: bool = Form(False)):
    try:
        revisions = json.loads(revisoes_json)
        if not isinstance(revisions, list):
            raise ValueError('Revisões inválidas.')
        book, summary = base.review(company_code, await file.read(), file.filename or 'modelo.xlsx', revisions, aprender)
        return download(book, f'RAZYNC_{company_code}_REVISADO.xlsx', summary)
    except Exception as exc:
        raise HTTPException(422, str(exc)) from exc

@router.post('/modelo-dominio-txt')
async def export_txt(
    file: UploadFile = File(...),
    data_inicial: str = Form(''),
    data_final: str = Form(''),
):
    try:
        from app import engine
        import pandas as pd
        def convert(content):
            book = pd.ExcelFile(io.BytesIO(content))
            frames = []
            for name in book.sheet_names:
                if name.casefold() == 'principal' or 'retir' in name.casefold():
                    continue
                raw = pd.read_excel(book, sheet_name=name, header=None)
                header = next((i for i in range(min(30,len(raw))) if {'DATA','VALOR','HISTÓRICO'} <= {str(v or '').strip().upper() for v in raw.iloc[i]}), None)
                if header is None:
                    continue
                frame = raw.iloc[header+1:].copy()
                frame.columns = [str(v or '').strip().upper() for v in raw.iloc[header]]
                if all(c in frame for c in services.COLUNAS):
                    frame = frame[services.COLUNAS].dropna(subset=['DATA','VALOR'])
                    frame['DATA'] = frame['DATA'].map(lambda v: pd.to_datetime(v,unit='D',origin='1899-12-30') if isinstance(v,(int,float)) else v)
                    parsed = pd.to_datetime(frame['DATA'],dayfirst=True,errors='coerce')
                    if bool(data_inicial) != bool(data_final):
                        raise ValueError('Informe as duas datas do período.')
                    if data_inicial:
                        inicio = pd.Timestamp(data_inicial).normalize()
                        fim = pd.Timestamp(data_final).normalize()
                        if fim < inicio:
                            raise ValueError('A Data Final não pode ser anterior à Data Inicial.')
                        mask = parsed.between(inicio, fim, inclusive='both')
                        frame = frame.loc[mask].copy()
                        parsed = parsed.loc[mask]
                    frame['DATA'] = parsed.dt.strftime('%d/%m/%Y')
                    frames.append(frame)
            if not frames:
                raise ValueError('Nenhuma aba do Modelo Domínio encontrada.')
            return engine.gerar_txt_dominio(pd.concat(frames,ignore_index=True)).encode('utf-8')
        content = await run_in_threadpool(convert, await file.read())
        return download(content, 'RAZYNC_MODELO_DOMINIO.txt', mime='text/plain; charset=utf-8')
    except Exception as exc:
        raise HTTPException(422,str(exc)) from exc
