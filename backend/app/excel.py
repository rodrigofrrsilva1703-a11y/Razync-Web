from __future__ import annotations

import base64
import io
import re
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

COLUNAS = ["DESCRIÇÃO", "DATA", "VALOR", "DÉBITO", "CRÉDITO", "HISTÓRICO"]
RESOURCE = Path(__file__).resolve().parents[1] / "resources" / "modelo_dominio.b64"


def _norm(value: object) -> str:
    text = str(value or "").upper()
    table = str.maketrans("ÁÀÃÂÉÊÍÓÔÕÚÇ", "AAAAEEIOOOUC")
    return re.sub(r"[^A-Z0-9]", "", text.translate(table))


def _template_bytes() -> bytes:
    return base64.b64decode(RESOURCE.read_text(encoding="utf-8").strip())


def gerar_modelo_abas(dados: dict[str, pd.DataFrame]) -> bytes:
    wb = load_workbook(io.BytesIO(_template_bytes()))
    ws_base = wb[wb.sheetnames[0]]

    header_row = None
    mapping = None
    expected = [_norm(c) for c in COLUNAS]
    for row in range(1, min(ws_base.max_row, 25) + 1):
        m = {_norm(ws_base.cell(row, col).value): col for col in range(1, ws_base.max_column + 1)}
        if all(k in m for k in expected):
            header_row, mapping = row, m
            break
    if header_row is None or mapping is None:
        raise ValueError("Cabeçalho do Modelo Domínio não encontrado.")

    first_data = header_row + 1
    styles = {
        col: {
            "style": copy(ws_base.cell(first_data, mapping[_norm(col)])._style),
            "number_format": ws_base.cell(first_data, mapping[_norm(col)]).number_format,
        }
        for col in COLUNAS
    }

    created = []
    for sheet_name, df in dados.items():
        ws = wb.copy_worksheet(ws_base)
        safe = str(sheet_name)[:31] or "Banco"
        candidate = safe
        n = 2
        while candidate in created or candidate in wb.sheetnames:
            candidate = f"{safe[:27]} {n}"
            n += 1
        ws.title = candidate
        created.append(candidate)

        if ws.max_row > header_row:
            for r in range(first_data, ws.max_row + 1):
                for c in range(1, ws.max_column + 1):
                    ws.cell(r, c).value = None

        frame = df.copy()
        for col in COLUNAS:
            if col not in frame.columns:
                frame[col] = ""

        for idx, record in enumerate(frame[COLUNAS].to_dict("records"), start=first_data):
            for col in COLUNAS:
                cell = ws.cell(idx, mapping[_norm(col)])
                value = record.get(col, "")
                if pd.isna(value):
                    value = ""
                if col == "DATA" and value not in ("", None):
                    dt = pd.to_datetime(value, dayfirst=True, errors="coerce")
                    value = "" if pd.isna(dt) else dt.to_pydatetime()
                cell.value = value
                cell._style = copy(styles[col]["style"])
                cell.number_format = styles[col]["number_format"]

    wb.remove(ws_base)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
