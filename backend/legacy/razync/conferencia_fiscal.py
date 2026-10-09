"""Conferência inteligente entre acumuladores fiscais e razão do Domínio."""

from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


VERSAO_LEITOR_FISCAL = "2026.09.23.7-filial-relatorio-individual"


def _moeda(valor) -> str:
    numero = float(valor or 0)
    return f"R$ {numero:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _renderizar_conferencia_fiscal_contabil(prefixo: str, empresa: str) -> None:
    """Renderiza a conferência fiscal genérica com estado isolado por empresa."""
    import streamlit as st

    chave = re.sub(r"[^a-z0-9_]+", "_", str(prefixo).lower()).strip("_")
    base = f"fiscal_{chave}"
    codigo_empresa = ""
    codigo_no_nome = re.match(r"\s*(\d+)\s*-", str(empresa))
    if codigo_no_nome:
        codigo_empresa = codigo_no_nome.group(1)
    st.markdown("### Conferência Fiscal × Contábil")
    st.caption(
        f"Empresa: {empresa}. Envie o Resumo por Acumulador e o Razão do mesmo período. "
        "Somente acumuladores com conta preenchida são conferidos; movimentos não fiscais "
        "são separados em alertas."
    )
    col_fiscal, col_razao = st.columns(2)
    with col_fiscal:
        arquivo_acumuladores = st.file_uploader(
            "Relatório de acumuladores do Domínio", type=["xls", "xlsx"],
            key=f"{base}_acumuladores",
        )
    with col_razao:
        arquivo_razao = st.file_uploader(
            "Razão com todas as contas", type=["xls", "xlsx"], key=f"{base}_razao",
        )
    if arquivo_acumuladores is None or arquivo_razao is None:
        st.info("Envie os dois relatórios. A conferência começará automaticamente.")
        return

    import hashlib
    assinatura = hashlib.sha256(
        arquivo_acumuladores.getvalue() + arquivo_razao.getvalue()
        + VERSAO_LEITOR_FISCAL.encode("utf-8") + codigo_empresa.encode("utf-8")
    ).hexdigest()
    if st.session_state.get(f"{base}_assinatura") != assinatura:
        try:
            with st.spinner("Cruzando acumuladores, contas e lançamentos do razão..."):
                resultado = processar_conferencia(
                    arquivo_acumuladores.getvalue(), arquivo_acumuladores.name,
                    arquivo_razao.getvalue(), arquivo_razao.name,
                    filial_alvo=codigo_empresa,
                )
            st.session_state[f"{base}_resultado"] = resultado
            st.session_state[f"{base}_assinatura"] = assinatura
            st.session_state.pop(f"{base}_erro", None)
        except Exception as erro:
            st.session_state[f"{base}_erro"] = str(erro)
            st.session_state.pop(f"{base}_resultado", None)

    if st.session_state.get(f"{base}_erro"):
        st.error("Não foi possível concluir a conferência: " + st.session_state[f"{base}_erro"])
    resultado = st.session_state.get(f"{base}_resultado")
    if not isinstance(resultado, dict):
        return
    resumo = resultado.get("resumo", pd.DataFrame())
    detalhes = resultado.get("detalhes", pd.DataFrame())
    if resumo.empty:
        st.warning("Nenhuma conta pôde ser comparada.")
        return

    filial_aplicada = resultado.get("filial_aplicada", "")
    filiais_encontradas = resultado.get("filiais_encontradas", [])
    if filial_aplicada:
        st.success(
            f"Razão consolidado identificado. Conferência filtrada pela filial "
            f"{filial_aplicada}. Filiais presentes no arquivo: "
            f"{', '.join(filiais_encontradas)}."
        )

    conferidas = int(resumo["SITUAÇÃO"].astype(str).str.startswith("CONFERE").sum())
    alertas = int(resumo["LANÇAMENTOS EXTRAS"].sum())
    revisar = int((resumo["SITUAÇÃO"] == "REVISAR").sum())
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Contas analisadas", len(resumo))
    m2.metric("Fiscal conferido", conferidas)
    m3.metric("Lançamentos em alerta", alertas)
    m4.metric("Contas para revisar", revisar)

    periodo = resultado.get("periodo_fiscal", {})
    inicio = pd.to_datetime(periodo.get("inicio"), errors="coerce")
    periodo_nome = inicio.strftime("%m%Y") if pd.notna(inicio) else "PERIODO_ANALISADO"
    nome_empresa = re.sub(r"[^A-Za-z0-9]+", "_", str(empresa)).strip("_")[:45]
    st.download_button(
        "Baixar relatório completo da conferência",
        data=gerar_relatorio_excel(resultado),
        file_name=f"{nome_empresa}_CONFERENCIA_FISCAL_{periodo_nome}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True, key=f"{base}_download",
    )
    if revisar:
        st.error(f"{revisar} conta(s) possuem diferença fiscal e precisam de revisão.")
    elif alertas:
        st.warning("Os valores fiscais conferem, mas existem lançamentos contábeis adicionais para revisar.")
    else:
        st.success("Todas as contas conferem e não foram encontrados lançamentos adicionais.")

    aba_visao, aba_alertas, aba_todos = st.tabs([
        "Visão geral", f"Alertas ({alertas})", "Todos os lançamentos"
    ])
    with aba_visao:
        ordem = {"REVISAR": 0, "AUSENTE NO CONTÁBIL": 1, "CONFERE COM ALERTAS": 2, "CONFERE": 3}
        resumo_ordenado = resumo.assign(
            _ORDEM=resumo["SITUAÇÃO"].map(ordem).fillna(9)
        ).sort_values(["_ORDEM", "CONTA"])
        for _, item in resumo_ordenado.iterrows():
            conta, situacao = str(item["CONTA"]), str(item["SITUAÇÃO"])
            extras = int(item["LANÇAMENTOS EXTRAS"])
            titulo = f"Conta {conta} · {situacao}" + (f" · {extras} alerta(s)" if extras else "")
            with st.expander(titulo, expanded=situacao != "CONFERE"):
                c1, c2, c3, c4 = st.columns(4)
                c1.metric("Valor fiscal", _moeda(item["VALOR FISCAL"]))
                c2.metric("Contábil compatível", _moeda(item["CONTÁBIL COMPATÍVEL"]))
                c3.metric("Diferença fiscal", _moeda(item["DIFERENÇA FISCAL"]))
                c4.metric("Total movimentado", _moeda(item["TOTAL DA CONTA"]))
                coluna_analisada = "Crédito" if item["TIPO"] == "SAÍDAS" else "Débito"
                total_coluna = (
                    item["TOTAL CRÉDITOS"] if coluna_analisada == "Crédito"
                    else item["TOTAL DÉBITOS"]
                )
                st.caption(
                    f"Coluna analisada: {coluna_analisada} · "
                    f"Total líquido: {_moeda(total_coluna)} · "
                    f"Estornos/redutores: {_moeda(item['ESTORNOS/REDUTORES'])}"
                )
                movimentos = detalhes[detalhes["CONTA"].astype(str).eq(conta)].copy()
                if not movimentos.empty:
                    movimentos["DATA"] = pd.to_datetime(movimentos["DATA"]).dt.strftime("%d/%m/%Y")
                    st.dataframe(
                        movimentos[[
                            "DATA", "HISTÓRICO", "CONTRAPARTIDA", "DÉBITO",
                            "CRÉDITO", "NATUREZA", "VALOR", "CLASSIFICAÇÃO"
                        ]], use_container_width=True, hide_index=True
                    )
    with aba_alertas:
        quadro = detalhes[detalhes["CLASSIFICAÇÃO"].eq("ALERTA - NÃO FISCAL")].copy()
        if quadro.empty:
            st.success("Nenhum lançamento adicional foi encontrado nas contas conferidas.")
        else:
            quadro["DATA"] = pd.to_datetime(quadro["DATA"]).dt.strftime("%d/%m/%Y")
            st.dataframe(
                quadro[[
                    "CONTA", "DATA", "HISTÓRICO", "CONTRAPARTIDA", "DÉBITO",
                    "CRÉDITO", "NATUREZA", "VALOR"
                ]], use_container_width=True, hide_index=True
            )
    with aba_todos:
        quadro = detalhes.copy()
        if not quadro.empty:
            quadro["DATA"] = pd.to_datetime(quadro["DATA"]).dt.strftime("%d/%m/%Y")
            st.dataframe(quadro, use_container_width=True, hide_index=True)


def renderizar_conferencia_fiscal(prefixo: str, empresa: str) -> None:
    """Agrupa as conferências fiscal/contábil e de impostos na mesma aba."""
    import streamlit as st

    aba_fiscal, aba_impostos = st.tabs([
        "Fiscal × Contábil", "Impostos × Balancete",
    ])
    with aba_fiscal:
        _renderizar_conferencia_fiscal_contabil(prefixo, empresa)
    with aba_impostos:
        from razync.conferencia_impostos import renderizar_conferencia_impostos

        renderizar_conferencia_impostos(prefixo, empresa)


def _texto(valor) -> str:
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return ""
    return re.sub(r"\s+", " ", str(valor)).strip()


def _numero(valor) -> float:
    if isinstance(valor, (int, float)) and not isinstance(valor, bool) and not pd.isna(valor):
        return round(float(valor), 2)
    texto = _texto(valor).replace("R$", "").replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return round(float(texto), 2)
    except ValueError:
        return 0.0


def _codigo_dominio(valor) -> str:
    """Preserva códigos inteiros exportados como 22643.0 nos arquivos XLS."""
    texto = _texto(valor).strip()
    if not texto:
        return ""
    if re.fullmatch(r"\d+[.,]0+", texto):
        texto = re.split(r"[.,]", texto, maxsplit=1)[0]
    return texto if re.fullmatch(r"\d+", texto) else ""


def _data_dominio(valor):
    """Datas do Domínio podem ser datas Excel numéricas (ex.: 46237)."""
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        if pd.notna(valor) and 20000 <= float(valor) <= 100000:
            return pd.to_datetime(float(valor), unit="D", origin="1899-12-30", errors="coerce")
    if isinstance(valor, str):
        texto = valor.strip().replace(",", ".")
        if re.fullmatch(r"\d+(?:\.0+)?", texto):
            numero = float(texto)
            if 20000 <= numero <= 100000:
                return pd.to_datetime(numero, unit="D", origin="1899-12-30", errors="coerce")
    return pd.to_datetime(valor, dayfirst=True, errors="coerce")


def _rotulo(valor) -> str:
    texto = unicodedata.normalize("NFKD", _texto(valor))
    texto = "".join(letra for letra in texto if not unicodedata.combining(letra))
    return re.sub(r"[^A-Z0-9]+", " ", texto.upper()).strip()


def _primeiro_preenchido(valores: list, inicio: int | None, largura: int = 4):
    if inicio is None:
        return ""
    for indice in range(inicio, min(inicio + largura, len(valores))):
        if _texto(valores[indice]):
            return valores[indice]
    return ""


def _excel(conteudo: bytes, nome: str) -> pd.ExcelFile:
    try:
        return pd.ExcelFile(io.BytesIO(conteudo))
    except Exception as erro_original:
        if Path(nome).suffix.lower() != ".xls":
            raise ValueError(f"Não foi possível abrir {nome}: {erro_original}") from erro_original
        # Same BIFF recovery used by the working Domínio reader, preserving every sheet.
        from razync.dominio_ledger import _recuperar_xls_biff_irregular
        recuperado = _recuperar_xls_biff_irregular(conteudo, workbook=True)
        if recuperado is not None:
            return recuperado
        # Some exports named .xls are actually HTML tables.
        if b"<table" in conteudo[:65536].lower():
            try:
                try:
                    texto_html = conteudo.decode("utf-8-sig")
                except UnicodeDecodeError:
                    texto_html = conteudo.decode("cp1252")
                tabelas = pd.read_html(io.StringIO(texto_html), header=None, decimal=",", thousands=".")
                convertido = io.BytesIO()
                with pd.ExcelWriter(convertido, engine="openpyxl") as writer:
                    for indice, tabela in enumerate(tabelas):
                        tabela.to_excel(writer, sheet_name=f"Relatorio_{indice + 1}", index=False, header=False)
                return pd.ExcelFile(io.BytesIO(convertido.getvalue()))
            except Exception:
                pass
        conversor = shutil.which("soffice") or shutil.which("libreoffice")
        if not conversor:
            raise ValueError(
                f"Não foi possível ler {nome}, mesmo após tentar recuperar o XLS do Domínio. "
                "Envie o arquivo original para análise ou exporte novamente em XLSX."
            ) from erro_original
        pasta = tempfile.mkdtemp(prefix="razync_fiscal_")
        origem = Path(pasta) / "origem.xls"
        origem.write_bytes(conteudo)
        try:
            processo = subprocess.run(
                [conversor, "--headless", "--convert-to", "xlsx", "--outdir", pasta, str(origem)],
                capture_output=True, text=True, timeout=45, check=False,
                env={**os.environ, "HOME": pasta},
            )
            convertido = Path(pasta) / "origem.xlsx"
            if processo.returncode != 0 or not convertido.exists():
                raise ValueError("Não foi possível recuperar o XLS exportado pelo Domínio.")
            dados = convertido.read_bytes()
            return pd.ExcelFile(io.BytesIO(dados))
        finally:
            shutil.rmtree(pasta, ignore_errors=True)


def _empresa_no_cabecalho(valores: list, primeira_linha: bool = False) -> str:
    """Identifica empresa tanto no Razão ('Empresa:') quanto no fiscal (nome em A1)."""
    if primeira_linha and valores:
        primeira_celula = _texto(valores[0]).strip()
        # O Resumo por Acumulador do Domínio pode ter apenas a razão social
        # na célula A1, acima de CNPJ, período e título do relatório.
        nome_rotulado = _rotulo(primeira_celula)
        if (
            len(nome_rotulado.split()) >= 3
            and re.search(r"\b(LTDA|LIMITADA|EIRELI|SLU|EPP|S A|SA)\b", nome_rotulado)
            and not re.match(r"^EMPRESA\b", nome_rotulado)
        ):
            return primeira_celula
    for indice, valor in enumerate(valores[:6]):
        original = _texto(valor).strip()
        if not re.match(r"^empresa\b\s*:?", original, flags=re.IGNORECASE):
            continue
        resto = re.sub(r"^empresa\b\s*:?", "", original, flags=re.IGNORECASE).strip()
        candidatos = [resto] + [_texto(v).strip() for v in valores[indice + 1:indice + 10]]
        for texto in candidatos:
            texto = re.sub(r"^\d{1,8}\s*[-–:]\s*", "", texto).strip()
            texto = re.split(r"\s{2,}(?=Folha\b|P[aá]gina\b)", texto, maxsplit=1, flags=re.IGNORECASE)[0].strip()
            if len(texto) < 4 or not re.search(r"[A-Za-zÀ-ÿ]{2,}", texto):
                continue
            if _rotulo(texto) in {"CNPJ", "C N P J", "FOLHA", "RAZAO", "PERIODO", "CONSOLIDADO"}:
                continue
            return texto
    return ""


def _chave_empresa(nome: str) -> str:
    """Tolera pontuação e sufixos societários ao comparar a mesma razão social."""
    partes = _rotulo(nome).split()
    while partes and partes[-1] in {"LTDA", "LIMITADA", "EIRELI", "SLU", "ME", "EPP"}:
        partes.pop()
    if len(partes) >= 2 and partes[-2:] == ["S", "A"]:
        partes = partes[:-2]
    return " ".join(partes)


def _conferir_empresa_header(periodo: dict, valores: list, primeira_linha: bool = False) -> None:
    achada = _empresa_no_cabecalho(valores, primeira_linha=primeira_linha)
    if not achada:
        return
    anterior = periodo.get("empresa_nome", "")
    if anterior and _chave_empresa(anterior) != _chave_empresa(achada):
        raise ValueError("O relatório contém cabeçalhos de empresas diferentes. Envie apenas os dados de uma empresa por conferência.")
    periodo["empresa_nome"] = anterior or achada


def ler_acumuladores(conteudo: bytes, nome: str) -> tuple[pd.DataFrame, dict]:
    xls = _excel(conteudo, nome)
    periodo = {"inicio": None, "fim": None, "empresa_nome": ""}
    registros = []
    sem_conta = []
    for aba in xls.sheet_names:
        bruto = pd.read_excel(xls, sheet_name=aba, header=None, dtype=object)
        tipo = ""
        col_codigo, col_descricao = 0, None
        col_valor, col_conta = None, None
        for indice_linha, linha in bruto.iterrows():
            valores = linha.tolist()
            rotulos = [_rotulo(valor) for valor in valores]
            primeiro = rotulos[0] if rotulos else ""
            _conferir_empresa_header(periodo, valores, primeira_linha=indice_linha == 0)
            if primeiro == "PERIODO":
                datas = [pd.to_datetime(v, errors="coerce") for v in valores]
                datas = [d for d in datas if pd.notna(d)]
                if datas:
                    periodo.update({"inicio": min(datas), "fim": max(datas)})
            if primeiro in {"ENTRADAS", "SAIDAS", "SERVICOS"}:
                tipo = {"SAIDAS": "SAÍDAS", "SERVICOS": "SERVIÇOS"}.get(primeiro, primeiro)
                continue

            # O Domínio muda a posição das colunas conforme o relatório, a empresa
            # e a quantidade de tributos selecionados. Localizamos os cabeçalhos
            # pelo nome em vez de depender das antigas colunas fixas 13 e 49.
            cabecalho = any(r in {"COD", "CODIGO"} for r in rotulos)
            if cabecalho:
                for indice, rotulo in enumerate(rotulos):
                    if rotulo in {"COD", "CODIGO"}:
                        col_codigo = indice
                    elif rotulo == "DESCRICAO":
                        col_descricao = indice
                    elif rotulo in {"VLR CONTABIL", "VALOR CONTABIL"}:
                        col_valor = indice
                    elif rotulo in {"CONTA", "CONTA CONTABIL", "CODIGO CONTA"}:
                        col_conta = indice
                continue

            codigo = _codigo_dominio(_primeiro_preenchido(valores, col_codigo))
            if not codigo.isdigit():
                continue
            valor_indice = col_valor
            if valor_indice is None:
                valor_indice = 10 if tipo == "SERVIÇOS" else 13
            valor = _numero(_primeiro_preenchido(valores, valor_indice))
            if abs(valor) < 0.005:
                continue

            conta = _codigo_dominio(_primeiro_preenchido(valores, col_conta))
            if not re.fullmatch(r"\d+", conta):
                # Em algumas exportações o título "Conta" desaparece, embora o
                # código permaneça na última coluna preenchida do acumulador.
                candidatos = [
                    _codigo_dominio(valor_bruto) for indice, valor_bruto in enumerate(valores)
                    if indice >= valor_indice + 20
                    and _codigo_dominio(valor_bruto) not in {"", "0"}
                ]
                conta = candidatos[-1] if candidatos else ""
            if col_descricao is not None and len(valores) > col_descricao:
                descricao = _texto(_primeiro_preenchido(valores, col_descricao))
            else:
                descricao = next(
                    (_texto(v) for v in valores[col_codigo + 1:valor_indice] if _texto(v)), ""
                )
            if not re.fullmatch(r"\d+", conta):
                if tipo:
                    sem_conta.append({
                        "tipo": tipo, "acumulador": codigo,
                        "descricao": descricao, "valor": valor,
                    })
                continue
            registros.append({
                "TIPO": tipo, "ACUMULADOR": codigo, "DESCRIÇÃO": descricao,
                "CONTA": conta, "VALOR_FISCAL": valor,
            })
    if not registros:
        raise ValueError("Nenhum acumulador com conta contábil preenchida foi encontrado.")
    quadro = pd.DataFrame(registros)
    quadro.attrs["sem_conta"] = sem_conta
    return quadro, periodo


def ler_razao(conteudo: bytes, nome: str) -> tuple[pd.DataFrame, dict]:
    xls = _excel(conteudo, nome)
    registros = []
    periodo = {"inicio": None, "fim": None, "empresa_nome": ""}
    for aba in xls.sheet_names:
        conta = ""
        descricao_conta = ""
        bruto = pd.read_excel(xls, sheet_name=aba, header=None, dtype=object)
        colunas = {
            "DATA": 0, "LOTE": 1, "HISTÓRICO": 2, "CONTRAPARTIDA": 7,
            "DÉBITO": 8, "CRÉDITO": 9, "FILIAL": None,
        }
        for indice_linha, linha in bruto.iterrows():
            valores = linha.tolist()
            rotulos = [_rotulo(valor) for valor in valores]
            primeiro = rotulos[0] if rotulos else ""
            _conferir_empresa_header(periodo, valores, primeira_linha=indice_linha == 0)
            if primeiro == "PERIODO":
                achado = re.findall(r"\d{2}/\d{2}/\d{4}", " ".join(_texto(v) for v in valores))
                if len(achado) >= 2:
                    periodo.update({
                        "inicio": pd.to_datetime(achado[0], dayfirst=True),
                        "fim": pd.to_datetime(achado[1], dayfirst=True),
                    })
            if primeiro == "CONTA":
                conta = _codigo_dominio(_primeiro_preenchido(valores, 1))
                descricao_conta = _texto(_primeiro_preenchido(valores, 5, 5))
                continue

            # O razão consolidado pode acrescentar a coluna Filial e deslocar as
            # demais colunas. A posição é obtida pelos títulos de cada página.
            if "DATA" in rotulos and any(r == "HISTORICO" for r in rotulos):
                for indice, rotulo in enumerate(rotulos):
                    if rotulo == "DATA":
                        colunas["DATA"] = indice
                    elif rotulo == "LOTE":
                        colunas["LOTE"] = indice
                    elif rotulo == "HISTORICO":
                        colunas["HISTÓRICO"] = indice
                    elif rotulo in {"CTA C PART", "CONTRAPARTIDA", "CONTA CONTRAPARTIDA"}:
                        colunas["CONTRAPARTIDA"] = indice
                    elif rotulo == "DEBITO":
                        colunas["DÉBITO"] = indice
                    elif rotulo == "CREDITO":
                        colunas["CRÉDITO"] = indice
                    elif rotulo in {
                        "FILIAL", "COD FILIAL", "CODIGO FILIAL", "COD DA FILIAL",
                        "CODIGO DA FILIAL", "EMPRESA FILIAL", "ESTABELECIMENTO", "ESTAB",
                    } or rotulo.startswith("FILIAL "):
                        colunas["FILIAL"] = indice
                continue

            data = _data_dominio(
                valores[colunas["DATA"]] if len(valores) > colunas["DATA"] else None
            )
            if not conta or pd.isna(data):
                continue
            debito = _numero(valores[colunas["DÉBITO"]] if len(valores) > colunas["DÉBITO"] else 0)
            credito = _numero(valores[colunas["CRÉDITO"]] if len(valores) > colunas["CRÉDITO"] else 0)
            if abs(debito) < 0.005 and abs(credito) < 0.005:
                continue
            filial = ""
            if colunas["FILIAL"] is not None and len(valores) > colunas["FILIAL"]:
                # Nos XLS do Domínio o cabeçalho pode ocupar várias células e o
                # valor da filial ficar uma ou duas posições após o início dele.
                filial = _codigo_dominio(
                    _primeiro_preenchido(valores, colunas["FILIAL"], 3)
                )
                filial = filial.lstrip("0") or ("0" if filial else "")
            registros.append({
                "CONTA": conta, "DESCRIÇÃO_CONTA": descricao_conta, "FILIAL": filial,
                "DATA": data.normalize(),
                "LOTE": _texto(valores[colunas["LOTE"]] if len(valores) > colunas["LOTE"] else ""),
                "HISTÓRICO": _texto(valores[colunas["HISTÓRICO"]] if len(valores) > colunas["HISTÓRICO"] else ""),
                "CONTRAPARTIDA": _codigo_dominio(valores[colunas["CONTRAPARTIDA"]] if len(valores) > colunas["CONTRAPARTIDA"] else ""),
                "DÉBITO": debito, "CRÉDITO": credito,
            })
    if not registros:
        raise ValueError("Nenhum lançamento foi encontrado no razão.")
    return pd.DataFrame(registros), periodo


def _parece_fiscal(historico: str) -> bool:
    texto = _texto(historico).upper()
    if re.match(r"^(PAGO|PAGAMENTO|RECEBIDO|RECEBIMENTO|BAIXA)\b", texto):
        return False
    sinais = (" CF NF ", "COMPRA DE ", "VENDA DE ", "SERVIÇOS DE TERCEIROS", "SERVICOS DE TERCEIROS", "BENEFICIAMENTO")
    return any(sinal in f" {texto} " for sinal in sinais)


def conferir_fiscal_contabil(acumuladores: pd.DataFrame, razao: pd.DataFrame):
    resumos, detalhes = [], []
    agrupado = acumuladores.groupby(["CONTA", "TIPO"], as_index=False).agg(
        VALOR_FISCAL=("VALOR_FISCAL", "sum"),
        ACUMULADORES=("ACUMULADOR", lambda x: ", ".join(map(str, x))),
        DESCRIÇÕES=("DESCRIÇÃO", lambda x: " | ".join(map(str, x))),
    )
    for _, fiscal in agrupado.iterrows():
        conta = str(fiscal["CONTA"])
        lado = "CRÉDITO" if fiscal["TIPO"] == "SAÍDAS" else "DÉBITO"
        movimentos = razao[razao["CONTA"].astype(str).eq(conta)].copy()
        movimentos["DÉBITO"] = pd.to_numeric(
            movimentos.get("DÉBITO"), errors="coerce"
        ).fillna(0.0)
        movimentos["CRÉDITO"] = pd.to_numeric(
            movimentos.get("CRÉDITO"), errors="coerce"
        ).fillna(0.0)
        # Cada grupo do acumulador determina a coluna correta do razão:
        # ENTRADAS/SERVIÇOS = Débito; SAÍDAS = Crédito. O lado oposto não participa
        # da conferência. Valores negativos na coluna escolhida reduzem como estorno.
        movimentos["VALOR_ANALISADO"] = movimentos[lado]
        movimentos = movimentos[movimentos["VALOR_ANALISADO"].abs() >= 0.005].copy()
        movimentos["COMPATÍVEL_FISCAL"] = movimentos["HISTÓRICO"].map(_parece_fiscal)
        movimentos["SEM_EVIDÊNCIA_HISTÓRICO"] = False
        movimentos["NATUREZA"] = movimentos.apply(
            lambda mov: (
                f"ESTORNO/REDUTOR ({lado})"
                if float(mov["VALOR_ANALISADO"]) < 0
                else lado
            ),
            axis=1,
        )
        valor_fiscal = round(float(fiscal["VALOR_FISCAL"]), 2)
        valor_total = round(float(movimentos["VALOR_ANALISADO"].sum()), 2)
        # Fechamento aritmético não comprova origem fiscal. Mantemos o valor
        # na comparação, mas destacamos históricos sem evidência suficiente.
        if not movimentos.empty and abs(valor_total - valor_fiscal) <= 0.01:
            movimentos["SEM_EVIDÊNCIA_HISTÓRICO"] = ~movimentos["COMPATÍVEL_FISCAL"]
            movimentos["COMPATÍVEL_FISCAL"] = True
        valor_compativel = round(float(movimentos.loc[movimentos["COMPATÍVEL_FISCAL"], "VALOR_ANALISADO"].sum()), 2)
        total_debitos = round(float(movimentos["DÉBITO"].sum()), 2) if lado == "DÉBITO" else 0.0
        total_creditos = round(float(movimentos["CRÉDITO"].sum()), 2) if lado == "CRÉDITO" else 0.0
        total_redutores = round(float(-movimentos.loc[movimentos["VALOR_ANALISADO"] < 0, "VALOR_ANALISADO"].sum()), 2)
        diferenca = round(valor_compativel - valor_fiscal, 2)
        extras = movimentos[~movimentos["COMPATÍVEL_FISCAL"]].copy()
        nao_comprovados = int(movimentos["SEM_EVIDÊNCIA_HISTÓRICO"].sum())
        if abs(diferenca) <= 0.01:
            situacao = "CONFERE COM ALERTAS" if not extras.empty or nao_comprovados else "CONFERE"
        elif movimentos.empty:
            situacao = "AUSENTE NO CONTÁBIL"
        else:
            situacao = "REVISAR"
        resumos.append({
            "CONTA": conta, "TIPO": fiscal["TIPO"], "ACUMULADORES": fiscal["ACUMULADORES"],
            "VALOR FISCAL": valor_fiscal, "CONTÁBIL COMPATÍVEL": valor_compativel,
            "TOTAL DA CONTA": valor_total, "DIFERENÇA FISCAL": diferenca,
            "TOTAL DÉBITOS": total_debitos, "TOTAL CRÉDITOS": total_creditos,
            "ESTORNOS/REDUTORES": total_redutores,
            "LANÇAMENTOS EXTRAS": len(extras),
            "FECHAMENTOS SEM EVIDÊNCIA": nao_comprovados,
            "SITUAÇÃO": situacao,
        })
        for _, mov in movimentos.iterrows():
            detalhes.append({
                "CONTA": conta, "DATA": mov["DATA"], "LOTE": mov["LOTE"],
                "HISTÓRICO": mov["HISTÓRICO"], "CONTRAPARTIDA": mov["CONTRAPARTIDA"],
                "DÉBITO": mov["DÉBITO"], "CRÉDITO": mov["CRÉDITO"],
                "NATUREZA": mov["NATUREZA"], "VALOR": mov["VALOR_ANALISADO"],
                "CLASSIFICAÇÃO": ("FECHAMENTO POR VALOR - VALIDAR"
                                  if mov["SEM_EVIDÊNCIA_HISTÓRICO"] else
                                  "FISCAL PROVÁVEL" if mov["COMPATÍVEL_FISCAL"]
                                  else "ALERTA - NÃO FISCAL"),
            })
    return pd.DataFrame(resumos), pd.DataFrame(detalhes)


def processar_conferencia(
    acumuladores_bytes: bytes, acumuladores_nome: str,
    razao_bytes: bytes, razao_nome: str, filial_alvo: str | None = None,
):
    acumuladores, periodo_fiscal = ler_acumuladores(acumuladores_bytes, acumuladores_nome)
    razao, periodo_razao = ler_razao(razao_bytes, razao_nome)
    empresa_fiscal = periodo_fiscal.get("empresa_nome", "")
    empresa_razao = periodo_razao.get("empresa_nome", "")
    if empresa_fiscal and empresa_razao and _chave_empresa(empresa_fiscal) != _chave_empresa(empresa_razao):
        raise ValueError(
            "Os documentos pertencem a empresas diferentes: Resumo por Acumulador: "
            f"{empresa_fiscal}; Razão: {empresa_razao}. Envie os dois relatórios da mesma empresa."
        )
    empresa_nome = empresa_fiscal or empresa_razao
    filiais_encontradas = sorted(
        filial for filial in razao.get("FILIAL", pd.Series(dtype=str)).astype(str).unique()
        if filial
    )
    filial_normalizada = re.sub(r"\D", "", str(filial_alvo or "")).lstrip("0")
    if len(filiais_encontradas) > 1 and not filial_normalizada:
        raise ValueError(
            "O Razão contém várias filiais (" + ", ".join(filiais_encontradas) +
            "). Informe o código da filial mostrado no Razão para evitar somar estabelecimentos diferentes."
        )
    filial_aplicada = filiais_encontradas[0] if len(filiais_encontradas) == 1 and not filial_normalizada else ""
    if filiais_encontradas and filial_normalizada:
        if filial_normalizada in filiais_encontradas:
            filial_aplicada = filial_normalizada
            razao = razao[razao["FILIAL"].astype(str).eq(filial_aplicada)].copy()
        elif len(filiais_encontradas) == 1:
            # Em relatórios individuais do Domínio, a coluna Filial costuma
            # trazer o estabelecimento interno (normalmente 1), e não o código
            # da empresa no escritório. Como não existe outra filial no arquivo,
            # todo o Razão pertence à empresa selecionada.
            filial_aplicada = filiais_encontradas[0]
            razao = razao[razao["FILIAL"].astype(str).eq(filial_aplicada)].copy()
        else:
            raise ValueError(
                f"O razão contém as filiais {', '.join(filiais_encontradas)}, mas não possui "
                f"lançamentos da empresa {filial_normalizada}."
            )
    sem_conta = acumuladores.attrs.get("sem_conta", [])
    resumo, detalhes = conferir_fiscal_contabil(acumuladores, razao)
    return {
        "sem_conta": sem_conta,
        "resumo": resumo, "detalhes": detalhes, "acumuladores": acumuladores,
        "periodo_fiscal": periodo_fiscal, "periodo_razao": periodo_razao,
        "empresa_nome": empresa_nome,
        "empresa_fiscal": empresa_fiscal, "empresa_razao": empresa_razao,
        "filial_aplicada": filial_aplicada,
        "filiais_encontradas": filiais_encontradas,
    }


def gerar_relatorio_excel(resultado: dict) -> bytes:
    """Gera o relatório da conferência em abas auditáveis."""
    wb = Workbook()
    wb.remove(wb.active)
    azul = PatternFill("solid", fgColor="123047")
    azul_claro = PatternFill("solid", fgColor="DCEAF4")
    amarelo = PatternFill("solid", fgColor="FFF1CC")
    vermelho = PatternFill("solid", fgColor="FADBD8")
    verde = PatternFill("solid", fgColor="DDF2E1")

    def adicionar_aba(nome: str, quadro: pd.DataFrame):
        ws = wb.create_sheet(nome)
        if quadro is None or quadro.empty:
            ws.append(["Nenhum registro encontrado"])
            return ws
        colunas = list(quadro.columns)
        ws.append(colunas)
        for celula in ws[1]:
            celula.fill = azul
            celula.font = Font(color="FFFFFF", bold=True)
            celula.alignment = Alignment(horizontal="center")
        for registro in quadro.itertuples(index=False, name=None):
            valores = []
            for valor in registro:
                if isinstance(valor, pd.Timestamp):
                    valor = valor.to_pydatetime()
                elif pd.isna(valor):
                    valor = ""
                valores.append(valor)
            ws.append(valores)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for indice, coluna in enumerate(colunas, start=1):
            valores = [str(ws.cell(linha, indice).value or "") for linha in range(1, ws.max_row + 1)]
            ws.column_dimensions[get_column_letter(indice)].width = min(max(len(v) for v in valores) + 2, 55)
            if any(chave in str(coluna).upper() for chave in ("VALOR", "TOTAL", "DIFERENÇA", "CONTÁBIL")):
                for linha in range(2, ws.max_row + 1):
                    ws.cell(linha, indice).number_format = 'R$ #,##0.00'
            if str(coluna).upper() == "DATA":
                for linha in range(2, ws.max_row + 1):
                    ws.cell(linha, indice).number_format = "dd/mm/yyyy"
        return ws

    resumo = resultado.get("resumo", pd.DataFrame()).copy()
    detalhes = resultado.get("detalhes", pd.DataFrame()).copy()
    acumuladores = resultado.get("acumuladores", pd.DataFrame()).copy()
    ws_resumo = adicionar_aba("Resumo por conta", resumo)
    if not resumo.empty and "SITUAÇÃO" in resumo.columns:
        coluna_situacao = list(resumo.columns).index("SITUAÇÃO") + 1
        for linha in range(2, ws_resumo.max_row + 1):
            situacao = str(ws_resumo.cell(linha, coluna_situacao).value or "")
            preenchimento = (
                verde if situacao == "CONFERE" else
                amarelo if situacao == "CONFERE COM ALERTAS" else vermelho
            )
            for celula in ws_resumo[linha]:
                celula.fill = preenchimento
    alertas = detalhes[
        detalhes.get("CLASSIFICAÇÃO", pd.Series(dtype=str)).eq("ALERTA - NÃO FISCAL")
    ].copy() if not detalhes.empty else pd.DataFrame()
    ws_alertas = adicionar_aba("Alertas", alertas)
    if ws_alertas.max_row > 1:
        for linha in range(2, ws_alertas.max_row + 1):
            for celula in ws_alertas[linha]:
                celula.fill = amarelo
    adicionar_aba("Todos os lançamentos", detalhes)
    adicionar_aba("Acumuladores considerados", acumuladores)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()
