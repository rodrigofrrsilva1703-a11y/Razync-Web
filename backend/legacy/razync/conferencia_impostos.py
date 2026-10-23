"""Conferência dos impostos da Receita/DCTFWeb com o balancete contábil.

Versão sincronizada com o Conector Windows 0.4.
"""

from __future__ import annotations

import io
import re
import unicodedata
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pypdf import PdfReader


IMPOSTOS = {
    "INSS / Previdenciários": ("INSS", "PREVIDENCIARIA", "PREVIDENCIARIO", "CONTRIBUICAO PREVIDENCIARIA"),
    "IRRF": ("IRRF", "IMPOSTO DE RENDA RETIDO"),
    "PIS/PASEP": ("PIS", "PASEP"),
    "COFINS": ("COFINS",),
    "CSLL": ("CSLL", "CONTRIBUICAO SOCIAL SOBRE O LUCRO"),
    "IRPJ": ("IRPJ", "IMPOSTO DE RENDA PESSOA JURIDICA"),
    "CPRB": ("CPRB", "RECEITA BRUTA PREVIDENCIARIA"),
}


def _rotulo(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(letra for letra in texto if not unicodedata.combining(letra))
    return re.sub(r"\s+", " ", texto.upper()).strip()


def _numero(valor) -> float | None:
    if isinstance(valor, (int, float)) and not pd.isna(valor):
        return round(float(valor), 2)
    texto = str(valor or "").strip().replace("R$", "").replace(" ", "")
    if not texto:
        return None
    negativo = texto.startswith("(") and texto.endswith(")")
    texto = texto.strip("()")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        numero = float(texto)
    except ValueError:
        return None
    return round(-numero if negativo else numero, 2)


def _linhas_documento(conteudo: bytes, nome: str) -> list[list]:
    extensao = Path(nome).suffix.lower()
    if extensao == ".pdf":
        leitor = PdfReader(io.BytesIO(conteudo))
        linhas = [
            [linha]
            for pagina in leitor.pages
            for linha in (pagina.extract_text() or "").splitlines()
            if linha.strip()
        ]
        if linhas:
            return linhas
        # Alguns balancetes são PDFs digitalizados. Nesse caso usamos OCR,
        # mantendo o mesmo fluxo de extração por descrições e valores.
        try:
            import fitz
            import pytesseract
            from PIL import Image

            documento = fitz.open(stream=conteudo, filetype="pdf")
            linhas_ocr = []
            for pagina in documento:
                pixmap = pagina.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                imagem = Image.frombytes(
                    "RGB", [pixmap.width, pixmap.height], pixmap.samples
                )
                texto = pytesseract.image_to_string(imagem, lang="por")
                linhas_ocr.extend([[linha] for linha in texto.splitlines() if linha.strip()])
            return linhas_ocr
        except Exception as erro:
            raise ValueError(
                "O balancete PDF não possui texto legível e o OCR não conseguiu processá-lo."
            ) from erro
    if extensao == ".csv":
        quadro = pd.read_csv(io.BytesIO(conteudo), header=None, sep=None, engine="python", dtype=object)
        return quadro.fillna("").values.tolist()
    xls = pd.ExcelFile(io.BytesIO(conteudo))
    linhas: list[list] = []
    for aba in xls.sheet_names:
        quadro = pd.read_excel(xls, sheet_name=aba, header=None, dtype=object)
        linhas.extend(quadro.fillna("").values.tolist())
    return linhas


def _valor_da_linha(linha: list) -> float | None:
    numeros = [_numero(valor) for valor in linha]
    numeros = [valor for valor in numeros if valor is not None]
    if numeros:
        return abs(numeros[-1])
    texto = " ".join(str(valor) for valor in linha)
    encontrados = re.findall(r"(?:R\$\s*)?-?\d{1,3}(?:\.\d{3})*,\d{2}", texto)
    return abs(_numero(encontrados[-1]) or 0) if encontrados else None


def extrair_impostos(conteudo: bytes, nome: str) -> dict[str, float]:
    totais = {imposto: 0.0 for imposto in IMPOSTOS}
    encontrados = {imposto: False for imposto in IMPOSTOS}
    for linha in _linhas_documento(conteudo, nome):
        texto = _rotulo(" ".join(str(valor) for valor in linha))
        valor = _valor_da_linha(linha)
        if valor is None:
            continue
        for imposto, aliases in IMPOSTOS.items():
            if any(alias in texto for alias in aliases):
                totais[imposto] += valor
                encontrados[imposto] = True
                break
    return {imposto: round(valor, 2) for imposto, valor in totais.items() if encontrados[imposto]}


def processar_conferencia_impostos(
    receita: bytes, nome_receita: str, balancete: bytes, nome_balancete: str,
) -> pd.DataFrame:
    valores_receita = extrair_impostos(receita, nome_receita)
    valores_contabeis = extrair_impostos(balancete, nome_balancete)
    impostos = sorted(set(valores_receita) | set(valores_contabeis))
    if not impostos:
        raise ValueError(
            "Não encontrei linhas de impostos nos arquivos. Envie o relatório detalhado "
            "da DCTFWeb/Receita e um balancete com as descrições das contas."
        )
    registros = []
    for imposto in impostos:
        receita_valor = valores_receita.get(imposto, 0.0)
        contabil_valor = valores_contabeis.get(imposto, 0.0)
        diferenca = round(contabil_valor - receita_valor, 2)
        registros.append({
            "IMPOSTO": imposto,
            "RECEITA / DCTFWEB": receita_valor,
            "BALANCETE": contabil_valor,
            "DIFERENÇA": diferenca,
            "SITUAÇÃO": "CONFERE" if abs(diferenca) <= 0.01 else "REVISAR",
        })
    return pd.DataFrame(registros)


def gerar_relatorio_impostos(resultado: pd.DataFrame, empresa: str, competencia: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Conferência de Impostos"
    ws.append(["EMPRESA", empresa])
    ws.append(["COMPETÊNCIA", competencia.strftime("%m/%Y")])
    ws.append([])
    ws.append(resultado.columns.tolist())
    for celula in ws[4]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="147DA4")
    for linha in resultado.itertuples(index=False, name=None):
        ws.append(list(linha))
    for coluna in ("B", "C", "D"):
        for celula in ws[coluna][4:]:
            celula.number_format = 'R$ #,##0.00'
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 14
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _cnpj_valido(valor: str) -> bool:
    cnpj = re.sub(r"\D", "", str(valor or ""))
    if len(cnpj) != 14 or len(set(cnpj)) == 1:
        return False
    numeros = [int(digito) for digito in cnpj]
    for tamanho, pesos in (
        (12, [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]),
        (13, [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]),
    ):
        soma = sum(numero * peso for numero, peso in zip(numeros[:tamanho], pesos))
        digito = 11 - soma % 11
        esperado = 0 if digito >= 10 else digito
        if numeros[tamanho] != esperado:
            return False
    return True


def _formatar_cnpj(valor: str) -> str:
    cnpj = re.sub(r"\D", "", str(valor or ""))
    if len(cnpj) != 14:
        return str(valor or "")
    return f"{cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]}"


def extrair_cnpjs(conteudo: bytes, nome: str) -> list[str]:
    """Localiza CNPJs válidos no balancete."""
    encontrados: list[str] = []
    padrao = re.compile(r"(?<!\d)\d{2}[\.\s]?\d{3}[\.\s]?\d{3}[\/\s]?\d{4}[-\s]?\d{2}(?!\d)")
    for linha in _linhas_documento(conteudo, nome):
        texto = " ".join(str(valor) for valor in linha)
        for candidato in padrao.findall(texto):
            cnpj = re.sub(r"\D", "", candidato)
            if _cnpj_valido(cnpj) and cnpj not in encontrados:
                encontrados.append(cnpj)
    return encontrados
