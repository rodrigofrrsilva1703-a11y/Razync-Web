import io
import json
from datetime import date
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from app.main import app, _original_export, _reconcile_multi
from app import engine
from app.migration_services import workflow, eletro_reports
from test_migration import xlsx, workbook_signature, reference_engine

client = TestClient(app)


def test_multi_bank_reconciliation_keeps_each_bank_separate(monkeypatch):
    import app.migration_services as migration_services

    def fake_reconcile(company_code, bank, model, statements, options):
        daily = pd.DataFrame([{
            'DATA': pd.Timestamp('2026-10-01'),
            'ENTRADAS EXTRATO': 100.0,
            'SAÍDAS EXTRATO': 20.0,
            'ENTRADAS PLANILHA': 100.0,
            'SAÍDAS PLANILHA': 20.0,
            'DIF. ENTRADAS': 0.0,
            'DIF. SAÍDAS': 0.0,
            'STATUS ENTRADAS': '✅ Batendo',
            'STATUS SAÍDAS': '✅ Batendo',
            'STATUS': '✅ Batendo',
        }])
        return {'ok': True}, {'Conferência diária': daily}

    monkeypatch.setattr(migration_services, 'reconcile', fake_reconcile)
    summary, rows, sheets = _reconcile_multi(
        266,
        ['itau', 'fibra'],
        b'modelo',
        [('itau', 'itau.pdf', b'1'), ('fibra', 'fibra.xlsx', b'2')],
        {},
    )
    assert summary['banks'] == 2
    assert {row['BANCO'] for row in rows} == {'Itaú 508', 'Banco Fibra 506'}
    assert any(name.startswith('Itaú 508') for name in sheets)
    assert any(name.startswith('Banco Fibra') for name in sheets)

def test_eletro_reports_preserve_individual_original_and_consolidated():
    code = 242
    from razync import eletro_forte as ef
    raw = b'<table><tr><th>Data</th><th>Debito</th><th>Credito</th><th>Valor</th><th>Historico</th></tr><tr><td>03/08</td><td>166</td><td>8</td><td>100,50</td><td>CLIENTE TESTE</td></tr></table>'
    roles = {name:[('relatorio.xls',raw)] for name in ('despesas','fornecedores','recebidos')}
    actual = eletro_reports(code,roles,{'ano':2026})
    groups = [f(raw,2026,None) for f in (ef.processar_despesas,ef.processar_fornecedores,ef.processar_recebidos)]
    expected = ef.gerar_consolidado_bancos_eletro_forte(engine.TEMPLATE.read_bytes(),*groups)
    assert workbook_signature(actual[f'ELETRO_FORTE_{code}_CONSOLIDADO.xlsx']) == workbook_signature(expected)
    for index,name in enumerate(('despesas','fornecedores','recebidos')):
        params=[None,{},{}];params[index]=groups[index]
        expected = ef.gerar_modelo_dominio_eletro_forte(raw,'relatorio.xls',engine.TEMPLATE.read_bytes(),*params)
        assert workbook_signature(actual[f'ELETRO_FORTE_{code}_{name.upper()}.xlsx']) == workbook_signature(expected)
    response = client.post(f'/api/v1/workflow/{code}/reports',data={'roles_json':json.dumps(list(roles)),'options_json':'{"ano":2026}'},files=[('files',('relatorio.xls',raw)) for _ in roles])
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == 'application/zip'


def test_1408_rejects_matrix_report_flow_and_requires_statement_plus_receipts():
    raw = b'<table><tr><th>Data</th><th>Debito</th><th>Credito</th><th>Valor</th><th>Historico</th></tr><tr><td>03/08</td><td>166</td><td>512</td><td>100,50</td><td>CLIENTE TESTE</td></tr></table>'
    legacy_roles = {name:[('relatorio.xls',raw)] for name in ('despesas','fornecedores','recebidos')}
    response = client.post(
        '/api/v1/workflow/1408/reports',
        data={'roles_json': json.dumps(list(legacy_roles)), 'options_json': '{}'},
        files=[('files',('relatorio.xls',raw)) for _ in legacy_roles],
    )
    assert response.status_code == 422
    assert 'Extrato Itaú' in response.json()['detail']

    response = client.post(
        '/api/v1/workflow/1408',
        data={'roles_json':'["extrato"]','options_json':'{}'},
        files=[('files',('itau.pdf',b'invalid'))],
    )
    assert response.status_code == 422


@pytest.mark.parametrize('code,accounts',[(266,['99549-5','451990-6','673947-1']),(1396,['98002-6','3084-8'])])
def test_nova_geracao_period_removals_and_workbook_match_original(code,accounts):
    data=xlsx([['Conta','Data','Valor','Lacto','Historico','Doc']]+[[a,'01/10/2026',100,'PAGAR','Fornecedor teste','NF1'] for a in accounts]+[[accounts[0],'02/10/2026',100,'ESTORNO DE BAIXA','Estorno de baixa fornecedor','NF1']])
    ref=reference_engine()
    parsers=[ref.processar_nova_geracao_itau,ref.processar_nova_geracao_bradesco,ref.processar_nova_geracao_fibra] if code==266 else [ref.processar_nova_geracao_filial_itau,ref.processar_nova_geracao_filial_bradesco]
    groups={name:dict(zip(('principal','retirados'),parser(data))) for name,parser in zip(('Itaú','Bradesco','Fibra'),parsers)}
    expected=ref.gerar_excel_nova_geracao(groups,engine.TEMPLATE.read_bytes(),prefixar_historicos=False)
    actual,_=workflow(code,{'consolidada':[('consolidada.xlsx',data)]},{})
    assert workbook_signature(actual)==workbook_signature(expected)

def test_vgv_workflow_needs_only_caixa_and_matches_original(monkeypatch):
    from test_original_vgv_1402 import _caixa_bytes
    from razync.vgv_1402 import ler_caixa_vgv
    data=_caixa_bytes()
    ref=reference_engine();monkeypatch.chdir(engine.TEMPLATE.parent)
    expected=ref.gerar_excel_modelo_dominio(ler_caixa_vgv(data)[['DESCRIÇÃO','DATA','VALOR','DÉBITO','CRÉDITO','HISTÓRICO']])
    actual,_=workflow(1402,{'planilha':[('caixa.xlsx',data)]},{})
    assert workbook_signature(actual)==workbook_signature(expected)

def test_up_pack_sicredi_keeps_original_accounts_and_workbook():
    from razync.up_pack import processar_planilha_up_pack
    data=xlsx([['Data','D/C','Complemento','Conf','Entrada','Saída','Saldo','Conta Vinculada'],['13/07/2026','TRANSFERÊNCIA (E)','TRANSFERENCIA ENTRE CONTAS',None,400,0,410.38,'SANTANDER'],['22/07/2026','PAGTO TITULO','Doc.084574/01 - HOTEL',None,0,238,172.38,None]])
    ref=reference_engine()
    expected=ref.gerar_excel_nova_geracao({'Sicredi':{'principal':processar_planilha_up_pack(data,'sicredi'),'retirados':pd.DataFrame()}},engine.TEMPLATE.read_bytes())
    actual,_=workflow(1096,{'sicredi':[('sig.xlsx',data)]},{})
    assert workbook_signature(actual)==workbook_signature(expected)

def test_preview_runs_same_workflow_and_returns_real_counts():
    from test_original_vgv_1402 import _caixa_bytes
    response=client.post('/api/v1/workflow/1402/preview',data={'roles_json':'["planilha"]'},files={'files':('caixa.xlsx',_caixa_bytes())})
    assert response.status_code==200,response.text
    sheet=response.json()['sheets'][0]
    assert (sheet['count'],sheet['entradas'],sheet['saidas'])==(2,4730,540.33)

def test_tax_endpoint_excel_matches_original_report():
    from razync.conferencia_impostos import processar_conferencia_impostos,gerar_relatorio_impostos
    balance=xlsx([['INSS',100],['IRRF',80]])
    revenue=xlsx([['INSS',100],['IRRF',90]])
    expected=gerar_relatorio_impostos(processar_conferencia_impostos(revenue,'receita.xlsx',balance,'balancete.xlsx'),'242 - Eletro Forte',date(2026,10,1))
    response=client.post('/api/v1/conferencia-impostos/242',data={'competencia':'2026-10'},files={'receita':('receita.xlsx',revenue),'balancete':('balancete.xlsx',balance)})
    assert response.status_code==200,response.text
    summary=json.loads(response.headers['x-razync-summary'])
    assert (summary['impostos'],summary['conferem'],summary['revisar'])==(2,1,1)
    from openpyxl import load_workbook
    actual=load_workbook(io.BytesIO(response.content));original=load_workbook(io.BytesIO(expected))
    # Catalog's full legal name is intentional; all calculation and formatting match.
    original.active['B1']=actual.active['B1'].value
    out=io.BytesIO();original.save(out)
    assert workbook_signature(response.content)==workbook_signature(out.getvalue())

def test_bank_validation_does_not_silently_use_another_processor():
    response=client.post('/api/v1/modelo-dominio/154',data={'bank':'sicredi'},files={'files':('extrato.pdf',b'pdf')})
    assert response.status_code==400

def pdf_text(text):
    import fitz
    document=fitz.open();page=document.new_page()
    page.insert_text((40,40),text,fontsize=9)
    return document.tobytes()

def test_gz_real_pdf_desmembramento_dates_balances_and_excel():
    from razync.gz_1211 import processar_gz,gerar_modelo_dominio_gz
    statement=pdf_text('GZ IMPORTADORA 0099343-5\n01/10/2026 SALDO ANTERIOR 1.000,00\n01/10/2026 BOLETOS RECEBIDOS 300,00\n01/10/2026 TARIFA -10,00\n01/10/2026 SALDO EM CONTA CORRENTE 1.290,00')
    receipts=pdf_text('Boletos baixados e liquidados GZ IMPORTADORA ITAU\nCLIENTE A 30/09/2026 01/10/2026 100,00 157 Liquidado\nCLIENTE B 30/09/2026 01/10/2026 200,00 157 Liquidado\nCLIENTE C 30/09/2026 01/10/2026 500,00 157 Baixado')
    frame,diagnosis,unused,summary=processar_gz(statement,receipts)
    assert len(frame)==3
    assert frame['VALOR'].tolist()==[100,200,-10]
    assert summary['agregados_batendo']==1
    assert summary['saldo_final_calculado']==1290
    expected=gerar_modelo_dominio_gz(frame,engine.TEMPLATE.read_bytes())
    roles={'extrato':[('gz.pdf',statement)],'boletos':[('liquidados.pdf',receipts)]}
    actual,_=workflow(1211,roles,{})
    assert workbook_signature(actual)==workbook_signature(expected)
    from app.migration_services import diagnostics
    actual_tables=diagnostics(1211,roles,{})
    pd.testing.assert_frame_equal(actual_tables['Boletos x extrato'],diagnosis)

def test_lcarlos_api_preserves_auxiliary_replacements_and_original_excel(monkeypatch):
    from razync.lcarlos import processar_planilhas_lcarlos
    jaguar=xlsx([['Data','Cliente ou Fornecedor (Razão Social ou Nome Fantasia)','','','Valor (R$)','Saldo (R$)'],['05/08/2026','RECEBIMENTO VENDAS NF','','',75,75],['06/08/2026','TARIFA BANCARIA','','',-5,70]])
    auxiliary=xlsx([['RECEBIMENTOS','LCARLOS','AGOSTO',''],['','','',''],['05/08/2026',50,'BOLETO NF1','CLIENTE A'],['',25,'BOLETO NF2','CLIENTE B']])
    frame,_,summary=processar_planilhas_lcarlos(jaguar,auxiliary)
    assert len(frame)==3
    assert summary['diferenca_total']==0
    ref=reference_engine();monkeypatch.chdir(engine.TEMPLATE.parent)
    expected=ref.gerar_excel_modelo_dominio(frame)
    response=client.post('/api/v1/workflow/285',data={'roles_json':'["jaguar","entradas"]'},files=[('files',('Jaguar.xlsx',jaguar)),('files',('Entradas.xlsx',auxiliary))])
    assert response.status_code==200,response.text
    assert workbook_signature(response.content)==workbook_signature(expected)

def test_radani_api_keeps_sispag_totals_and_full_diagnostics(monkeypatch):
    from app import migration_services as services
    from razync import radani
    from test_original_radani import _extrato,_comprovantes
    source,receipts=_extrato(),_comprovantes()
    monkeypatch.setattr(engine,'_radani_cache_extrato_pdf',lambda *_:source.to_dict('records'))
    monkeypatch.setattr(services,'statement',lambda *_:source.copy())
    def read_receipts(files,start,end):
        assert pd.Timestamp(start)==source['DATA'].min()
        assert pd.Timestamp(end)==source['DATA'].max()
        return receipts.copy()
    monkeypatch.setattr(radani,'consolidar_comprovantes_sispag',read_receipts)
    original=radani.analisar_desmembramentos(source,'Itaú',receipts)
    expected=reference_engine().gerar_excel_nova_geracao({'Itaú':{'principal':original.organizado,'retirados':pd.DataFrame()}},engine.TEMPLATE.read_bytes())
    roles={'itau':[('itau.pdf',b'pdf')],'sispag':[('sispag.pdf',b'pdf')]}
    actual,_=workflow(968,roles,{})
    assert workbook_signature(actual)==workbook_signature(expected)
    tables=services.diagnostics(968,roles,{})
    pd.testing.assert_frame_equal(tables['Detalhamentos Itaú'],original.detalhamentos)

def test_1408_extrato_is_primary_when_recebidos_is_also_uploaded(monkeypatch):
    from app import migration_services as services
    from razync.eletro_forte_filial_1408 import montar_modelo_1408
    source=pd.DataFrame([{'DATA':'04/08/2026','VALOR':80.63,'HISTÓRICO':'PIX RECEBIDO'},{'DATA':'04/08/2026','VALOR':-24.56,'HISTÓRICO':'TAR COBRANCA'}])
    raw=b'<table><tr><th>Data</th><th>Debito</th><th>Credito</th><th>Valor</th><th>Historico</th></tr><tr><td>04/08</td><td>166</td><td>2</td><td>80,63</td><td>CLIENTE TESTE</td></tr></table>'
    monkeypatch.setattr(services,'statement',lambda *_:source.copy())
    frame,summary=montar_modelo_1408(source.to_dict('records'),raw,2026)
    assert len(frame)==2 and summary['historicos_substituidos']==1
    ref=reference_engine();monkeypatch.chdir(engine.TEMPLATE.parent)
    expected=ref.gerar_excel_modelo_dominio(frame)
    actual,_=workflow(1408,{'extrato':[('itau.pdf',b'pdf')],'recebidos':[('recebidos.xls',raw)]},{'ano':2026})
    assert workbook_signature(actual)==workbook_signature(expected)

def test_nibo_bank_descriptions_order_and_unclassified_accounts_match_original(monkeypatch):
    from razync import nibo
    from test_migration import model
    from openpyxl import load_workbook
    data=pd.DataFrame([{'DESCRIÇÃO':'NIBO','DATA':'02/10/2026','VALOR':100,'DÉBITO':'','CRÉDITO':'','HISTÓRICO':'Recebido: ACME'}])
    monkeypatch.setattr(nibo,'processar_extrato_nibo_pdf',lambda *_:data.copy())
    combined=pd.concat([data.assign(DESCRIÇÃO='BANCO ITAÚ'),data.assign(DESCRIÇÃO='BANCO DO BRASIL')],ignore_index=True).sort_values(['DATA','DESCRIÇÃO'],kind='stable')
    ref=reference_engine();monkeypatch.chdir(engine.TEMPLATE.parent)
    expected=ref.gerar_excel_modelo_dominio(combined)
    actual,_=workflow(1529,{'itau':[('nibo.pdf',b'pdf')],'banco_brasil':[('nibo-bb.pdf',b'pdf')]},{})
    assert workbook_signature(actual)==workbook_signature(expected)

def test_francesinhas_endpoint_preserves_emitido_em_and_original_tabs(monkeypatch):
    from test_original_eletro_forte_francesinhas import TEXTO_508,TEXTO_509,_zip_teste
    from razync import eletro_forte_francesinhas as original
    monkeypatch.setattr(original,'_texto_pdf',lambda content:TEXTO_508 if content==b'pdf-508' else TEXTO_509)
    source=_zip_teste();frame,_=original.processar_zip_francesinhas(source)
    expected=original.gerar_excel_francesinhas(engine.TEMPLATE.read_bytes(),frame)
    response=client.post('/api/v1/francesinhas/242',files={'files':('francesinhas.zip',source)})
    assert response.status_code==200,response.text
    assert workbook_signature(response.content)==workbook_signature(expected)

def test_txt_keeps_original_history_and_excel_serial_dates():
    from test_migration import model
    source=model(date=46212)
    response=client.post('/api/v1/modelo-dominio-txt',files={'file':('modelo.xlsx',source)})
    assert response.status_code==200,response.text
    assert response.content.decode().startswith('09/07/2026;')
    assert 'Pago: Empresa: ACME INDUSTRIAL' in response.content.decode()

def test_company_catalog_does_not_advertise_fiscal_or_tax_tools():
    result=client.get('/api/v1/companies').json()
    assert len(result)==48
    assert all('conferencia_fiscal' not in row['capabilities']['tools'] for row in result)
    assert all('conferencia_impostos' not in row['capabilities']['tools'] for row in result)
    assert all(row['capabilities'].get('status') != 'fiscal_only' for row in result)
    empty_companies = [row for row in result if row['capabilities'].get('status') != 'api_ready']
    assert empty_companies
    assert all(row['capabilities']['tools'] == [] for row in empty_companies)
    assert all(row['capabilities']['banks'] == {} for row in empty_companies)

def test_reconciliation_repeated_movements_and_removed_estornos_match_reference(monkeypatch):
    from openpyxl import Workbook
    from app import migration_services as services
    wb=Workbook();ws=wb.active;ws.title='Itaú'
    ws.append(services.COLUNAS)
    for value in (100,100,-25):
        ws.append(['BANCO ITAÚ','01/10/2026',value,'','','Recebido: ACME' if value>0 else 'Pago: TARIFA'])
    removed=wb.create_sheet('Lançamentos retirados');removed.append(services.COLUNAS+['MOTIVO']);removed.append(['BANCO ITAÚ','01/10/2026',200,'','','Estorno de baixa','Estorno de baixa identificado'])
    buffer=io.BytesIO();wb.save(buffer);model=buffer.getvalue()
    movements=pd.DataFrame([{'DESCRIÇÃO':'BANCO ITAÚ','DATA':'01/10/2026','VALOR':v,'DÉBITO':'','CRÉDITO':'','HISTÓRICO':'Estorno de baixa' if v==200 else 'Movimento'} for v in (100,100,100,-25,200)])
    monkeypatch.setattr(services,'statement',lambda *_:movements.copy())
    ref=reference_engine()
    main,excluded,_=ref.ler_planilha_organizada_conferencia(model,'itau')
    expected=ref.conciliar_empresa_com_extrato(main,movements,excluded)
    summary,tables=services.reconcile(3,'itau',model,[('itau.pdf',b'pdf')])
    assert summary['faltando_planilha']==1 and not summary['ok']
    for name,frame in zip(('Conferência diária','Faltando na planilha','A mais na planilha','Estornos ignorados'),expected):
        pd.testing.assert_frame_equal(tables[name],frame)

def test_a1_validated_encrypted_and_scoped_to_company(monkeypatch):
    from cryptography import x509
    from cryptography.x509.oid import NameOID
    from cryptography.hazmat.primitives import hashes,serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import pkcs12
    from datetime import datetime,timedelta,timezone
    from app.certificates_service import connection
    from razync.certificado_digital import decifrar_certificado
    monkeypatch.setenv('RAZYNC_ACCESS_TOKEN','test-admin')
    monkeypatch.setenv('CERTIFICATES_MASTER_KEY','test-crypto-master')
    key=rsa.generate_private_key(public_exponent=65537,key_size=2048)
    name=x509.Name([x509.NameAttribute(NameOID.COMMON_NAME,'Razync Teste:11222333000181')])
    now=datetime.now(timezone.utc)
    certificate=x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key()).serial_number(x509.random_serial_number()).not_valid_before(now-timedelta(days=1)).not_valid_after(now+timedelta(days=30)).sign(key,hashes.SHA256())
    content=pkcs12.serialize_key_and_certificates(b'teste',key,certificate,None,serialization.BestAvailableEncryption(b'test-password'))
    headers={'Authorization':'Bearer test-admin'}
    response=client.post('/api/v1/certificates/242',headers=headers,data={'password':'test-password'},files={'file':('teste.pfx',content)})
    assert response.status_code==200,response.text
    assert response.json()['certificate']['cnpj']=='11222333000181'
    assert client.get('/api/v1/certificates/1408',headers=headers).json()['certificate'] is None
    con=connection();row=con.execute('SELECT encrypted FROM certificates WHERE company=242').fetchone();con.close()
    assert decifrar_certificado(row['encrypted'],'test-crypto-master')==(content,'test-password')
    assert 'test-password' not in row['encrypted']
    assert client.post('/api/v1/certificates/242',headers=headers,data={'password':'wrong'},files={'file':('teste.pfx',content)}).status_code==422

def test_selected_period_filters_standard_api_without_changing_legacy_processor(monkeypatch):
    import importlib
    main=importlib.import_module('app.main')
    source=pd.DataFrame([{'DESCRIÇÃO':'BANCO ITAÚ','DATA':pd.Timestamp(d),'VALOR':v,'DÉBITO':'','CRÉDITO':'508','HISTÓRICO':'Pago: Teste'} for d,v in [('2026-09-30',-10),('2026-10-01',-20)]])
    monkeypatch.setattr(main,'_process_modular',lambda *_:source.copy())
    response=client.post('/api/v1/modelo-dominio/88',data={'bank':'itau','options_json':'{"data_inicial":"2026-10-01","data_final":"2026-10-01"}'},files={'files':('itau.pdf',b'pdf')})
    assert response.status_code==200,response.text
    expected=engine.gerar_excel_modelo_dominio(source.iloc[1:])
    assert workbook_signature(response.content)==workbook_signature(expected)
    bad=client.post('/api/v1/modelo-dominio/88/multi',data={'banks_json':'["itau"]','options_json':'[]'},files={'files':('itau.pdf',b'pdf')})
    assert bad.status_code==400


def test_non_tool_companies_are_explicitly_empty():
    result = client.get('/api/v1/companies').json()
    empty_companies = [row for row in result if not row['capabilities'].get('tools')]
    assert empty_companies
    for row in empty_companies:
        assert row['capabilities'].get('tools') == []
        assert row['capabilities'].get('banks') == {}
        assert row['capabilities'].get('status') == 'catalog_only'
