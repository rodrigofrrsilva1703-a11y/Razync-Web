import concurrent.futures
import threading
import time
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from app import fiscal_ai, fiscal_ai_sessions as sessions
from app.main import app


@pytest.fixture(autouse=True)
def isolate(monkeypatch):
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.setattr(sessions, "_dispatch_cursor", {"curto": 0, "medio": 0})
    with sessions._guard:
        sessions._sessions.clear()
        sessions._completed_cache.clear()
    monkeypatch.setattr(fiscal_ai, "_groq_enabled", lambda: True)
    monkeypatch.setattr(fiscal_ai, "_gemini_free_key", lambda: "synthetic")
    yield
    with sessions._guard:
        sessions._sessions.clear()
        sessions._completed_cache.clear()


def report(n=17):
    return {"contas": [{"conta": str(i), "tipo": "ENTRADAS", "situacao": "REVISAR",
             "fiscal": 10, "contabil": 8, "total_conta": 8, "diferenca": -2, "extras": 0}
             for i in range(n)],
            "lancamentos": [{"conta": str(i % n), "historico": "Compra fictícia" + ("x" * 181 if n == 4 and i % n >= 2 else ""), "debito": 1, "credito": 0}
                            for i in range(2001 if n == 65 else n * 2)]}


def test_lotes_preservam_mais_de_12_grupos_e_1500_registros():
    source = report(65)
    manifest = sessions.prepare(source)
    batches = sessions._get(manifest["sessao"])["batches"]
    assert manifest["grupos"] == 65
    assert manifest["registros"] == 2001
    assert manifest["cooperacao"] is True
    assert {batch["provedor"] for batch in batches} == {"gemini"}
    assert sum(len(batch["report"]["contas"]) for batch in batches) == 65
    refs = [row["_referencia_ia"] for batch in batches for row in batch["report"]["lancamentos"]]
    assert len(refs) == len(set(refs)) == 2001
    assert set(refs) == {f"L{i}" for i in range(1, 2002)}
    assert "_referencia_ia" not in source["lancamentos"][0]
    assert all(batch["grupos"] <= 2 for batch in manifest["lotes"])
    assert all("report" not in batch for batch in manifest["lotes"])


def test_provedores_processam_simultaneamente_sem_mudar_ambiente(monkeypatch):
    monkeypatch.setenv("RAZYNC_AI_PRIMARY", "groq")
    manifest = sessions.prepare(report(4))
    barrier = threading.Barrier(2)
    seen = []
    def explain(batch, preferred_provider, **kwargs):
        seen.append(preferred_provider)
        barrier.wait(timeout=3)
        return {"analises": [{"conta": row["conta"]} for row in batch["contas"]],
                "provedor": preferred_provider, "modelo_usado": "synthetic", "gratuito": True}
    monkeypatch.setattr(fiscal_ai, "explain", explain)
    with concurrent.futures.ThreadPoolExecutor(2) as pool:
        futures = [pool.submit(sessions.process, manifest["sessao"], i) for i in (0,1)]
        results = [future.result(timeout=5) for future in futures]
    assert set(seen) == {"gemini", "groq"}
    assert {row["conta"] for result in results for row in result["analises"]} == {"0","1","2","3"}
    assert results[0]["analises"][0]["provedor"] == "groq"
    import os
    assert os.environ["RAZYNC_AI_PRIMARY"] == "groq"


def test_lote_concluido_e_cacheado_sem_gastar_tokens_novamente(monkeypatch):
    manifest = sessions.prepare(report(2))
    calls = []
    monkeypatch.setattr(fiscal_ai, "explain", lambda source, **k: calls.append(1) or {"analises": [{"conta": r["conta"]} for r in source["contas"]], "provedor":"groq"})
    first = sessions.process(manifest["sessao"], 0)
    assert sessions.process(manifest["sessao"], 0) == first
    assert calls == [1, 1]  # Uma chamada isolada por conta; repetir não chama a IA.


def test_falha_nao_descarta_outro_lote_e_pode_ser_repetida(monkeypatch):
    manifest = sessions.prepare(report(4))
    def explain(*_a, preferred_provider, **kwargs):
        if preferred_provider == "gemini":
            raise HTTPException(429, "Cota esgotada")
        return {"analises": [{"conta": r["conta"]} for r in _a[0]["contas"]], "provedor":"groq"}
    monkeypatch.setattr(fiscal_ai, "explain", explain)
    first = sessions.process(manifest["sessao"], 0)
    with pytest.raises(HTTPException) as error:
        sessions.process(manifest["sessao"], 1)
    assert error.value.status_code == 429
    assert sessions.process(manifest["sessao"], 0) == first
    assert sessions._get(manifest["sessao"])["batches"][1]["running"] is False


def test_cancelamento_e_expiracao_descartam_contexto():
    manifest = sessions.prepare(report(2))
    sessions.discard(manifest["sessao"])
    with pytest.raises(HTTPException) as error:
        sessions.process(manifest["sessao"], 0)
    assert error.value.status_code == 410
    manifest = sessions.prepare(report(2))
    sessions._get(manifest["sessao"])["expires"] = time.monotonic() - 1
    with pytest.raises(HTTPException):
        sessions._get(manifest["sessao"])
    assert manifest["sessao"] not in sessions._sessions


def test_lote_invalido_e_requisicao_duplicada_nao_chamam_ia(monkeypatch):
    manifest = sessions.prepare(report(2))
    monkeypatch.setattr(fiscal_ai, "explain", lambda *a, **k: pytest.fail("Não deve chamar IA"))
    with pytest.raises(HTTPException) as error:
        sessions.process(manifest["sessao"], 99)
    assert error.value.status_code == 404
    session = sessions._get(manifest["sessao"])
    session["active_providers"].add("groq")
    with pytest.raises(HTTPException) as error:
        sessions.process(manifest["sessao"], 0)
    assert error.value.status_code == 409


def test_refs_do_contexto_permanecem_globais():
    manifest = sessions.prepare(report(4))
    second = sessions._get(manifest["sessao"])["batches"][1]["report"]
    _, _, refs, _ = fiscal_ai.detailed_context(second)
    assert "L3" in refs and "L1" not in refs
    assert refs["L3"]["conta"] == "2"


def test_excel_exporta_mais_de_15_contas_sem_ia(monkeypatch):
    monkeypatch.setattr(fiscal_ai, "explain", lambda *a, **k: pytest.fail("Exportação não chama IA"))
    payload = {"analises": [{"conta":str(i), "tipo":"ENTRADAS", "explicacao":"Exemplo", "verificar":"Conferir"} for i in range(65)]}
    r = TestClient(app).post("/api/v1/conferencia-fiscal/ia/exportar", json=payload)
    assert r.status_code == 200
    import io
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(r.content))
    assert workbook.worksheets[0].max_row == 70


def test_preferencias_explicitas_preservam_provedores(monkeypatch):
    monkeypatch.setenv("GROQ_API_KEY", "synthetic")
    monkeypatch.setenv("GROQ_MODEL", "openai/gpt-oss-120b")
    monkeypatch.setenv("GEMINI_FREE_MODEL", "gemini-3.1-flash-lite")
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.setattr(fiscal_ai, "gemini_cooldown_until", 0)
    answer = {"analises": [{"grupo":"G1", "explicacao":"Revisar", "verificar":"Conferir", "evidencias":[]}]}
    calls = []
    monkeypatch.setattr(fiscal_ai, "_groq_completion", lambda *a, **k: calls.append("groq") or (answer, "openai/gpt-oss-120b"))
    monkeypatch.setattr(fiscal_ai, "_gemini_completion", lambda *a, **k: calls.append("gemini") or (answer, "gemini-3.1-flash-lite"))
    assert fiscal_ai.explain(report(1), preferred_provider="gemini")["provedor"] == "gemini"
    assert fiscal_ai.explain(report(1), preferred_provider="groq")["provedor"] == "groq"
    assert calls == ["gemini", "groq"]



def test_api_le_arquivos_uma_vez_e_cacheia_resultado(monkeypatch):
    read_calls = []
    ai_calls = []
    monkeypatch.setattr(fiscal_ai, "conferencia_fiscal_preview", lambda *a: read_calls.append(1) or report(4))
    monkeypatch.setattr(fiscal_ai, "explain", lambda batch, preferred_provider, **kwargs: ai_calls.append(preferred_provider) or {
        "analises": [{"conta": row["conta"]} for row in batch["contas"]], "provedor": preferred_provider, "gratuito": True})
    client = TestClient(app)
    response = client.post("/api/v1/conferencia-fiscal/ia/sessoes", files={"acumuladores": ("fiscal.xlsx", b"synthetic"), "razao": ("razao.xlsx", b"synthetic")})
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    token = response.json()["sessao"]
    for batch in (0, 1, 0):
        response = client.post(f"/api/v1/conferencia-fiscal/ia/sessoes/{token}/lotes/{batch}")
        assert response.status_code == 200
        assert response.headers["cache-control"] == "no-store"
    assert read_calls == [1]
    assert ai_calls == ["groq", "groq", "gemini", "gemini"]
    assert client.delete(f"/api/v1/conferencia-fiscal/ia/sessoes/{token}").status_code == 200
    assert client.post(f"/api/v1/conferencia-fiscal/ia/sessoes/{token}/lotes/0").status_code == 410



def test_limpeza_cancela_temporizador_da_sessao(monkeypatch):
    timers = []
    class Timer:
        def __init__(self, *args, **kwargs):
            self.started = self.cancelled = False
            timers.append(self)
        def start(self): self.started = True
        def cancel(self): self.cancelled = True
    monkeypatch.setattr(sessions.threading, "Timer", Timer)
    manifest = sessions.prepare(report(1))
    assert timers[0].started is True
    sessions.discard(manifest["sessao"])
    assert timers[0].cancelled is True


def test_openrouter_divide_carga_curta_inclusive_entre_sessoes(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "synthetic")
    providers = []
    for _ in range(4):
        manifest = sessions.prepare(report(2))
        providers.append(manifest["lotes"][0]["provedor"])
        assert set(manifest["provedores"]) == {"groq", "gemini", "openrouter"}
        sessions.discard(manifest["sessao"])
    assert providers == ["groq", "groq", "groq", "openrouter"]


def test_contas_extensas_nunca_usam_reserva_com_amostragem(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "synthetic")
    source = report(2)
    for row in source["lancamentos"]:
        row["historico"] = "x" * 600
    manifest = sessions.prepare(source)
    batch = sessions._get(manifest["sessao"])["batches"][0]
    assert batch["provedor"] == "gemini"
    assert batch["reservas"] == ["openrouter"]
    assert batch["report"]["lancamentos"][0]["historico"] == "x" * 600


def test_reserva_valida_so_lote_pendente_e_cache_nao_reenvia(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "synthetic")
    manifest = sessions.prepare(report(2))
    seen = []
    def explain(source, preferred_provider, *, only_provider):
        assert preferred_provider == only_provider
        seen.append(only_provider)
        if only_provider == "groq":
            raise HTTPException(429, "Cota", headers={"Retry-After":"30"})
        return {"analises":[{"conta":row["conta"]} for row in source["contas"]],
                "provedor":only_provider, "gratuito":True}
    monkeypatch.setattr(fiscal_ai, "explain", explain)
    first = sessions.process(manifest["sessao"], 0)
    assert first["provedor"] == "openrouter" and first["fallback_usado"] is True
    assert sessions.process(manifest["sessao"], 0) == first
    assert seen == ["groq", "openrouter", "openrouter"]


def test_sessao_nao_ativa_gemini_legado_sem_confirmacao_gratuita(monkeypatch):
    monkeypatch.setattr(fiscal_ai, "_groq_enabled", lambda:False)
    monkeypatch.setattr(fiscal_ai, "_gemini_free_key", lambda:"")
    monkeypatch.setenv("RAZYNC_AI_LEGACY_GEMINI", "1")
    monkeypatch.setenv("GEMINI_API_KEY", "synthetic")
    with pytest.raises(HTTPException) as error:
        sessions.prepare(report(2))
    assert error.value.status_code == 503


def test_reservas_ocupadas_pedem_espera_sem_reenviar_outros_lotes(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "synthetic")
    manifest = sessions.prepare(report(2))
    session = sessions._get(manifest["sessao"])
    session["active_providers"].update({"gemini", "openrouter"})
    calls = []
    def fail(*args, only_provider, **kwargs):
        calls.append(only_provider)
        raise HTTPException(502, "Indisponível")
    monkeypatch.setattr(fiscal_ai, "explain", fail)
    with pytest.raises(HTTPException) as error:
        sessions.process(manifest["sessao"], 0)
    assert error.value.status_code == 429
    assert error.value.headers["Retry-After"] == "15"
    assert calls == ["groq"]
    assert session["batches"][0]["result"] is None
    assert session["batches"][0]["running"] is False


def test_cache_entre_sessoes_reutiliza_sem_nova_chamada(monkeypatch):
    calls = []
    def explain(source, preferred_provider, **kwargs):
        calls.append(preferred_provider)
        return {"analises": [{"conta": r["conta"]} for r in source["contas"]],
                "provedor": preferred_provider, "gratuito": True}
    monkeypatch.setattr(fiscal_ai, "explain", explain)
    first = sessions.prepare(report(2))
    sessions.process(first["sessao"], 0)
    sessions.discard(first["sessao"])
    second = sessions.prepare(report(2))
    cached = sessions.process(second["sessao"], 0)
    assert cached["cache_reutilizado"] is True
    assert len(calls) == 2


def test_resposta_incompleta_rejeitada_sem_cache(monkeypatch):
    manifest = sessions.prepare(report(2))
    monkeypatch.setattr(fiscal_ai, "explain", lambda *a, **k: {"analises": []})
    with pytest.raises(HTTPException):
        sessions.process(manifest["sessao"], 0)
    assert sessions._get(manifest["sessao"])["batches"][0]["result"] is None
    assert sessions._completed_cache == {}


def test_mesma_conta_em_duas_naturezas_e_aceita(monkeypatch):
    source = report(2)
    source["contas"][1].update(conta="0", tipo="SAÍDAS")
    def explain(batch, **kwargs):
        return {"analises": [dict(conta=r["conta"],tipo=r["tipo"]) for r in batch["contas"]],
                "provedor":"groq","gratuito":True}
    monkeypatch.setattr(fiscal_ai,"explain",explain)
    manifest = sessions.prepare(source)
    result = sessions.process(manifest["sessao"],0)
    assert {(i["conta"],i["tipo"]) for i in result["analises"]} == {("0","ENTRADAS"),("0","SAÍDAS")}


def test_alterar_uma_conta_reusa_outra_mesmo_em_outro_lote(monkeypatch):
    calls = []
    def explain(batch, **kwargs):
        calls.append([r["conta"] for r in batch["contas"]])
        return {"analises": [dict(conta=r["conta"],tipo=r["tipo"]) for r in batch["contas"]],
                "provedor":"groq","gratuito":True}
    monkeypatch.setattr(fiscal_ai,"explain",explain)
    initial = sessions.prepare(report(4))
    for i in (0,1):
        sessions.process(initial["sessao"],i)
    sessions.discard(initial["sessao"])
    changed = report(4)
    changed["contas"][0]["fiscal"] = 99
    changed["contas"] = list(reversed(changed["contas"]))
    repeated = sessions.prepare(changed)
    results = [sessions.process(repeated["sessao"],i) for i in (0,1)]
    assert calls == [["0"],["1"],["2"],["3"],["0"]]
    assert sum(len(result["analises"]) for result in results) == 4


def test_nao_reutiliza_parecer_de_outra_empresa(monkeypatch):
    calls = []
    def explain(batch, **kwargs):
        calls.append(1)
        return {"analises": [dict(conta=r["conta"],tipo=r["tipo"]) for r in batch["contas"]],
                "provedor":"groq","gratuito":True}
    monkeypatch.setattr(fiscal_ai,"explain",explain)
    for name in ("EMPRESA A","EMPRESA B"):
        source = dict(report(2),empresa_nome=name)
        manifest = sessions.prepare(source)
        sessions.process(manifest["sessao"],0)
        sessions.discard(manifest["sessao"])
    assert len(calls) == 4


def test_duas_pessoas_com_mesma_conta_nao_duplicam_chamada(monkeypatch):
    first = sessions.prepare(report(2))
    second = sessions.prepare(report(2))
    entered, release = threading.Event(), threading.Event()
    calls = []
    def explain(batch, **kwargs):
        calls.append(1)
        entered.set()
        assert release.wait(5)
        return {"analises": [dict(conta=r["conta"],tipo=r["tipo"]) for r in batch["contas"]],
                "provedor":"groq","gratuito":True}
    monkeypatch.setattr(fiscal_ai,"explain",explain)
    with concurrent.futures.ThreadPoolExecutor(1) as executor:
        future = executor.submit(sessions.process,first["sessao"],0)
        assert entered.wait(3)
        try:
            with pytest.raises(HTTPException) as exc:
                sessions.process(second["sessao"],0)
            assert exc.value.status_code == 429
        finally:
            release.set()
        future.result(timeout=5)
    assert sessions.process(second["sessao"],0)["cache_reutilizado"]
    assert calls == [1, 1]


def test_provedor_com_fila_cheia_e_reserva_nao_primeira_opcao(monkeypatch):
    monkeypatch.setattr(sessions,"provider_wait_seconds",lambda p: 30 if p == "groq" else 0)
    manifest = sessions.prepare(report(2))
    assert manifest["lotes"][0]["provedor"] == "gemini"


def test_seis_pessoas_podem_manter_as_duas_etapas_abertas():
    tokens = []
    try:
        for _ in range(12):
            tokens.append(sessions.prepare(report(2))["sessao"])
        assert len(sessions._sessions) == 12
    finally:
        for token in tokens:
            sessions.discard(token)
