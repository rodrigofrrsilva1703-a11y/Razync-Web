"""Regressões empresa 266: BB 200076-8 na mesma consolidada dos demais bancos."""
import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import engine, operating_hours
from app.advanced import workflow_modelo
from app.main import app
from app.migration_services import workflow
from app.nova_geracao_bb_266 import (
    ABA_MODELO, CONTA_BANCARIA, CONTA_CONTABIL,
    processar_nova_geracao_bb_266,
)
from app.registry import CAPABILITIES
from test_migration import xlsx, workbook_signature


def consolidated():
    return xlsx([
        ["Conta", "Data", "Valor", "Lacto", "Historico", "Doc"],
        ["99549-5", "02/10/2026", "100,00", "RECEBER", "Cliente Itaú", "I1"],
        ["451990-6", "02/10/2026", "60,00", "PAGAR", "Fornecedor Bradesco", "B1"],
        ["673947-1", "02/10/2026", "70,00", "RECEBER", "Cliente Fibra", "F1"],
        ["200076-8", "03/10/2026", "1.234,56", "RECEBER", "Cliente BB", "BB1"],
        ["2000768", "04/10/2026", "25,00", "PAGAR", "Fornecedor BB", "BB2"],
        ["200076-9", "04/10/2026", "9.999,99", "RECEBER", "OUTRA CONTA", "X1"],
        ["1200076-8", "04/10/2026", "9.999,99", "RECEBER", "CONTA SEMELHANTE", "X2"],
        ["200076-8", "05/10/2026", "90,00", "ESTORNO DE BAIXA", "Estorno baixa BB", "BB3"],
    ], title="Consolidada")


def test_cadastro_266_inclui_bb_conta_contabil_8():
    banks = CAPABILITIES[266]["banks"]
    assert banks == {"itau": "508", "bradesco": "9", "fibra": "506", "banco_brasil": "8"}
    assert CAPABILITIES[1396]["banks"] == {"itau": "515", "bradesco": "514"}
    assert CONTA_BANCARIA == "200076-8"
    assert CONTA_CONTABIL == "8"


def test_leitor_bb_filtra_apenas_conta_200076_8_e_separa_estornos():
    with engine.processing_context("nova_geracao_matriz"):
        principal, retirados = processar_nova_geracao_bb_266(consolidated())
    assert principal["VALOR"].tolist() == [1234.56, -25.0]
    assert principal["DESCRIÇÃO"].tolist() == ["BANCO DO BRASIL"] * 2
    assert principal["DÉBITO"].tolist() == ["8", ""]
    assert principal["CRÉDITO"].tolist() == ["", "8"]
    assert "Cliente BB" in principal.iloc[0]["HISTÓRICO"]
    assert "Fornecedor BB" in principal.iloc[1]["HISTÓRICO"]
    assert len(retirados) == 1
    assert retirados.iloc[0]["MOTIVO"] == "Estorno de baixa identificado"
    assert "OUTRA CONTA" not in " ".join(principal["HISTÓRICO"])


def test_selecionar_bb_cria_aba_propria_com_conta_bancaria_no_titulo():
    result, filename = workflow(
        266, {"consolidada": [("consolidado.xlsx", consolidated())]},
        {"bancos": ["banco_brasil"]}
    )
    assert filename == "RAZYNC_266_MODELO_DOMINIO.xlsx"
    wb = load_workbook(io.BytesIO(result))
    assert wb.sheetnames == [ABA_MODELO, "Lançamentos retirados"]
    rows = list(wb[ABA_MODELO].values)
    assert len(rows) == 3
    assert rows[1][0] == "BANCO DO BRASIL"
    assert rows[1][2] == 1234.56 and str(rows[1][3]) == "8" and not rows[1][4]
    assert rows[2][2] == -25 and not rows[2][3] and str(rows[2][4]) == "8"


def test_266_gera_quatro_abas_sem_alterar_as_tres_existentes():
    result, _ = workflow(266, {"consolidada": [("consolidado.xlsx", consolidated())]}, {})
    wb = load_workbook(io.BytesIO(result))
    assert wb.sheetnames == [
        "Itaú", "Bradesco", "Fibra", ABA_MODELO, "Lançamentos retirados"
    ]
    assert wb["Itaú"]["A2"].value == "BANCO ITAÚ"
    assert wb["Bradesco"]["A2"].value == "BANCO BRADESCO"
    assert wb["Fibra"]["A2"].value == "BANCO FIBRA"
    assert wb[ABA_MODELO]["A2"].value == "BANCO DO BRASIL"


def test_bb_exige_conta_correta_quando_selecionado():
    bad = xlsx([["Conta", "Data", "Valor", "Lacto", "Historico", "Doc"],
                ["200076-9", "01/10/2026", 500, "RECEBER", "NÃO É BB 200076-8", "X"]])
    with pytest.raises(ValueError, match="200076-8"):
        workflow(266, {"consolidada": [("relatorio.xlsx", bad)]},
                 {"bancos": ["banco_brasil"]})


def test_266_regra_mesma_no_fluxo_alternativo_e_periodo():
    options = {"bancos": ["banco_brasil"], "data_inicial": "2026-10-03",
               "data_final": "2026-10-03"}
    roles = {"consolidada": [("consolidado.xlsx", consolidated())]}
    actual, _ = workflow_modelo(266, roles, options)
    expected, _ = workflow(266, roles, options)
    assert workbook_signature(actual) == workbook_signature(expected)
    wb = load_workbook(io.BytesIO(actual))
    assert wb[ABA_MODELO].max_row == 2
    assert wb[ABA_MODELO]["C2"].value == 1234.56


def test_endpoint_api_retorna_bb_8_e_aceita_aba_bb(monkeypatch):
    monkeypatch.setattr(operating_hours, "is_open", lambda *args: True)
    client = TestClient(app)
    detail = client.get("/api/v1/companies/266")
    assert detail.status_code == 200
    assert detail.json()["capabilities"]["banks"]["banco_brasil"] == "8"
    response = client.post(
        "/api/v1/workflow/266",
        data={"roles_json": json.dumps(["consolidada"]),
              "options_json": json.dumps({"bancos": ["banco_brasil"]})},
        files={"files": ("relatorio.xlsx", consolidated(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    assert ABA_MODELO in load_workbook(io.BytesIO(response.content)).sheetnames
