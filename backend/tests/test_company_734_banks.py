"""734 BTG: apenas PDF, leitura de todas as linhas e reconciliação de saldos."""
import io
from unittest.mock import patch

import fitz
import pytest
import pandas as pd
from openpyxl import load_workbook

from app.bank_734 import read_bank_pdf, workflow_734, _clean_history
from app.registry import CAPABILITIES
from app.migration_services import _workflow


def btg_pdf(with_bad_total=False):
    doc = fitz.open()
    page = doc.new_page()
    first = [
        "Conta corrente - PJ",
        "Razao social",
        "EMPRESA DE EXEMPLO LTDA",
        "Periodo do extrato: 01/07/2026 - 31/07/2026",
        "Período do extrato: 01/07/2026 - 31/07/2026",
        "Saldo de abertura em 01/07/2026:",
        "R$ 100,00",
        "Saldo de fechamento em 31/07/2026:",
        "R$ 150,27",
        "Total de entradas",
        "R$ 100,27",
        "Total de saídas",
        ("R$ 50,01" if with_bad_total else "R$ 50,00"),
        "Data lançamento",
        "Descrição do lançamento",
        "31/07/2026",
        "Saldo de fechamento",
        "150,27",
        "31/07/2026",
        "Valor de Rendimento Remunera+",
        "0,27",
        "150,27",
        "30/07/2026",
        "Pix recebido de CLIENTE TESTE",
        "Banco 341 | CPF ***.123.456-**",
        "100,00",
        "150,00",
        "30/07/2026",
        "Pagamento de Conta / Tributo - TESTE",
        "-50,00",
        "50,00",
        "Fale com nossa central",
    ]
    for i, line in enumerate(first):
        page.insert_text((28, 36 + 18 * i), line, fontsize=8)
    page = doc.new_page()
    second = [
        "Data lançamento",
        "Descrição do lançamento",
        "01/07/2026",
        "Saldo de abertura",
        "100,00",
        "Fale com nossa central",
    ]
    for i, line in enumerate(second):
        page.insert_text((28, 36 + 18 * i), line, fontsize=8)
    content = doc.tobytes()
    doc.close()
    return content


def test_734_tem_somente_btg_e_conta_510():
    cap = CAPABILITIES[734]
    assert cap["banks"] == {"btg": "510"}
    assert [role["name"] for role in cap["roles"]] == ["btg"]
    assert cap["roles"][0]["accept"] == ".pdf"
    assert cap["roles"][0]["multiple"] is True
    assert CAPABILITIES[1402]["banks"]["btg"] == "510"


def test_leitor_btg_confere_totais_saldos_e_multilinhas():
    result = read_bank_pdf("btg", btg_pdf(), "extrato.pdf")
    assert len(result) == 3
    assert result["VALOR"].tolist() == [0.27, 100, -50]
    assert result["DÉBITO"].tolist() == ["510", "510", ""]
    assert result["CRÉDITO"].tolist() == ["", "", "510"]
    assert "Banco 341" in result.iloc[1]["HISTÓRICO"]
    assert "Saldo" not in " ".join(result["HISTÓRICO"])


def test_btg_recusa_saldo_que_nao_fecha_e_arquivo_invalido():
    with pytest.raises(ValueError, match="não conferem"):
        read_bank_pdf("btg", btg_pdf(with_bad_total=True), "extrato.pdf")
    with pytest.raises(ValueError, match="PDF"):
        read_bank_pdf("btg", b"not pdf", "extrato.xlsx")
    with pytest.raises(ValueError, match="somente o BTG"):
        read_bank_pdf("bradesco", btg_pdf(), "extrato.pdf")


def test_modelo_734_tem_uma_aba_btg_e_respeita_periodo():
    result, filename = _workflow(734, {"btg": [("extrato.pdf", btg_pdf())]}, {
        "data_inicial": "2026-07-01", "data_final": "2026-07-31",
    })
    assert filename == "RAZYNC_734_BTG_510_MODELO_DOMINIO.xlsx"
    wb = load_workbook(io.BytesIO(result))
    assert any("BTG 510" in name for name in wb.sheetnames)
    assert not any("Bradesco" in name for name in wb.sheetnames)
    with pytest.raises(ValueError, match="período"):
        workflow_734({"btg": [("extrato.pdf", btg_pdf())]}, {
            "data_inicial": "2026-08-01", "data_final": "2026-08-31",
        })


def test_upload_repetido_recusado_e_bradesco_adiado():
    file = btg_pdf()
    with pytest.raises(ValueError, match="mais de uma vez"):
        workflow_734({"btg": [("a.pdf", file), ("b.pdf", file)]}, {})
    with pytest.raises(ValueError, match="somente para o BTG"):
        workflow_734({"bradesco": [("bradesco.pdf", file)]}, {})


def test_remove_cpf_mascarado_mantendo_nome_banco_e_mensagem():
    source = "Pix recebido de CLIENTE TESTE Banco 341 | CPF ***.558.948-** | Mensagem - consulta julho"
    cleaned = _clean_history(source)
    assert "CPF" not in cleaned
    assert "558.948" not in cleaned
    assert "CLIENTE TESTE" in cleaned
    assert "Banco 341" in cleaned
    assert "Mensagem - consulta julho" in cleaned


def test_remove_cpf_completo_e_sem_pontuacao():
    assert "123.456.789-09" not in _clean_history("Pagamento CPF 123.456.789-09 referente junho")
    assert "12345678909" not in _clean_history("CPF 12345678909 Pix recebido")
    assert _clean_history("Rendimento Remunera+") == "Rendimento Remunera+"


def test_historico_do_excel_734_nao_contem_cpf():
    # Usa o mesmo formato de PIX e CPF mascarado presente no extrato real.
    from app.bank_734 import _clean_history
    history = _clean_history("Pix recebido de CLIENTE Banco 341 | CPF ***.558.948-** | Mensagem - julho")
    assert "CPF" not in history
    assert "Mensagem - julho" in history


def test_btg_descarta_debito_e_credito_na_conta_corrente_sem_descartar_outros():
    from app.bank_734 import _skip_current_account_transfer
    assert _skip_current_account_transfer("Débito na conta corrente")
    assert _skip_current_account_transfer("Credito na conta corrente")
    assert _skip_current_account_transfer("Crédito na conta corrente - movimento BTG")
    assert _skip_current_account_transfer("  DEBITO  NA CONTA CORRENTE ")
    assert not _skip_current_account_transfer("Pix recebido na conta corrente")
    assert not _skip_current_account_transfer("Pagamento de Conta / Tributo")
    assert not _skip_current_account_transfer("Rendimento Remunera+")


def test_btg_exclui_movimentos_somente_depois_de_validar_saldo(monkeypatch):
    from app import bank_734
    base = [
        {"DESCRIÇÃO": "BANCO BTG", "DATA": pd.Timestamp("2026-07-03"), "VALOR": 200,
         "DÉBITO": "510", "CRÉDITO": "", "HISTÓRICO": "Credito na conta corrente"},
        {"DESCRIÇÃO": "BANCO BTG", "DATA": pd.Timestamp("2026-07-03"), "VALOR": -100,
         "DÉBITO": "", "CRÉDITO": "510", "HISTÓRICO": "Débito na conta corrente"},
        {"DESCRIÇÃO": "BANCO BTG", "DATA": pd.Timestamp("2026-07-04"), "VALOR": -50,
         "DÉBITO": "", "CRÉDITO": "510", "HISTÓRICO": "Pix enviado"},
    ]
    monkeypatch.setattr(bank_734, "_parse_pdf", lambda pdf: base)
    result = read_bank_pdf("btg", btg_pdf(), "extrato.pdf")
    assert len(result) == 1
    assert result.iloc[0]["HISTÓRICO"] == "Pix enviado"
    assert result.iloc[0]["CRÉDITO"] == "510"
