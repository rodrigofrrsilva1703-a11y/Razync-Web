"""Investigações explicativas sem enviar arquivos financeiros a provedores."""
import copy
import io
import pytest
from openpyxl import load_workbook
from app.fiscal_investigation import build_cases, case_context, explain_cases
from app.fiscal_ai_audit import protect_publication, structured_analysis
from app import fiscal_ai, fiscal_ai_sessions, groq_fiscal


def source(status="VALOR_DIFERENTE"):
    return {"conta":"470", "comparacoes":[{"id":"E-1","nota":"100","serie":"1","data":"01/09/2026",
        "parceiro":"Fornecedor", "situacao":status,"criterio":"Nota e data comparadas",
        "valor_fiscal":100,"valor_razao":110,"diferenca":10,
        "linhas_fiscais":[{"id":"E-1"},{"id":"E-2"}],
        "lancamentos":[{"id":"R-7","conta":"470","data":"01/09/2026","valor":110,
                          "historico":"COMPRA NF 100","contrapartida":"508"}]}],
        "razao_sem_vinculo":[{"id":"R-9","valor":9.99,"historico":"Pago: tarifa","data":"01/09/2026","contrapartida":"508"}]}


@pytest.mark.parametrize("status,grade",[("VALOR_DIFERENTE","Candidato forte"),("DATA_DIFERENTE","Candidato forte"),
    ("CONTA_DIFERENTE","Candidato forte"),("AMBIGUO","Ambíguo"),("PROVAVEL","Ambíguo"),
    ("ESTORNO_EM_ANALISE","Ambíguo"),("SEM_LANCAMENTO_IDENTIFICADO","Evidência insuficiente")])
def test_roteiros_distintos_explicam_evidencia_sem_confirmar_causa(status, grade):
    cases = build_cases(source(status))
    assert cases[0]["grau_evidencia"] == grade
    assert cases[0]["referencias_fiscais"] == ["E-1","E-2"]
    assert cases[0]["diferenca"] == 10
    assert cases[0]["acao"] and cases[0]["falta_confirmar"]
    assert cases[1]["diferenca"] is None
    assert cases[1]["candidatos"][0]["referencia"] == "L9"


def test_parcelas_agrupadas_geram_um_caso_e_conciliados_nao_entram():
    group = source()
    assert len(build_cases(group)) == 2
    group["comparacoes"][0]["situacao"] = "CONCILIADO"
    assert len(build_cases(group)) == 1


def test_contexto_so_inclui_caso_quando_todas_evidencias_estao_disponiveis():
    cases = build_cases(source())
    docs = [{"referencia":"E-1"},{"referencia":"E-2"}]
    assert not case_context(cases, docs[:1], [{"referencia":"L7"}], "470")
    assert not case_context(cases, docs, [], "470")
    assert case_context(cases, docs, [{"referencia":"L7"}], "470")[0]["id"] == "P-E-1"
    assert "candidatos" in case_context(cases,docs,[{"referencia":"L7"}],"470")[0]


def test_trecho_incorreto_e_retirado_sem_apagar_a_explicacao_valida():
    item = {"conta":"470","explicacao":"Hipóteses: E-1 e L7 podem corresponder à mesma operação; confirme o documento. A nota 999 pertence a este caso. Limitações: Falta conferir a série.",
            "verificar":"1. Compare E-1 com L7.\n2. Confira a nota 999.","evidencias":[]}
    protect_publication(item,[{"referencia":"L7","historico":"Compra NF 100"}],[{"referencia":"E-1","nota":"100"}])
    assert "E-1 e L7" in item["explicacao"]
    assert "999" not in item["explicacao"] + item["verificar"]
    assert "Compare E-1" in item["verificar"]
    assert item["checagem_semantica"]["trechos_retirados"] == 2
    saved = copy.deepcopy(item)
    protect_publication(item,[{"referencia":"L7","historico":"Compra NF 100"}],[{"referencia":"E-1","nota":"100"}])
    assert saved == item


def test_interpretacoes_so_se_associam_a_referencias_do_proprio_caso():
    cases = build_cases(source())
    linked = explain_cases(cases,"E-1 pode corresponder a L7. L9 exige comprovante. E-1 e L999 precisam de revisão.","1. Compare E-2 e L7.")
    assert len(linked[0]["interpretacao_ia"]) == 1
    assert "L9" not in linked[0]["interpretacao_ia"][0]
    assert len(linked[1]["interpretacao_ia"]) == 1
    assert len(linked[0]["verificacoes_ia"]) == 1


def test_detalhes_locais_e_acoes_validadas_chegam_ao_parecer_e_excel():
    cases = build_cases(source())
    item = {"conta":"470","tipo":"ENTRADAS","valores":{"fiscal":100,"contabil":110,"total_conta":119.99,"diferenca":10},
        "explicacao":"Hipóteses: E-1 pode corresponder a L7. Limitações: confira a origem.","verificar":"1. Compare E-1 com L7.",
        "investigacoes_locais":cases,"evidencias":[{"referencia":"L7","historico":"Compra NF 100"}],
        "resumo_documental":{"documentos_lidos":2,"conciliados":0,"documentos_pendentes":2}}
    item["parecer"] = structured_analysis(item,[],documentary=True)
    assert item["parecer"]["leitura_comparacao"]["orientacoes_ia"] == item["verificar"]
    book = load_workbook(io.BytesIO(fiscal_ai.gerar_excel_analise({"analises":[item]})))
    sheet = book["Pendências explicadas"]
    assert sheet.max_row == 3
    assert sheet["G2"].value == 100 and sheet["I2"].value == 10
    assert "E-1" in sheet["K2"].value
    assert book["Evidências por pendência"]["G2"].value == "COMPRA NF 100"


def test_escopo_e_cache_preservam_todos_os_casos_so_da_conta_escolhida():
    report = {"contas":[{"conta":"470","tipo":"ENTRADAS"}],"investigacoes_locais":{"470:ENTRADAS":build_cases(source()),"417:ENTRADAS":[]}}
    scoped = fiscal_ai_sessions._scope_report(report,report["contas"])
    assert list(scoped["investigacoes_locais"]) == ["470:ENTRADAS"]
    assert len(scoped["investigacoes_locais"]["470:ENTRADAS"]) == 2


def test_groq_nao_recebe_caso_cujo_candidato_foi_cortado():
    cases = build_cases(source())
    group = {"grupo":"G1","resumo":{"conta":"470"},"lancamentos":[],
        "roteiro_investigacao":case_context(cases,[{"referencia":"E-1"},{"referencia":"E-2"}],[{"referencia":"L7"}],"470")}
    assert not groq_fiscal._compact_group(group)["roteiro_investigacao"]


def test_campo_desconhecido_nao_e_classificado_como_coincidente():
    case = build_cases(source())[0]
    assert "Número da nota" in case["campos_a_confirmar"]
    assert "Série" in case["campos_a_confirmar"]
    assert "Data" in case["campos_coincidentes"]
    assert "Valor do conjunto" in case["campos_divergentes"]


def test_candidato_em_outra_conta_nao_vira_referencia_local_da_ia():
    group = source("CONTA_DIFERENTE")
    group["comparacoes"][0]["lancamentos"][0]["conta"] = "508"
    cases = build_cases(group)
    selected = case_context(cases,[{"referencia":"E-1"},{"referencia":"E-2"}],[],"470")
    assert selected[0]["candidatos"][0]["conta"] == "508"
    assert selected[0]["candidatos"][0]["referencia"] == ""
    assert "Conta" in cases[0]["campos_divergentes"]


def test_valores_de_compensacao_nao_significam_causa_comprovada():
    item = {"valores":{"diferenca":750.50},"explicacao":"Hipóteses: causa indeterminada.",
        "resumo_documental":{"composicao":{"itens":[{"descricao":"Diferenças","valor":760.49},
                                                     {"descricao":"Fora da comparação inicial","valor":-9.99}]}}}
    observations = structured_analysis(item,[],documentary=True)["leitura_comparacao"]["observacoes"]
    assert any("fecha a diferença" in text and "não comprova" in text for text in observations)
    item["valores"]["diferenca"] = 751.50
    observations = structured_analysis(item,[],documentary=True)["leitura_comparacao"]["observacoes"]
    assert any("Parcela ainda não explicada" in text and "1,00" in text for text in observations)


def test_roteiro_curto_cobre_categorias_diferentes_antes_de_repetir_notas():
    group = source()
    extra = copy.deepcopy(group["comparacoes"][0])
    extra.update(id="E-3",situacao="DATA_DIFERENTE",valor_fiscal=1,valor_razao=1,diferenca=0,
                 linhas_fiscais=[{"id":"E-3"}])
    group["comparacoes"].append(extra)
    rows = [{"referencia":"L7"},{"referencia":"L9"}]
    selected = case_context(build_cases(group),[{"referencia":ref} for ref in ("E-1","E-2","E-3")],rows,"470")
    assert len(selected) == 3
    assert {s["id"] for s in selected} == {"P-E-1","P-E-3","P-R-9"}


def test_orcamento_so_cresce_quando_ha_roteiro_e_tem_teto():
    from app.fiscal_ai_strategy import output_budget
    base = {"situacao":"REVISAR"}
    assert output_budget([dict(base,roteiro_investigacao=[{}])]) == output_budget([base]) + 120
    assert output_budget([dict(base,roteiro_investigacao=[{}]*30)]) == output_budget([base]) + 360
