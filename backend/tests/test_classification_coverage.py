"""Cobertura de linhas da Base Inteligente, independente das regras de classificação."""
import io

from openpyxl import Workbook, load_workbook

from app.classification_coverage import coverage
from app import classification_service


HEADERS = ["DESCRIÇÃO", "DATA", "VALOR", "DÉBITO", "CRÉDITO", "HISTÓRICO"]


def workbook(sheets):
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        ws = book.create_sheet(name)
        ws.append(HEADERS)
        for row in rows:
            ws.append(row)
    buff = io.BytesIO()
    book.save(buff)
    return buff.getvalue()


def row(value=-100, debit="", credit="", history="PAGO TESTE", desc="BANCO DO BRASIL"):
    return [desc, "15/10/2026", value, debit, credit, history]


def test_conta_linhas_classificadas_antes_depois_sem_incluir_retirados():
    before = workbook({
        "Banco do Brasil": [
            row(debit="", credit=8, history="PAGO PRIMEIRO"),
            row(debit=12, credit=8, history="PAGO JA REVISADO"),
            row(value=80, debit=8, credit="", history="RECEBIDO PENDENTE"),
            row(history=""),  # Não é lançamento elegível.
        ],
        "Retirados": [row(debit=8, credit=55, history="ESTORNO")],
    })
    after = workbook({
        "Banco do Brasil": [
            row(debit=187, credit=8, history="PAGO PRIMEIRO"),
            row(debit=12, credit=8, history="PAGO JA REVISADO"),
            row(value=80, debit=8, credit="", history="RECEBIDO PENDENTE"),
            row(history=""),
        ],
        "Retirados": [row(debit=8, credit=55, history="ESTORNO")],
    })
    progress = coverage(before, after, "nova_geracao_matriz")
    assert progress == {
        "total": 3, "classificadas": 2, "pendentes": 1,
        "ja_classificadas": 1, "novas_classificadas": 1,
        "percentual": 66.7,
    }


def test_conta_placeholders_das_abas_242_na_modalidade_consolidada():
    source = workbook({
        "Principal": [row(debit=187, credit=508, history="LINHA NAO CONSIDERADA")],
        "Itaú 508": [
            row(debit=166, credit=508, history="PAGO ALTERADO", desc="BANCO ITAU"),
            row(value=100, debit=508, credit=14, history="RECEBIDO PENDENTE", desc="BANCO ITAU"),
            row(value=100, debit=508, credit=22, history="RECEBIDO COMPLETO", desc="BANCO ITAU"),
        ],
    })
    after = workbook({
        "Principal": [row(debit=187, credit=508, history="LINHA NAO CONSIDERADA")],
        "Itaú 508": [
            row(debit=187, credit=508, history="PAGO ALTERADO", desc="BANCO ITAU"),
            row(value=100, debit=508, credit=14, history="RECEBIDO PENDENTE", desc="BANCO ITAU"),
            row(value=100, debit=508, credit=22, history="RECEBIDO COMPLETO", desc="BANCO ITAU"),
        ],
    })
    result = coverage(source, after, "eletro_forte", {"modo_consolidado": True})
    assert result["total"] == 3
    assert result["ja_classificadas"] == 1
    assert result["novas_classificadas"] == 1
    assert result["pendentes"] == 1
    assert result["percentual"] == 66.7


def test_fornecedor_242_distingue_conta_provisoria_por_modalidade():
    sample = workbook({"Itaú": [row(debit=166, credit=508)]})
    fornecedor = coverage(sample, sample, "eletro_forte",
                          {"coluna_substituir": "debito", "valores_substituiveis": ["166", "0", ""]})
    despesa = coverage(sample, sample, "eletro_forte",
                       {"coluna_substituir": "debito", "valores_substituiveis": ["0", ""]})
    assert fornecedor["percentual"] == 0
    assert despesa["percentual"] == 100


def test_1408_usa_apenas_placeholders_zero_e_preserva_conta_166():
    source = workbook({"Itaú 512": [
        row(debit=166, credit=512, desc="BANCO ITAU"),
        row(debit=0, credit=512, history="PAGO ZERO", desc="BANCO ITAU"),
    ]})
    result = workbook({"Itaú 512": [
        row(debit=166, credit=512, desc="BANCO ITAU"),
        row(debit=25, credit=512, history="PAGO ZERO", desc="BANCO ITAU"),
    ]})
    output = coverage(source, result, "eletro_forte_filial_1408",
                      {"modo_consolidado": True})
    assert output["total"] == 2
    assert output["ja_classificadas"] == 1
    assert output["novas_classificadas"] == 1
    assert output["percentual"] == 100


def test_sem_linhas_nao_anuncia_100_porcento():
    source = workbook({"Itaú": [row(history="")]})
    result = coverage(source, source, "nova_geracao_matriz")
    assert result["total"] == 0
    assert result["percentual"] is None
    assert result["classificadas"] == result["pendentes"] == 0


def test_integracao_classify_retorna_cobertura_sem_alterar_regra_existente(monkeypatch):
    monkeypatch.setattr(classification_service, "records", lambda code: [])
    source = workbook({"Banco do Brasil": [
        row(debit=187, credit=8, history="PAGO JA REVISADO"),
        row(debit="", credit=8, history="PAGO SEM CONTRAPARTIDA"),
    ]})
    result, summary = classification_service.classify(266, source, "modelo.xlsx", {})
    assert summary["cobertura"]["total"] == 2
    assert summary["cobertura"]["classificadas"] == 1
    assert summary["cobertura"]["pendentes"] == 1
    assert summary["cobertura"]["ja_classificadas"] == 1
    assert summary["cobertura"]["novas_classificadas"] == 0
    book = load_workbook(io.BytesIO(result), read_only=True)
    assert book["Banco do Brasil"]["D2"].value == 187
    assert book["Banco do Brasil"]["E2"].value == 8
    book.close()
