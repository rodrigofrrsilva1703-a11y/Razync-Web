import io
import struct

import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from razync import conferencia_fiscal, dominio_ledger


def workbook(sheets):
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for name, rows in sheets.items():
            pd.DataFrame(rows).to_excel(writer, sheet_name=name, header=False, index=False)
    return out.getvalue()


def reports():
    fiscal = workbook({"Fiscal": [
        ["ENTRADAS"], ["Codigo", "Descrição", "Valor Contabil", "Conta"],
        [1152, "Mercadorias", 1234.56, 22643],
    ]})
    ledger = workbook({"Razão": [
        ["Período:", "01/08/2026 - 31/08/2026"],
        ["Conta:", 22643.0, "", "", "", "ESTOQUE"],
        ["Data", "Lote", "Histórico", "Cta.C.Part.", "Filial", "Débito", "Crédito"],
        [46237, 10, "COMPRA DE MERCADORIA CF NF 1", 22644.0, 242, 1234.56, 0],
        [46237, 11, "COMPRA DE MERCADORIA CF NF 2", 22644.0, 1408, 999, 0],
        ["Saldo anterior", "", "", "", "", 99999, 0],
    ]})
    return fiscal, ledger


def test_fiscal_242_preview_and_excel_same_values():
    fiscal, ledger = reports()
    client = TestClient(app)
    files = {"acumuladores": ("fiscal.xlsx", fiscal), "razao": ("razao.xlsx", ledger)}
    preview = client.post("/api/v1/conferencia-fiscal/242/preview", files=files)
    assert preview.status_code == 200, preview.text
    data = preview.json()
    assert data["filial_aplicada"] == "242"
    assert len(data["lancamentos"]) == 1
    assert data["lancamentos"][0]["data"] == "03/08/2026"
    assert data["lancamentos"][0]["contrapartida"] == "22644"
    assert data["contas"][0]["contabil"] == 1234.56
    assert data["contas"][0]["situacao"] == "CONFERE"
    response = client.post("/api/v1/conferencia-fiscal/242", files=files)
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content), data_only=True)
    resumo = pd.read_excel(io.BytesIO(response.content), sheet_name=book.sheetnames[0])
    assert resumo.iloc[0]["CONTÁBIL COMPATÍVEL"] == 1234.56
    assert resumo.iloc[0]["SITUAÇÃO"] == "CONFERE"
    assert client.post("/api/v1/conferencia-fiscal/1408/preview", files=files).status_code == 422


def test_fiscal_uses_shared_biff_recovery_preserving_all_sheets(monkeypatch):
    recovered = pd.ExcelFile(io.BytesIO(reports()[1]))
    original = conferencia_fiscal.pd.ExcelFile
    def open_excel(source, *args, **kwargs):
        if source.getvalue() == b"irregular":
            raise ValueError("invalid BOUNDSHEET")
        return original(source, *args, **kwargs)
    monkeypatch.setattr(conferencia_fiscal.pd, "ExcelFile", open_excel)
    calls = []
    def recover(content, *, workbook=False):
        calls.append((content, workbook))
        return recovered
    monkeypatch.setattr(dominio_ledger, "_recuperar_xls_biff_irregular", recover)
    movements, _ = conferencia_fiscal.ler_razao(b"irregular", "Domínio.xls")
    assert calls == [(b"irregular", True)]
    assert movements["DÉBITO"].tolist() == [1234.56, 999]


def test_html_xls_brazilian_values_and_all_tables():
    table = '''<table><tr><td>Conta:</td><td>22643</td></tr>
    <tr><td>Data</td><td>Lote</td><td>Histórico</td><td>Cta.C.Part.</td><td>Débito</td><td>Crédito</td></tr>
    <tr><td>03/08/2026</td><td>1</td><td>COMPRA CF NF 1</td><td>22644</td><td>1.234,56</td><td>0,00</td></tr></table>'''
    movements, _ = conferencia_fiscal.ler_razao((table + table).encode(), "razao.xls")
    assert len(movements) == 2
    assert movements["DÉBITO"].tolist() == [1234.56, 1234.56]


def test_new_sheet_does_not_inherit_previous_account():
    content = workbook({
        "Conta": [["Conta:", 22643], ["Data", "Lote", "Histórico", "Cta.C.Part.", "Débito", "Crédito"],
                  ["03/08/2026", 1, "COMPRA CF NF 1", 22644, 100, 0]],
        "Sem conta": [["Data", "Lote", "Histórico", "Cta.C.Part.", "Débito", "Crédito"],
                      ["03/08/2026", 2, "FORA DA CONTA", 1, 500, 0]],
    })
    movements, _ = conferencia_fiscal.ler_razao(content, "razao.xlsx")
    assert movements["DÉBITO"].tolist() == [100]


def test_biff_recovery_keeps_multiple_sheets_and_legacy_dataframe(monkeypatch):
    def record(code, content=b""):
        return struct.pack("<HH", code, len(content)) + content
    def bof(kind):
        return record(0x809, struct.pack("<HHHHII", 0x600, kind, 0x0DBB, 1996, 0x41, 6))
    def sheet(number):
        return (bof(0x10) + record(0x200, struct.pack("<IIHHH", 0, 1, 0, 1, 0))
                + record(0x203, struct.pack("<HHHd", 0, 0, 0, number)) + record(0xA))
    stream = bof(5)
    for name in (b"One", b"Two"):
        stream += record(0x85, struct.pack("<IBBBB", 0, 0, 0, len(name), 0) + name)
    stream += record(0xA) + sheet(123) + sheet(456)
    class Document:
        def __init__(self, *args, **kwargs): pass
        def get_named_stream(self, name): return stream
    monkeypatch.setattr("xlrd.compdoc.CompDoc", Document)
    book = dominio_ledger._recuperar_xls_biff_irregular(b"synthetic-container", workbook=True)
    assert book.sheet_names == ["One", "Two"]
    assert pd.read_excel(book, sheet_name="Two", header=None).iloc[0, 0] == 456
    first = dominio_ledger._recuperar_xls_biff_irregular(b"synthetic-container")
    assert first.iloc[0, 0] == 123
