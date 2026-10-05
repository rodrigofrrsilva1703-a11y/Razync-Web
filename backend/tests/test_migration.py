import ast
import io
import json
import hashlib
import builtins
import symtable
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from app import engine
from app import classification_service as base
from app.main import app
from app.migration_services import workflow, slug, accounts, reconcile

ROOT = Path(__file__).resolve().parents[1]
client = TestClient(app)

def xlsx(rows, title='Itaú'):
    wb = Workbook(); ws = wb.active; ws.title = title
    for row in rows:
        ws.append(row)
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()

def model(history='Pago: Empresa: ACME INDUSTRIAL', debit='', credit='508', value=-100, date='01/10/2026', title='Itaú'):
    return xlsx([['DESCRIÇÃO','DATA','VALOR','DÉBITO','CRÉDITO','HISTÓRICO'], ['BANCO ITAÚ',date,value,debit,credit,history]], title)

def workbook_signature(content):
    wb = load_workbook(io.BytesIO(content))
    return [(ws.title, [(c.coordinate, c.value, c.number_format, str(c.font), str(c.fill), str(c.border), str(c.alignment)) for row in ws for c in row],
             ws.freeze_panes, ws.auto_filter.ref, [(k,v.width) for k,v in ws.column_dimensions.items()]) for ws in wb]

def reference_engine():
    namespace = {k: v for k,v in vars(engine).items() if not k.startswith('__')}
    namespace['st'] = SimpleNamespace(session_state={}, warning=lambda value: None)
    manifest = json.loads((ROOT/'resources/engine_manifest.json').read_text())
    for filename in ('app_legacy.py','app.py'):
        source = (ROOT/'tests/reference'/f'{filename}.txt').read_text(encoding='utf-8')
        if filename == 'app.py':
            namespace['_identificar_chave_banco_legado'] = namespace['identificar_chave_banco_empresa']
            namespace['_nome_banco_por_chave_legado'] = namespace['nome_banco_por_chave']
            namespace['_processar_extrato_conferencia_legado'] = namespace['processar_extrato_conferencia_empresa']
        names = {item['name'] for item in manifest['functions'] if item['source'] == filename}
        definitions = []
        for node in ast.parse(source).body:
            if isinstance(node, ast.FunctionDef) and node.name in names:
                node.decorator_list = []
                definitions.append(node)
        exec(compile(ast.Module(body=definitions, type_ignores=[]), filename, 'exec'), namespace)
    return SimpleNamespace(**namespace)

def test_extracted_rules_match_reference_ast_and_hashes():
    manifest = json.loads((ROOT/'resources/engine_manifest.json').read_text())
    migrated_source = (ROOT/'app/engine.py').read_text(encoding='utf-8')
    migrated_functions = [n for n in ast.parse(migrated_source).body if isinstance(n, ast.FunctionDef)]
    for filename in ('app_legacy.py','app.py'):
        source = (ROOT/'tests/reference'/f'{filename}.txt').read_text(encoding='utf-8')
        assert hashlib.sha256(source.encode()).hexdigest() == manifest['sources'][filename]
        originals = {n.name:n for n in ast.parse(source).body if isinstance(n,ast.FunctionDef)}
        for item in (r for r in manifest['functions'] if r['source'] == filename):
            node = originals[item['name']]
            text = '\n'.join(source.splitlines()[node.lineno-1:node.end_lineno])
            assert hashlib.sha256(text.encode()).hexdigest() == item['source_sha256']
            expected = text.replace('st.session_state','request_state.get()').replace('st.warning(', 'record_warning(')
            expected = expected.replace("'Modelo dominio.xlsx'", 'str(TEMPLATE)').replace('"Modelo dominio.xlsx"','str(TEMPLATE)')
            expected_ast = ast.dump(ast.parse(expected).body[0], include_attributes=False)
            assert any(ast.dump(n,include_attributes=False) == expected_ast for n in migrated_functions), item['name']

def test_extracted_engine_has_no_unresolved_globals():
    table = symtable.symtable((ROOT/'app/engine.py').read_text(encoding='utf-8'), 'engine.py', 'exec')
    def scan(scope):
        for symbol in scope.get_symbols():
            if symbol.is_global() and symbol.is_referenced():
                assert symbol.get_name() in vars(engine) or hasattr(builtins,symbol.get_name()), (scope.get_name(),symbol.get_name())
        for child in scope.get_children():
            scan(child)
    for child in table.get_children():
        scan(child)

@pytest.mark.parametrize('code',[3,178,343])
def test_autokraft_workbook_matches_original(code):
    data = xlsx([['','','','','',''],['','','01/10/2026','','',''],['ITAU','','Cliente',100,'Fornecedor',40],['','','TOTAL DE CREDITOS',100,'TOTAL DE DEBITOS',40]], '01-10')
    reference = reference_engine()
    groups, _ = reference.processar_mapa_autokraft(data,'MAPA_2026.xlsx')
    expected = reference.gerar_excel_nova_geracao(groups)
    actual, _ = workflow(code, {'mapa':[('MAPA_2026.xlsx',data)]}, {})
    assert workbook_signature(actual) == workbook_signature(expected)

@pytest.mark.parametrize('code',[1000,1001])
def test_accede_detailed_groups_match_original(code):
    data = xlsx([['Data','D/C','Complemento','Conf','Entrada','Saida'],['01/10/2026','D','SISPAG','',None,300],[None,'DOC1',100,'SALARIO MARIA',None,None],[None,'DOC2',200,'SALARIO JOAO',None,None]])
    reference = reference_engine()
    expected_frame = reference.processar_planilha_accede_sig(data,'itau',slug(code))
    assert len(expected_frame) == 2
    assert expected_frame['VALOR'].sum() == -300
    expected = reference.gerar_excel_nova_geracao({'Itau':{'principal':expected_frame,'retirados':pd.DataFrame()}},engine.TEMPLATE.read_bytes())
    actual,_ = workflow(code, {'itau':[('SIG.xlsx',data)]},{})
    assert workbook_signature(actual) == workbook_signature(expected)

def test_classification_preserves_full_original_rules_and_backup():
    history = 'Pago: Empresa: ACME INDUSTRIAL'
    for month in (7,8,9):
        content = model(history,'166','508',date=f'01/{month:02d}/2026')
        assert base.learn(3,content,f'REVISADO_{month}.xlsx') == 1
        assert base.learn(3,content,f'REVISADO_{month}.xlsx') == 0
    source = model(history)
    actual, summary = base.classify(3,source,'modelo.xlsx')
    expected, expected_summary = reference_engine().classificar_planilha_final(source,'modelo.xlsx',base.records(3),accounts(3),empresa_classificacao=slug(3))
    assert summary == expected_summary
    assert summary['automaticos'] == 1
    assert workbook_signature(actual) == workbook_signature(expected)
    assert base.status(178)['patterns'] == 0

def test_name_based_classification_in_two_periods_and_conflict():
    for month in (7,8):
        base.learn(3,model('Pago: Empresa: ACME INDUSTRIAL OBS: NF','166',date=f'01/{month:02d}/2026'),f'M{month}.xlsx')
    result, summary = base.classify(3,model('Pago: Empresa: ACME INDUSTRIAL OBS: DOCUMENTO'),'modelo.xlsx')
    assert summary['por_nome_empresa'] == 1
    base.learn(3,model('Pago: Empresa: ACME INDUSTRIAL OBS: NF','200',date='01/09/2026'),'conflict.xlsx')
    result, summary = base.classify(3,model('Pago: Empresa: ACME INDUSTRIAL'),'modelo.xlsx')
    assert summary['automaticos'] == 0

def test_eletro_principal_is_preserved_and_placeholder_is_replaced():
    wb = load_workbook(io.BytesIO(model('Pago: Fornecedor','166','508',title='Despesa - Itau 508')))
    principal = wb.copy_worksheet(wb.active); principal.title='Principal'
    source = io.BytesIO(); wb.save(source)
    actual, summary = base.classify(242,source.getvalue(),'modelo.xlsx')
    original = load_workbook(io.BytesIO(source.getvalue()))
    result = load_workbook(io.BytesIO(actual))
    assert list(result['Principal'].values) == list(original['Principal'].values)
    assert result['Despesa - Itau 508']['D2'].value is None
    assert result['Despesa - Itau 508']['E2'].value == '508'

def test_api_health_catalog_cors_and_safe_error():
    assert client.get('/health').status_code == 200
    companies = client.get('/api/v1/companies').json()
    assert len(companies) == 48
    assert next(c for c in companies if str(c['codigo']) == '626')['capabilities']['banks']['sicredi'] == '1155'
    response = client.get('/health',headers={'Origin':'https://rodrigofrrsilva1703-a11y.github.io'})
    assert 'X-Razync-Summary' in response.headers['access-control-expose-headers']
    response = client.post('/api/v1/workflow/242',data={'roles_json':'["despesas"]'},files=[('files',('invalid.xlsx',b'invalid'))])
    assert response.status_code == 422
    assert 'detail' in response.json()

def test_review_does_not_accept_client_coordinates():
    source = model()
    pending = base.pending(3,source)
    assert len(pending) == 1
    pending[0]['_col_destino'] = 6
    pending[0]['Conta da contrapartida'] = '166'
    book, summary = base.review(3,source,'modelo.xlsx',pending,True)
    wb = load_workbook(io.BytesIO(book))
    assert wb.active['D2'].value == 166
    assert wb.active['F2'].value == 'Pago: Empresa: ACME INDUSTRIAL'
    assert summary['aplicadas'] == 1

def test_connector_package_includes_installer_and_automation():
    response = client.get('/api/v1/connector/download')
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert {'connector.py','automate_ecac.ps1','chrome_extension/manifest.json','INSTALAR_RAZYNC_WEB.bat'} <= set(archive.namelist())

def test_sensitive_operations_require_configured_auth(monkeypatch):
    monkeypatch.delenv('RAZYNC_ACCESS_TOKEN',raising=False)
    monkeypatch.delenv('CLASSIFICATION_ADMIN_PASSWORD',raising=False)
    assert client.get('/api/v1/certificates/242').status_code == 503
    monkeypatch.setenv('RAZYNC_ACCESS_TOKEN','test-access')
    assert client.get('/api/v1/certificates/242').status_code == 401
    assert client.get('/api/v1/certificates/242',headers={'Authorization':'Bearer test-access'}).json() == {'certificate':None}
