import io
from datetime import datetime
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from fastapi.testclient import TestClient
from app.main import app
from app.model_editor import inspect, apply
from app.conferencia import ler_modelo_excel


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = 'BTG'
    sheet.append(['DESCRIÇÃO', 'DATA', 'VALOR', 'DÉBITO', 'CRÉDITO', 'HISTÓRICO'])
    sheet.append(['BANCO BTG', datetime(2026, 7, 9), 100, '510', '', 'Entrada'])
    sheet.append(['BANCO BTG', datetime(2026, 7, 10), -40, '', '510', 'Saída'])
    sheet['C2'].number_format = '#,##0.00'
    sheet['C2'].fill = PatternFill('solid', fgColor='ABCDEF')
    book.create_sheet('Auxiliar').append(['Preservar', 123])
    out = io.BytesIO(); book.save(out)
    return out.getvalue()


def test_open_lists_all_rows_and_only_accounting_sheets():
    result = inspect(workbook())['sheets']
    assert len(result) == 1
    assert result[0]['rows'][0][1] == '09/07/2026'
    assert len(result[0]['rows']) == 2


def test_edits_insert_delete_keep_styles_and_change_reconciliation_input():
    operations = [
        {'kind':'set','sheet':'BTG','row':2,'col':3,'value':'1.234,56'},
        {'kind':'delete','sheet':'BTG','row':3},
        {'kind':'insert','sheet':'BTG','row':2},
        {'kind':'set','sheet':'BTG','row':2,'col':1,'value':'BANCO BTG'},
        {'kind':'set','sheet':'BTG','row':2,'col':2,'value':'08/07/2026'},
        {'kind':'set','sheet':'BTG','row':2,'col':3,'value':'-50,25'},
        {'kind':'set','sheet':'BTG','row':2,'col':6,'value':'Inserida'},
    ]
    output = apply(workbook(), operations)
    book = load_workbook(io.BytesIO(output))
    assert book['BTG']['C3'].value == 1234.56
    assert book['BTG']['C3'].fill.fgColor.rgb == '00ABCDEF'
    assert book['Auxiliar']['B1'].value == 123
    frame = ler_modelo_excel(output)
    assert list(frame['VALOR']) == [-50.25, 1234.56]
    assert list(frame['HISTÓRICO']) == ['Inserida','Entrada']


@pytest.mark.parametrize('op', [
    {'kind':'delete','sheet':'BTG','row':1},
    {'kind':'set','sheet':'BTG','row':2,'col':7,'value':'x'},
    {'kind':'set','sheet':'BTG','row':2,'col':2,'value':'31/02/2026'},
    {'kind':'set','sheet':'BTG','row':2,'col':3,'value':'nan'},
    {'kind':'delete','sheet':'Auxiliar','row':2},
    {'kind':'delete','sheet':'BTG','row':999},
])
def test_rejects_invalid_edits_without_changing_original(op):
    source = workbook()
    with pytest.raises(ValueError):
        apply(source, [op])
    assert inspect(source)['sheets'][0]['rows'][0][2] == 100


def test_text_is_literal_not_formula():
    out = apply(workbook(), [{'kind':'set','sheet':'BTG','row':2,'col':6,'value':'=1+1'}])
    cell = load_workbook(io.BytesIO(out))['BTG']['F2']
    assert cell.value == '=1+1' and cell.data_type == 's'


def test_api_open_and_apply():
    import json
    client = TestClient(app)
    response = client.post('/api/v1/model-editor/open',files={'file':('modelo.xlsx',workbook())})
    assert response.status_code == 200
    response = client.post('/api/v1/model-editor/apply',files={'file':('modelo.xlsx',workbook())},data={'operations_json':json.dumps([{'kind':'delete','sheet':'BTG','row':3}])})
    assert response.status_code == 200
    assert len(ler_modelo_excel(response.content)) == 1


def test_formula_workbook_is_explicitly_rejected():
    book = load_workbook(io.BytesIO(workbook()))
    book['BTG']['C2'] = '=10+20'
    out = io.BytesIO(); book.save(out)
    with pytest.raises(ValueError, match='fórmulas'):
        inspect(out.getvalue())
