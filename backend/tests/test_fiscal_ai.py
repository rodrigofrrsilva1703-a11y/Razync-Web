import json
import urllib.error
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from app import fiscal_ai
from app.main import app


@pytest.fixture(autouse=True)
def _openrouter_disabled_by_default(monkeypatch):
    """Os testes legados usam Gemini; OpenRouter é ativado explicitamente nos novos casos."""
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_FREE_TIER_CONFIRMED", raising=False)
    # Exercita integrações anteriores apenas em testes; produção usa OpenRouter.
    monkeypatch.setenv("RAZYNC_AI_LEGACY_GEMINI", "1")


def report():
    return {"empresa":"CONFIDENTIAL", "lancamentos":[{"historico":"JOAO CPF 12345678900", "valor":9123.45}],
        "contas":[{"conta":"22643", "tipo":"ENTRADAS", "descricao":"EMPRESA PRIVADA",
        "fiscal":9123.45,"contabil":0,"diferenca":-9123.45,"extras":2,"situacao":"REVISAR"}]}


def test_external_payload_has_no_identifiers_free_text_or_values():
    groups, mapping = fiscal_ai.sanitized(report())
    assert groups == [{"grupo":"G1", "situacao":"REVISAR", "natureza":"debito", "diferenca":"falta", "adicionais":True}]
    assert mapping == {"G1":{"conta":"22643","tipo":"ENTRADAS"}}
    for secret in ("22643", "JOAO", "12345678900", "9123", "EMPRESA", "CONFIDENTIAL"):
        assert secret not in json.dumps(groups)


def gateway(monkeypatch, answer):
    monkeypatch.setenv("GEMINI_API_KEY", "test-secret")
    monkeypatch.setattr(fiscal_ai, "last_request", 0)
    class Response:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def read(self, limit):
            return json.dumps({"candidates":[{"finishReason":"STOP","content":{"parts":[{"text":json.dumps(answer)}]}}]}).encode()
    def request(req, timeout):
        assert "test-secret" not in req.full_url
        assert timeout == 45
        payload = req.data.decode()
        assert "9123" in payload and "22643" in payload and "JOAO" not in payload
        return Response()
    monkeypatch.setattr(fiscal_ai.urllib.request, "urlopen", request)


def test_valid_response_maps_to_local_account_without_changing_numbers(monkeypatch):
    gateway(monkeypatch, {"analises":[{"grupo":"G1","explicacao":"Há falta de valor compatível.","verificar":"Verifique a integração no período."}]})
    source = report()
    result = fiscal_ai.explain(source)
    assert result["analises"][0]["conta"] == "22643"
    assert source["contas"][0]["fiscal"] == 9123.45
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(source)
    assert error.value.status_code == 429


def test_invalid_group_is_rejected(monkeypatch):
    gateway(monkeypatch, {"analises":[{"grupo":"G999","explicacao":"Erro","verificar":"Teste"}]})
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert error.value.status_code == 502


def test_missing_key_does_not_require_admin(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    client = TestClient(app)
    assert client.get("/api/v1/conferencia-fiscal/242/ia/status").json() == {"configurado":False, "provedor":None, "gratuito":True, "fallback_gemini":False}
    files = {"acumuladores":("fiscal.xlsx",b"fake"),"razao":("razao.xlsx",b"fake")}
    assert client.post("/api/v1/conferencia-fiscal/242/ia",files=files).status_code == 503


def test_analysis_endpoint_accepts_no_admin_password(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "test-secret")
    monkeypatch.setattr(fiscal_ai, "conferencia_fiscal_preview", lambda *args: report())
    monkeypatch.setattr(fiscal_ai, "explain", lambda report: {"analises":[],"aviso":"Teste"})
    files = {"acumuladores":("fiscal.xlsx",b"fake"),"razao":("razao.xlsx",b"fake")}
    response = TestClient(app).post("/api/v1/conferencia-fiscal/242/ia", files=files)
    assert response.status_code == 200
    assert response.json()["aviso"] == "Teste"


def test_matching_groups_do_not_trigger_external_call(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "test-secret")
    source = report(); source["contas"][0]["situacao"] = "CONFERE"
    monkeypatch.setattr(fiscal_ai.urllib.request, "urlopen", lambda *_: pytest.fail("Must not call Google"))
    assert fiscal_ai.explain(source)["analises"] == []


@pytest.mark.parametrize("code,reason,expected", [
    (400, "API_KEY_INVALID", "recusou a chave"),
    (403, "API_KEY_SERVICE_BLOCKED", "restrições"),
    (403, "PERMISSION_DENIED", "negou acesso"),
    (404, "NOT_FOUND", "modelo configurado"),
    (429, "RESOURCE_EXHAUSTED", "cota"),
    (400, "INVALID_ARGUMENT", "JSON estruturadas"),
    (503, "UNAVAILABLE", "HTTP 503"),
])
def test_provider_errors_are_actionable_without_leaking_secrets(monkeypatch, code, reason, expected):
    import io
    monkeypatch.setenv("GEMINI_API_KEY", "test-secret")
    monkeypatch.setattr(fiscal_ai, "last_request", 0)
    def fail(*args, **kwargs):
        body = json.dumps({"error": {"message": "test-secret private provider text", "details": [{"reason": reason}]}}).encode()
        raise urllib.error.HTTPError("https://example.invalid", code, "failure", {}, io.BytesIO(body))
    monkeypatch.setattr(fiscal_ai.urllib.request, "urlopen", fail)
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert expected in error.value.detail
    assert "test-secret" not in error.value.detail
    assert "private provider text" not in error.value.detail
    assert error.value.status_code == (429 if code == 429 else 502)


def test_unavailable_configured_model_retries_flash_lite(monkeypatch):
    import io
    answer = {"analises":[{"grupo":"G1","explicacao":"Revisar diferença.","verificar":"Verifique os lançamentos."}]}
    gateway(monkeypatch, answer)
    monkeypatch.setenv("GEMINI_MODEL", "gemini-2.5-flash")
    original = fiscal_ai.urllib.request.urlopen
    calls = []
    def request(req, timeout):
        calls.append(req.full_url)
        if len(calls) == 1:
            raise urllib.error.HTTPError(req.full_url, 404, "missing", {}, io.BytesIO(b"{}"))
        return original(req, timeout)
    monkeypatch.setattr(fiscal_ai.urllib.request, "urlopen", request)
    assert fiscal_ai.explain(report())["analises"][0]["conta"] == "22643"
    assert len(calls) == 2
    assert "gemini-3.1-flash-lite" in calls[1]


def test_detailed_context_includes_authorized_records_and_totals():
    source = report()
    source["lancamentos"] = [{"conta":"22643", "historico":"Fornecedor JOAO CPF 12345678900", "data":"03/08/2026", "debito":1000, "credito":0}]
    context, mapping, refs, coverage = fiscal_ai.detailed_context(source)
    assert context["grupos"][0]["resumo"]["diferenca"] == -9123.45
    assert refs["L1"]["historico"] == "Fornecedor JOAO CPF 12345678900"
    assert refs["L1"]["debito"] == 1000
    assert "1 lançamentos" in coverage
    assert "sem documentos fiscais individuais" in context["fonte_fiscal"]


def test_invented_evidence_is_rejected(monkeypatch):
    gateway(monkeypatch, {"analises":[{"grupo":"G1", "explicacao":"Teste", "verificar":"Conferir", "evidencias":["L999"]}]})
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert error.value.status_code == 502


def test_evidence_returns_original_record_and_full_history(monkeypatch):
    gateway(monkeypatch, {"analises":[{"grupo":"G1", "explicacao":"Revisar L1", "verificar":"Conferir documento", "evidencias":["L1"]}]})
    source = report()
    source["lancamentos"] = [{"conta":"22643", "historico":"Histórico completo do fornecedor", "debito":1000, "credito":0}]
    result = fiscal_ai.explain(source)
    assert result["analises"][0]["evidencias"][0]["historico"] == "Histórico completo do fornecedor"
    assert result["analises"][0]["evidencias"][0]["debito"] == 1000


def test_coverage_explicit_when_context_is_limited():
    source = report()
    source["contas"] = [dict(source["contas"][0], conta=str(i)) for i in range(15)]
    source["lancamentos"] = [dict(conta="0", historico="Original", debito=1) for _ in range(1501)]
    context, mapping, refs, coverage = fiscal_ai.detailed_context(source)
    assert len(mapping) == 12
    assert len(refs) == 1500
    assert context["grupos"][0]["cobertura"] == {"enviados":1500, "existentes":1501}
    assert "12 de 15" in coverage


def test_moeda_do_gemini_sempre_em_formato_brasileiro():
    assert fiscal_ai.formatar_brl(1234567.8) == "R$ 1.234.567,80"
    assert fiscal_ai.formatar_brl(-4282.5) == "-R$ 4.282,50"
    assert fiscal_ai.formatar_brl(0) == "R$ 0,00"
    fonte = "Conta 22643: R$ 1234.56; R$ 1,234.56; R$ 1.234,56. Diferença R$ -91.52."
    corrigido = fiscal_ai.padronizar_moeda(fonte)
    assert corrigido == (
        "Conta 22643: R$ 1.234,56; R$ 1.234,56; R$ 1.234,56. "
        "Diferença -R$ 91,52."
    )


def test_gemini_recebe_formatos_brl_e_devolve_valores_oficiais(monkeypatch):
    source = report()
    captured = {}
    monkeypatch.setenv("GEMINI_API_KEY", "test-secret")
    monkeypatch.setattr(fiscal_ai, "last_request", 0)
    class Resposta:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def read(self, limit):
            return json.dumps({"candidates":[{"finishReason":"STOP","content":{"parts":[{"text":json.dumps({
                "analises":[{"grupo":"G1","explicacao":"Falta R$ 9123.45.",
                             "verificar":"Verifique R$ 9,123.45.", "evidencias":[]}],
            })}]}}]}).encode()
    def fake_urlopen(req, timeout):
        body = json.loads(req.data)
        captured["payload"] = body
        assert "R$ 9.123,45" in body["contents"][0]["parts"][0]["text"]
        return Resposta()
    monkeypatch.setattr(fiscal_ai.urllib.request, "urlopen", fake_urlopen)
    result = fiscal_ai.explain(source)
    item = result["analises"][0]
    assert item["valores"]["fiscal"] == 9123.45
    assert item["valores"]["diferenca"] == -9123.45
    assert item["explicacao"] == "Falta R$ 9.123,45."
    assert item["verificar"] == "Verifique R$ 9.123,45."


def test_gemini_devolve_codigos_originais_de_acumuladores_e_conta(monkeypatch):
    source = report()
    source["contas"][0].update({
        "acumuladores": "1152, 1153",
        "detalhes_fiscais": [
            {"codigo": "1152", "descricao": "Mercadorias", "valor": 8000.00},
            {"codigo": "1153", "descricao": "Outras entradas", "valor": 1123.45},
        ],
    })
    context, mapping, *_ = fiscal_ai.detailed_context(source)
    fiscal = context["grupos"][0]["resumo"]["detalhes_fiscais"]
    assert [row["codigo"] for row in fiscal] == ["1152", "1153"]
    assert context["grupos"][0]["resumo"]["conta"] == "22643"
    gateway(monkeypatch, {"analises": [{
        "grupo": "G1", "explicacao": "Revisar os acumuladores 1152 e 1153.",
        "verificar": "Conferir a conta 22643.", "evidencias": [],
    }]})
    resposta = fiscal_ai.explain(source)["analises"][0]
    assert resposta["conta"] == "22643"
    assert resposta["acumuladores"] == "1152, 1153"
    assert resposta["detalhes_fiscais"] == fiscal


def test_gemini_universal_status_e_filial_independente_da_empresa(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "test-secret")
    calls = []
    def preview(*args):
        calls.append(args)
        return report()
    monkeypatch.setattr(fiscal_ai, "conferencia_fiscal_preview", preview)
    monkeypatch.setattr(fiscal_ai, "explain", lambda value: {"analises": [], "aviso": "Teste universal"})
    client = TestClient(app)
    assert client.get("/api/v1/conferencia-fiscal/ia/status").json()["configurado"] is True
    response = client.post(
        "/api/v1/conferencia-fiscal/ia",
        files={
            "acumuladores": ("fiscal.xlsx", b"fiscal"),
            "razao": ("razao.xlsx", b"razao"),
        },
        data={"empresa_codigo": "987", "filial": "242"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["aviso"] == "Teste universal"
    assert len(calls) == 1
    assert calls[0][-2:] == (987, "242")



def test_gemini_exige_analise_didatica_com_limites_e_verificacoes(monkeypatch):
    """Mantém a resposta atual, mas orienta o modelo a explicar os dados em profundidade."""
    source = report()
    monkeypatch.setenv("GEMINI_API_KEY", "test-secret")
    monkeypatch.setattr(fiscal_ai, "last_request", 0)

    class Resposta:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def read(self, limit):
            return json.dumps({"candidates": [{
                "finishReason": "STOP",
                "content": {"parts": [{"text": json.dumps({"analises": [{
                    "grupo": "G1", "explicacao": "Os valores precisam ser conferidos.",
                    "verificar": "1. Conferir a conta.", "evidencias": []
                }]})}]}
            }]}).encode()

    def confirmar_instrucoes(req, timeout):
        dados = json.loads(req.data)
        prompt = dados["systemInstruction"]["parts"][0]["text"]
        assert "3 a 4 parágrafos" in prompt
        assert "3 a 5 etapas objetivas e numeradas" in prompt
        assert "códigos e descrições dos acumuladores" in prompt
        assert "não invente citação" in prompt
        assert "Diferença" in prompt or "DIFERENÇA" in prompt
        assert "valores_brl" in prompt
        assert "explicacao" in str(dados["generationConfig"]["responseSchema"])
        assert timeout == 45
        return Resposta()

    monkeypatch.setattr(fiscal_ai.urllib.request, "urlopen", confirmar_instrucoes)
    analise = fiscal_ai.explain(source)["analises"][0]
    assert analise["conta"] == "22643"
    assert analise["verificar"] == "1. Conferir a conta."
    assert source["contas"][0]["fiscal"] == 9123.45



def test_exportar_analise_gemini_em_excel_sem_nova_chamada(monkeypatch):
    """Exporta o parecer já recebido, preservando moedas, acumuladores e evidências."""
    import io
    from openpyxl import load_workbook

    monkeypatch.setattr(
        fiscal_ai, "explain",
        lambda *_: pytest.fail("Exportação não pode chamar Gemini novamente"),
    )
    payload = {
        "empresa_nome": "EMPRESA EXEMPLO LTDA",
        "analises": [{
            "conta": "22643", "tipo": "ENTRADAS", "descricao": "Compras",
            "situacao": "REVISAR", "acumuladores": "1152",
            "detalhes_fiscais": [{"codigo": "1152", "descricao": "Compras para revenda",
                                  "valor": 1234.56}],
            "valores": {"fiscal": 1234.56, "contabil": 1200,
                        "total_conta": 1400, "diferenca": -34.56},
            "explicacao": "A diferença é R$ 34,56; confirme o histórico.",
            "verificar": "1. Conferir a NF.\n2. Validar acumulador.",
            "evidencias": [{
                "referencia": "L1", "data": "09/10/2026",
                "historico": "=HYPERLINK(\"http://invalid\", \"teste\")",
                "conta": "22643", "contrapartida": "508",
                "debito": 1200, "credito": 0
            }]
        }]
    }
    response = TestClient(app).post("/api/v1/conferencia-fiscal/ia/exportar", json=payload)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.content[:2] == b"PK"
    workbook = load_workbook(io.BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ["Análises IA", "Acumuladores", "Lançamentos citados"]
    main = workbook["Análises IA"]
    assert main["A2"].value == "Empresa: EMPRESA EXEMPLO LTDA"
    assert main["A6"].value == "22643"
    assert main["F6"].value == 1234.56
    assert main["G6"].value == 1200
    assert main["I6"].value == -34.56
    assert "confirme o histórico" in main["J6"].value
    assert "Conferir a NF" in main["K6"].value
    acc = workbook["Acumuladores"]
    assert acc["C2"].value == "1152"
    assert acc["E2"].value == 1234.56
    evidence = workbook["Lançamentos citados"]
    assert evidence["B2"].value == "L1"
    assert evidence["D2"].data_type != "f", "Históricos externos nunca viram fórmulas"
    assert "HYPERLINK" in evidence["D2"].value


def test_exportar_analise_gemini_rejeita_conteudo_vazio_ou_excessivo():
    client = TestClient(app)
    for payload in [
        {"analises":[]},
        {"analises":[{"conta":"1"}] * 16},
        {"analises":[{"conta":"1", "explicacao":"X" * 6001}]},
    ]:
        response = client.post("/api/v1/conferencia-fiscal/ia/exportar", json=payload)
        assert response.status_code == 422, response.text


def test_openrouter_usa_prompt_integral_fallback_e_contas_originais(monkeypatch):
    """A mesma conferência completa segue para o próximo modelo sem perder evidências."""
    monkeypatch.setenv("OPENROUTER_API_KEY", "openrouter-test-secret")
    monkeypatch.setenv("OPENROUTER_MODELS",
                       "openrouter/free,nvidia/nemotron-3-ultra-550b-a55b:free,nvidia/nemotron-3.5-lightning:free")
    monkeypatch.setenv("GEMINI_API_KEY", "legacy-gemini-test")
    monkeypatch.setattr(fiscal_ai, "model_rotation_index", 0)
    context = report()
    context["lancamentos"] = [{
        "conta":"22643", "historico":"Compra de mercadoria CF NF 100",
        "data":"03/08/2026", "debito":9123.45, "credito":0,
    }]
    calls = []
    class Response:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def read(self, n):
            return json.dumps({"model": "nvidia/nemotron-3-ultra-550b-a55b:free", "choices":[{
                "finish_reason": "stop",
                "message": {"content":json.dumps({"analises":[{
                    "grupo":"G1","explicacao":"Revisar L1 e acumulador.",
                    "verificar":"1. Conferir o lançamento L1.","evidencias":["L1"],
                }]})}
            }]}).encode()
    def urlopen(req, timeout):
        assert timeout == 90
        assert req.full_url == "https://openrouter.ai/api/v1/chat/completions"
        assert req.get_header("Authorization") == "Bearer openrouter-test-secret"
        body = json.loads(req.data)
        calls.append(body)
        assert body["models"] == [
            "openrouter/free", "nvidia/nemotron-3-ultra-550b-a55b:free", "nvidia/nemotron-3.5-lightning:free"
        ] if len(calls) == 1 else [
            "nvidia/nemotron-3-ultra-550b-a55b:free", "nvidia/nemotron-3.5-lightning:free", "openrouter/free"
        ]
        assert body["provider"]["require_parameters"] is True
        assert body["provider"]["data_collection"] == "deny"
        assert body["response_format"]["type"] == "json_schema"
        assert body["response_format"]["json_schema"]["strict"] is True
        assert "valores_brl" in body["messages"][0]["content"]
        assert "9123.45" in body["messages"][1]["content"]
        assert "22643" in body["messages"][1]["content"]
        assert "Compra de mercadoria" in body["messages"][1]["content"]
        assert "openrouter-test-secret" not in str(body)
        return Response()
    monkeypatch.setattr(fiscal_ai.urllib.request,"urlopen",urlopen)
    results = [fiscal_ai.explain(context),fiscal_ai.explain(context)]
    assert len(calls) == 2
    for result in results:
        assert result["provedor"] == "openrouter"
        assert result["modelo_usado"] == "nvidia/nemotron-3-ultra-550b-a55b:free"
        assert result["analises"][0]["conta"] == "22643"
        assert result["analises"][0]["valores"]["fiscal"] == 9123.45
        assert result["analises"][0]["evidencias"][0]["historico"] == "Compra de mercadoria CF NF 100"
        assert "L1" in result["analises"][0]["explicacao"]


def test_openrouter_status_tem_prioridade_e_gemini_continua_reserva(monkeypatch):
    client = TestClient(app)
    monkeypatch.setenv("GEMINI_API_KEY","gemini-legacy")
    assert client.get("/api/v1/conferencia-fiscal/ia/status").json() == {
        "configurado": True, "provedor": "gemini", "gratuito": False, "fallback_gemini":False
    }
    monkeypatch.setenv("OPENROUTER_API_KEY","or-key")
    assert client.get("/api/v1/conferencia-fiscal/ia/status").json() == {
        "configurado": True, "provedor": "openrouter", "gratuito": True, "fallback_gemini":False
    }


@pytest.mark.parametrize("code,needle", [
    (401,"recusou a chave"),
    (402,"créditos do OpenRouter"),
    (429,"limites de requisições"),
    (400,"formato de resposta"),
])
def test_openrouter_falhas_seguras_sem_expor_dados(monkeypatch,code,needle):
    import io
    monkeypatch.setenv("OPENROUTER_API_KEY","openrouter-test-secret")
    monkeypatch.setenv("OPENROUTER_MODELS","openrouter/free,nvidia/nemotron-3-ultra-550b-a55b:free")
    def fail(req,timeout):
        raise urllib.error.HTTPError(
            req.full_url,code,"Error",{},
            io.BytesIO(b'{"error":{"message":"openrouter-test-secret e dados sigilosos"}}')
        )
    monkeypatch.setattr(fiscal_ai.urllib.request,"urlopen",fail)
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert needle in error.value.detail
    assert "openrouter-test-secret" not in error.value.detail
    assert "sigilosos" not in error.value.detail


def test_openrouter_nao_aceita_evidencia_inventada(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY","or-key")
    monkeypatch.setattr(fiscal_ai,"_openrouter_completion",lambda *_: (
        {"analises":[{"grupo":"G1","explicacao":"Teste","verificar":"Conferir","evidencias":["L999"]}]},
        "openrouter/free",
    ))
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert error.value.status_code == 502


def test_openrouter_rejeita_lista_de_modelos_invalida(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY","or-key")
    monkeypatch.setenv("OPENROUTER_MODELS","openai/gpt-4.1-mini")
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert error.value.status_code == 503
    assert "OPENROUTER_MODELS" in error.value.detail


def test_openrouter_modo_gratuito_por_padrao_e_precos_zerados(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "free-key")
    monkeypatch.delenv("OPENROUTER_MODELS", raising=False)
    assert fiscal_ai._openrouter_models() == ["openrouter/free"]
    monkeypatch.setattr(fiscal_ai, "model_rotation_index", 0)
    seen = {}
    class Response:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def read(self, limit):
            return json.dumps({"model": "openrouter/free",
                "choices": [{"finish_reason": "stop", "message": {
                    "content":json.dumps({"analises":[{
                        "grupo":"G1","explicacao":"Revisar possível diferença.",
                        "verificar":"1. Conferir histórico.","evidencias":[]
                    }]})}}]}).encode()
    def check(req, timeout):
        assert req.full_url == "https://openrouter.ai/api/v1/chat/completions"
        body = json.loads(req.data)
        assert body["model"] == "openrouter/free"
        assert "models" not in body
        assert body["provider"]["max_price"] == {"prompt":0,"completion":0}
        assert body["provider"]["data_collection"] == "deny"
        assert body["provider"]["require_parameters"] is True
        seen["called"] = True
        return Response()
    monkeypatch.setattr(fiscal_ai.urllib.request, "urlopen", check)
    result = fiscal_ai.explain(report())
    assert seen["called"]
    assert result["gratuito"] is True
    assert result["analises"][0]["valores"]["fiscal"] == 9123.45


@pytest.mark.parametrize("model", [
    "openai/gpt-4.1-mini",
    "google/gemini-2.5-flash",
    "anthropic/claude-haiku-4.5",
    "openrouter/auto",
    "openrouter/free,openai/gpt-4.1-mini",
    "openrouter/free,",
])
def test_modo_gratuito_rejeita_modelos_pagos_ou_nao_explicitos(monkeypatch, model):
    monkeypatch.setenv("OPENROUTER_API_KEY", "free-key")
    monkeypatch.setenv("OPENROUTER_MODELS", model)
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert error.value.status_code == 503
    assert "Modelos pagos estão bloqueados" in error.value.detail


def test_sem_chave_openrouter_nao_ativa_gemini_por_padrao(monkeypatch):
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.delenv("RAZYNC_AI_LEGACY_GEMINI", raising=False)
    monkeypatch.setenv("GEMINI_API_KEY", "gemini-key")
    response = TestClient(app).get("/api/v1/conferencia-fiscal/ia/status")
    assert response.json() == {"configurado":False, "provedor":None, "gratuito":True, "fallback_gemini":False}
    files={"acumuladores":("fiscal.xlsx",b"x"),"razao":("razao.xlsx",b"y")}
    response = TestClient(app).post("/api/v1/conferencia-fiscal/ia",files=files)
    assert response.status_code == 503


def test_modo_gratuito_nunca_envia_modelo_pago(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "free-key")
    monkeypatch.setenv("OPENROUTER_MODELS", "openrouter/free,nvidia/nemotron-3-ultra-550b-a55b:free")
    models=fiscal_ai._openrouter_models()
    assert all(m=="openrouter/free" or m.endswith(":free") for m in models)


def test_quota_openrouter_troca_para_gemini_preservando_contexto(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-or")
    monkeypatch.setenv("GEMINI_API_KEY", "test-gemini")
    monkeypatch.setenv("GEMINI_FREE_TIER_CONFIRMED", "1")
    payloads = []
    def openrouter(payload, models, key):
        payloads.append(("or", payload))
        raise HTTPException(429, "Limite atingido")
    def gemini(payload, key, model, *, legacy):
        payloads.append(("gemini", payload))
        return {"analises":[{"grupo":"G1", "explicacao":"Análise consistente.",
                "verificar":"1. Conferir relatório.", "evidencias":[]}]}, model
    monkeypatch.setattr(fiscal_ai, "_openrouter_completion", openrouter)
    monkeypatch.setattr(fiscal_ai, "_gemini_completion", gemini)
    result = fiscal_ai.explain(report())
    assert len(payloads) == 2
    assert payloads[0][1] is payloads[1][1]
    assert result["provedor"] == "gemini"
    assert result["fallback_usado"] is True
    assert result["gratuito"] is True
    assert result["analises"][0]["conta"] == "22643"
    assert result["analises"][0]["valores"]["fiscal"] == 9123.45


def test_fallback_gemini_nao_configurado_nao_usa_google(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-or")
    monkeypatch.setenv("GEMINI_API_KEY", "test-gemini")
    monkeypatch.delenv("GEMINI_FREE_TIER_CONFIRMED", raising=False)
    def fail(*args):
        raise HTTPException(429, "Limite OpenRouter")
    monkeypatch.setattr(fiscal_ai, "_openrouter_completion", fail)
    monkeypatch.setattr(fiscal_ai, "_gemini_completion",
                        lambda *args, **kwargs: pytest.fail("Fallback não autorizado"))
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert error.value.status_code == 429
    assert fiscal_ai.status()["fallback_gemini"] is False


def test_status_ativa_reserva_apenas_em_projeto_confirmado(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-or")
    monkeypatch.setenv("GEMINI_API_KEY", "test-gemini")
    monkeypatch.setenv("GEMINI_FREE_TIER_CONFIRMED", "1")
    assert fiscal_ai.status()["fallback_gemini"] is True
    monkeypatch.delenv("GEMINI_FREE_TIER_CONFIRMED")
    assert fiscal_ai.status()["fallback_gemini"] is False


def test_fallback_recusa_resposta_fabricada_no_gemini(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "or-test")
    monkeypatch.setenv("GEMINI_API_KEY", "gemini-test")
    monkeypatch.setenv("GEMINI_FREE_TIER_CONFIRMED", "1")
    monkeypatch.setattr(fiscal_ai, "_openrouter_completion",
        lambda *args: (_ for _ in ()).throw(HTTPException(429, "limite")))
    monkeypatch.setattr(fiscal_ai, "_gemini_completion",
        lambda *args, **kwargs: ({"analises":[{
            "grupo":"G1", "explicacao":"Inconsistente", "verificar":"Confira",
            "evidencias":["L999"]
        }]}, "gemini-3.1-flash-lite"))
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert error.value.status_code == 429


def test_modelos_gemini_pagantes_nao_podem_ser_reserva(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "or-test")
    monkeypatch.setenv("GEMINI_API_KEY", "gemini-test")
    monkeypatch.setenv("GEMINI_FREE_TIER_CONFIRMED", "1")
    monkeypatch.setenv("GEMINI_FREE_MODEL", "gemini-3.1-pro")
    monkeypatch.setattr(fiscal_ai, "_openrouter_completion",
        lambda *args: pytest.fail("Não deve chamar antes de validar o modelo"))
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert error.value.status_code == 503


def test_duas_cotas_esgotadas_nao_alteram_relatorio(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "or-test")
    monkeypatch.setenv("GEMINI_API_KEY", "gemini-test")
    monkeypatch.setenv("GEMINI_FREE_TIER_CONFIRMED", "1")
    def limit(*args, **kwargs):
        raise HTTPException(429, "Limite temporário")
    monkeypatch.setattr(fiscal_ai, "_openrouter_completion", limit)
    monkeypatch.setattr(fiscal_ai, "_gemini_completion", limit)
    source = report()
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(source)
    assert error.value.status_code == 429
    assert "OpenRouter e Gemini" in error.value.detail
    assert source["contas"][0]["fiscal"] == 9123.45


def test_openrouter_formato_estrito_400_tenta_json_compativel_sem_mudar_dados(monkeypatch):
    """Corrige o erro 400 comum de modelo grátis sem remover proteção zero custo."""
    import io
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-or")
    monkeypatch.delenv("OPENROUTER_MODELS", raising=False)
    monkeypatch.delenv("GEMINI_FREE_TIER_CONFIRMED", raising=False)
    requests = []

    class Response:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def read(self, size):
            return json.dumps({"model":"nvidia/nemotron-3.5-lightning:free",
                "choices":[{"finish_reason":"stop", "message":{"content":
                    '```json\\n' + json.dumps({"analises":[{
                        "grupo":"G1", "explicacao":"Conferir o acumulador fiscal.",
                        "verificar":"1. Comparar documento.", "evidencias":[]
                    }]}) + '\\n```'
                }}]}).encode()

    def mock(req, timeout):
        body = json.loads(req.data)
        requests.append(body)
        assert req.get_header("Authorization") == "Bearer test-or"
        assert timeout == 90
        if len(requests) == 1:
            raise urllib.error.HTTPError(
                req.full_url, 400, "Invalid parameters", {},
                io.BytesIO(b'{"error":{"message":"unsupported schema"}}')
            )
        return Response()

    monkeypatch.setattr(fiscal_ai.urllib.request, "urlopen", mock)
    result=fiscal_ai.explain(report())
    assert len(requests) == 2
    first, second = requests
    assert first["model"] == "openrouter/free" == second["model"]
    assert first["response_format"]["type"] == "json_schema"
    assert "response_format" not in second
    assert first["provider"]["require_parameters"] is True
    assert "require_parameters" not in second["provider"]
    assert first["provider"]["data_collection"] == second["provider"]["data_collection"] == "deny"
    assert first["provider"]["max_price"] == second["provider"]["max_price"] == {"prompt":0,"completion":0}
    assert first["messages"][1]["content"] == second["messages"][1]["content"]
    assert first["messages"][0]["content"] in second["messages"][0]["content"]
    assert result["provedor"] == "openrouter"
    assert result["modelo_usado"] == "nvidia/nemotron-3.5-lightning:free"
    assert result["analises"][0]["conta"] == "22643"
    assert result["analises"][0]["valores"]["fiscal"] == 9123.45


def test_openrouter_auth_403_nao_ignora_erro_nem_faz_retry(monkeypatch):
    import io
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-or")
    counter=[]
    def fail(req, timeout):
        counter.append(True)
        raise urllib.error.HTTPError(req.full_url,403,"Forbidden",{},
             io.BytesIO(b'{"error":{"message":"bad credentials secret"}}'))
    monkeypatch.setattr(fiscal_ai.urllib.request,"urlopen", fail)
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert len(counter) == 1
    assert error.value.status_code == 403
    assert "bad credentials secret" not in error.value.detail


def test_openrouter_doispedidos_400_continua_gratuito_erro_legivel(monkeypatch):
    import io
    monkeypatch.setenv("OPENROUTER_API_KEY","test-or")
    monkeypatch.delenv("GEMINI_FREE_TIER_CONFIRMED",raising=False)
    seen=[]
    def fail(req, timeout):
        data=json.loads(req.data)
        seen.append(data)
        raise urllib.error.HTTPError(req.full_url,400,"Unsupported",{},
          io.BytesIO(b'{"error":{"message":"histórico de cliente sensível"}}'))
    monkeypatch.setattr(fiscal_ai.urllib.request,"urlopen",fail)
    with pytest.raises(HTTPException) as error:
        fiscal_ai.explain(report())
    assert len(seen)==2
    assert error.value.status_code==400
    assert "histórico de cliente sensível" not in error.value.detail
    assert "gratuito" in error.value.detail
    assert all(p["provider"]["max_price"]=={"prompt":0,"completion":0} for p in seen)
