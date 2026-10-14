"""Regressão Accede 1001: PDF real do aplicativo Santander + bancos SIG originais.

Usa PDF sintético, sem armazenar extratos confidenciais no repositório.
"""
import io
import json

import fitz
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import operating_hours
from app.main import app
from app.migration_services import workflow
from app.registry import CAPABILITIES
from razync.companies import CONFIGURACOES_ACCEDE
from test_migration import xlsx


def sample_sig(bank):
    return xlsx([
        ["Data", "D/C", "Complemento", "Conf", "Entrada", "Saída", "Saldo", "Conta Vinculada"],
        ["01/10/2026", "CRÉDITO", f"RECEBIMENTO {bank}", "DOC-01", 150.75, None, 150.75, None],
        ["02/10/2026", "DÉBITO", f"PAGAMENTO {bank}", "DOC-02", None, 45.20, 105.55, None],
    ], title=bank.title())


def sample_santander_pdf():
    """Saldos em ordem inversa: 920,00 -> 1000,00 -> 800,00 -> 799,95."""
    doc = fitz.open()
    first = doc.new_page()
    lines = [
        "Aplicativo Santander Empresas",
        "ACCEDE EQUIPAMENTOS INDUSTRIAIS LTDA Agencia: 4433 Conta: 11111",
        "Periodos: 01/07/2026 a 31/08/2026",
        "Data Historico Documento Valor (R$) Saldo (R$)",
        "03/08/2026 Pix Enviado Fornecedor Teste -80,00 920,00",
        "02/08/2026 Resgate Cdb/rdb 200,00 1.000,00",
    ]
    for i, line in enumerate(lines):
        first.insert_text((40, 50 + i * 25), line, fontsize=9)
    second = doc.new_page()
    lines = [
        "01/08/2026 Rendimento Liquido De Contamax 7000",
        "RENDIMENTO LIQUIDO DE CONTAMAX 0,05 800,00",
        "30/07/2026 Mensalidade De Seguro",
        "INCENDIO EMP 000111 -7,00 799,95",
        "Saldo de ContaMax 920,00",
        "Saldo Disponivel 0,00",
    ]
    for i, line in enumerate(lines):
        second.insert_text((40, 50 + i * 25), line, fontsize=9)
    result = doc.tobytes()
    doc.close()
    return result


def test_cadastro_santander_1001_somente_pdf_sem_alterar_outros():
    cap = CAPABILITIES[1001]
    assert cap["banks"] == {"itau": "508", "sicredi": "505", "santander": "513"}
    role = next(role for role in cap["roles"] if role["name"] == "santander")
    assert role["accept"] == ".pdf"
    assert role["multiple"] is True and role["optional"] is True
    assert CONFIGURACOES_ACCEDE["accede_equipamentos"]["contas_bancarias"]["santander"] == "513"
    assert "santander" not in CAPABILITIES[1000]["banks"]


def test_pdf_santander_1001_gera_aba_e_conta_513_com_datas_corretas():
    content, name = workflow(1001, {"santander": [("extrato.pdf", sample_santander_pdf())]}, {})
    assert name == "ACCEDE_1001_MODELO_DOMINIO.xlsx"
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Santander"]
    rows = list(workbook.active.values)[1:]
    assert len(rows) == 4
    assert [r[0] for r in rows] == ["BANCO SANTANDER"] * 4
    assert [r[1] for r in rows] == ["30/07/2026", "01/08/2026", "02/08/2026", "03/08/2026"]
    assert [r[2] for r in rows] == [-7, 0.05, 200, -80]
    assert [str(r[3]) if r[3] else "" for r in rows] == ["", "513", "513", ""]
    assert [str(r[4]) if r[4] else "" for r in rows] == ["513", "", "", "513"]


def test_pdf_santander_1001_filtro_agosto_exclui_julho():
    content, _ = workflow(1001, {"santander": [("extrato.pdf", sample_santander_pdf())]},
                          {"data_inicial": "2026-08-01", "data_final": "2026-08-31"})
    rows = list(load_workbook(io.BytesIO(content)).active.values)[1:]
    assert len(rows) == 3
    assert all(r[1].endswith("/08/2026") for r in rows)


def test_multibanco_com_santander_pdf_e_sig_itau_sicredi():
    roles = {bank: [(f"{bank}.xlsx", sample_sig(bank))] for bank in ("itau", "sicredi")}
    roles["santander"] = [("extrato-santander.pdf", sample_santander_pdf())]
    content, _ = workflow(1001, roles, {})
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Itau", "Sicredi", "Santander"]
    assert wb["Itau"]["A2"].value == "BANCO ITAÚ"
    assert wb["Sicredi"]["A2"].value == "SICREDI"
    assert wb["Santander"]["A2"].value == "BANCO SANTANDER"
    assert wb["Itau"]["C2"].value == wb["Sicredi"]["C2"].value == 150.75


def test_santander_1001_rejeita_sig_pdf_invalido_e_arquivo_repetido():
    with pytest.raises(ValueError, match="PDF"):
        workflow(1001, {"santander": [("sig.xlsx", sample_sig("santander"))]}, {})
    with pytest.raises(ValueError, match="PDF|layout|não foi possível"):
        workflow(1001, {"santander": [("extrato.pdf", b"%PDF-not-real")]}, {})
    pdf = sample_santander_pdf()
    with pytest.raises(ValueError, match="mais de uma vez"):
        workflow(1001, {"santander": [("a.pdf", pdf), ("b.pdf", pdf)]}, {})


def test_api_1001_aceita_extrato_pdf_e_retorna_modelo(monkeypatch):
    monkeypatch.setattr(operating_hours, "is_open", lambda *a: True)
    client = TestClient(app)
    response = client.get("/api/v1/companies/1001")
    assert response.status_code == 200
    assert next(r for r in response.json()["capabilities"]["roles"] if r["name"] == "santander")["accept"] == ".pdf"
    response = client.post(
        "/api/v1/workflow/1001",
        data={"roles_json": json.dumps(["santander"]), "options_json": "{}"},
        files={"files": ("santander.pdf", sample_santander_pdf(), "application/pdf")},
    )
    assert response.status_code == 200, response.text
    assert load_workbook(io.BytesIO(response.content)).sheetnames == ["Santander"]
