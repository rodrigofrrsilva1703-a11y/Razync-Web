import io
import pytest
from openpyxl import Workbook, load_workbook
from app import ai_classification_agent as agent, classification_service as base


def book(rows=None, principal=False):
    wb=Workbook();ws=wb.active;ws.title='Itaú 508'
    ws.append(['DESCRIÇÃO','DATA','VALOR','DÉBITO','CRÉDITO','HISTÓRICO'])
    for row in rows or [[-100,'',508,'Pago: Empresa: Fornecedor ABC']]:
        ws.append(['BANCO ITAU','15/10/2026',*row])
    ws['G2']='=1+1';ws['G2'].number_format='0.00'
    wb.create_sheet('Retirados')['A1']='Preservado'
    if principal: wb.create_sheet('Principal')['A1']='Preservado'
    out=io.BytesIO();wb.save(out);return out.getvalue()


def record(periods=('2026-07','2026-08'),account='412',bank='itau_508',nature='pago',signature=None):
    bankaccount={'itau_508':'508','bb':'8','itau_512':'512'}[bank]
    return dict(banco=bank,assinatura=signature or agent.engine.criar_assinatura_classificacao('Pago: Empresa: Fornecedor ABC'),
        debito=account if nature=='pago' else bankaccount,credito=bankaccount if nature=='pago' else account,
        periodos=list(periods),ocorrencias=20,exemplo_historico='Pago: Empresa: Fornecedor ABC')


def setup(monkeypatch,records):
    monkeypatch.setattr(base,'records',lambda company:records if int(company)==242 else [])
    agent._cache.clear()


@pytest.mark.parametrize('periods,strength',[(('2026-07',),'MODERADA'),(('2026-07','2026-08'),'FORTE'),(('2026-07','2026-08','2026-09'),'FORTE')])
def test_period_evidence(monkeypatch,periods,strength):
    setup(monkeypatch,[record(periods)])
    r=agent.investigate(242,book(),use_ai=False)
    assert r['sugestoes'][0]['situacao']==strength
    assert r['sugestoes'][0]['conta_sugerida']=='412'
    assert not r['sugestoes'][0]['aprovada']


def test_conflicts_require_manual(monkeypatch):
    setup(monkeypatch,[record(),record(account='413')])
    r=agent.investigate(242,book(),use_ai=False)
    assert r['sugestoes'][0]['conta_sugerida']==''
    assert r['sugestoes'][0]['situacao']=='MANUAL'


def test_crossbank_and_other_company(monkeypatch):
    setup(monkeypatch,[record(bank='bb')])
    assert agent.investigate(242,book(),use_ai=False)['sugestoes'][0]['situacao']=='MODERADA'
    assert agent.investigate(1408,book([[-100,'',512,'Pago: Empresa: Fornecedor ABC']]),use_ai=False)['sugestoes'][0]['conta_sugerida']==''


def test_nature_and_similar_names(monkeypatch):
    setup(monkeypatch,[record(nature='recebido')])
    assert agent.investigate(242,book(),use_ai=False)['sugestoes'][0]['conta_sugerida']==''
    setup(monkeypatch,[record(signature='pago|fornecedor abc outra entidade')])
    assert agent.investigate(242,book(),use_ai=False)['sugestoes'][0]['conta_sugerida']==''


def test_groups_and_cache(monkeypatch):
    setup(monkeypatch,[record()])
    data=book([[-100,'',508,'Pago: Empresa: Fornecedor ABC'],[-200,'',508,'Pago: Empresa: Fornecedor ABC'],[-100,'',508,'Pago: Empresa: Fornecedor ABC Obs: Devolução']])
    r=agent.investigate(242,data,use_ai=False)
    assert r['total_pendencias']==3 and r['grupos_investigados']==2
    assert agent.investigate(242,data,use_ai=False)['cache']


def test_profile_preserves_principal_and_completed(monkeypatch):
    setup(monkeypatch,[record()])
    data=book([[-100,166,508,'Pago: Empresa: Fornecedor ABC'],[-100,99,508,'Pago: Empresa: Fornecedor ABC']],True)
    r=agent.investigate(242,data,{'origem_242':'fornecedor'},False)
    assert r['total_pendencias']==1
    r2=agent.investigate(242,data,{'origem_242':'despesa'},False)
    assert r2['total_pendencias']==0


def test_apply_only_approved_and_no_learning(monkeypatch):
    setup(monkeypatch,[record()])
    data=book([[-100,'',508,'Pago: Empresa: Fornecedor ABC'],[-200,'',508,'Pago: Empresa: Fornecedor ABC']],True)
    r=agent.investigate(242,data,use_ai=False)
    monkeypatch.setattr(base,'_db',lambda:pytest.fail('Aprendizado não autorizado'))
    out,summary=agent.apply(242,data,'modelo.xlsx',[{'_aba':'Itaú 508','_linha':2,'Conta da contrapartida':'412'}],r['token'])
    wb=load_workbook(io.BytesIO(out));assert wb['Itaú 508']['D2'].value==412
    assert wb['Itaú 508']['D3'].value is None
    assert wb['Itaú 508']['G2'].value=='=1+1'
    assert wb['Itaú 508']['G2'].number_format=='0.00'
    assert wb['Principal']['A1'].value=='Preservado'
    assert wb['Retirados']['A1'].value=='Preservado'
    assert summary['aprendizados']==0 and summary['cobertura']['pendentes']==1


@pytest.mark.parametrize('account',['=1+1','999999','0','508','166'])
def test_reject_unknown_or_unsafe_accounts(monkeypatch,account):
    setup(monkeypatch,[record()]);data=book();r=agent.investigate(242,data,use_ai=False)
    with pytest.raises(ValueError): agent.apply(242,data,'modelo.xlsx',[{'_aba':'Itaú 508','_linha':2,'Conta da contrapartida':account}],r['token'])


def test_stale_file_company_base_and_duplicates(monkeypatch):
    records=[record()];setup(monkeypatch,records);data=book();r=agent.investigate(242,data,use_ai=False)
    revision={'_aba':'Itaú 508','_linha':2,'Conta da contrapartida':'412'}
    for company,content in [(1408,data),(242,book([[-101,'',508,'Pago: Empresa: Fornecedor ABC']]))]:
        with pytest.raises(ValueError):agent.apply(company,content,'modelo.xlsx',[revision],r['token'])
    with pytest.raises(ValueError):agent.apply(242,data,'modelo.xlsx',[revision,revision],r['token'])
    records.append(record(account='413'))
    with pytest.raises(ValueError):agent.apply(242,data,'modelo.xlsx',[revision],r['token'])


def test_ai_failure_keeps_local_evidence(monkeypatch):
    from app import classification_agent_ai as ai
    setup(monkeypatch,[record(),record(account='413')])
    monkeypatch.setattr(ai,'choose_provider',lambda:None)
    r=agent.investigate(242,book())
    assert r['ia']['estado']=='indisponivel' and len(r['sugestoes'][0]['evidencias'])==2


def test_learning_only_after_approval(monkeypatch):
    setup(monkeypatch,[record()]);data=book();r=agent.investigate(242,data,use_ai=False)
    calls=[]
    class Conn:
        def execute(self,*args):pass
        def commit(self):calls.append('commit')
        def close(self):pass
    monkeypatch.setattr(base,'_db',lambda:Conn())
    monkeypatch.setattr(base,'_save',lambda con,company,records:calls.append((company,len(records))))
    agent.apply(242,data,'modelo.xlsx',[{'_aba':'Itaú 508','_linha':2,'Conta da contrapartida':'412'}],r['token'],True)
    assert calls==[(242,1),'commit']

def test_1408_preserves_166_and_accepts_zero(monkeypatch):
    setup(monkeypatch,[])
    data=book([[-100,166,512,'Pago: Empresa: Fornecedor ABC'],[-100,0,512,'Pago: Empresa: Fornecedor ABC']])
    assert len(agent.eligible(1408,data,base.resolve_classification_options(1408)))==1


def test_api_investigation_and_review_contract(monkeypatch):
    from fastapi.testclient import TestClient
    from app.main import app
    setup(monkeypatch,[record()]);data=book();client=TestClient(app)
    response=client.post('/api/v1/base-inteligente/242/sugerir-pendentes',files={'file':('modelo.xlsx',data)},data={'usar_ia':'false'})
    assert response.status_code==200
    report=response.json()
    revision={'_aba':'Itaú 508','_linha':2,'Conta da contrapartida':'412'}
    import json
    result=client.post('/api/v1/base-inteligente/242/revisar',files={'file':('modelo.xlsx',data)},data={'investigacao_token':report['token'],'revisoes_json':json.dumps([revision]),'aprender':'false'})
    assert result.status_code==200
    assert json.loads(result.headers['x-razync-summary'])['aplicadas']==1


def test_ai_invalid_response_never_changes_candidates(monkeypatch):
    from app import classification_agent_ai as ai
    from contextlib import contextmanager
    setup(monkeypatch,[record(),record(account='413')])
    monkeypatch.setattr(ai,'choose_provider',lambda:('groq','test','synthetic-key'))
    captured=[]
    @contextmanager
    def slot(*args):yield lambda raw:None
    class Response:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def read(self,n):return b'{"choices":[{"finish_reason":"stop","message":{"content":"{\\"prioridades\\":[999]}"}}]}'
    monkeypatch.setattr(ai,'request_slot',slot)
    def request(req,**kwargs):
        captured.append(req.data.decode());return Response()
    monkeypatch.setattr(ai.urllib.request,'urlopen',request)
    r=agent.investigate(242,book())
    assert r['ia']['estado']=='indisponivel'
    assert r['sugestoes'][0]['conta_sugerida']==''
    assert 'Fornecedor' not in captured[0] and '412' not in captured[0]
    assert len(captured)==1

def test_ai_budget_refusal_never_calls_provider(monkeypatch):
    from app import classification_agent_ai as ai
    from contextlib import contextmanager
    from fastapi import HTTPException
    setup(monkeypatch,[record(),record(account='413')])
    monkeypatch.setattr(ai,'choose_provider',lambda:('groq','test','synthetic'))
    @contextmanager
    def refuse(*args):raise HTTPException(429,'Cota');yield
    monkeypatch.setattr(ai,'request_slot',refuse)
    monkeypatch.setattr(ai.urllib.request,'urlopen',lambda *a,**k:pytest.fail('Não deve chamar o provedor'))
    r=agent.investigate(242,book())
    assert r['ia']['chamadas_nesta_execucao']==0 and r['sugestoes']


def test_reject_wrong_nature_and_expired_ticket(monkeypatch):
    records=[record(),record(account='777',nature='recebido')];setup(monkeypatch,records)
    data=book();r=agent.investigate(242,data,use_ai=False)
    revision={'_aba':'Itaú 508','_linha':2,'Conta da contrapartida':'777'}
    with pytest.raises(ValueError):agent.apply(242,data,'modelo.xlsx',[revision],r['token'])
    for item in agent._cache.values():item['time']-=agent._TTL+1
    revision['Conta da contrapartida']='412'
    with pytest.raises(ValueError):agent.apply(242,data,'modelo.xlsx',[revision],r['token'])


def test_cross_bank_conflict_cannot_be_strong(monkeypatch):
    setup(monkeypatch,[record(),record(bank='bb',account='413')])
    r=agent.investigate(242,book(),use_ai=False)
    assert r['sugestoes'][0]['situacao']=='MANUAL'


def test_cache_invalidates_after_learning_changes(monkeypatch):
    records=[record()];setup(monkeypatch,records);data=book()
    r=agent.investigate(242,data,use_ai=False)
    records.append(record(account='413'))
    changed=agent.investigate(242,data,use_ai=False)
    assert not changed['cache'] and changed['token']!=r['token']


def test_received_credit_only_and_formula_protection(monkeypatch):
    setup(monkeypatch,[record(nature='recebido',account='777',signature=agent.engine.criar_assinatura_classificacao('Recebido: Empresa: Fornecedor ABC'))])
    data=book([[100,508,'','Recebido: Empresa: Fornecedor ABC'],[-100,'=1+1',508,'Pago: Empresa: Fornecedor ABC']])
    r=agent.investigate(242,data,use_ai=False)
    assert r['total_pendencias']==1 and r['sugestoes'][0]['conta_sugerida']=='777'
    out,_=agent.apply(242,data,'modelo.xlsx',[{'_aba':'Itaú 508','_linha':2,'Conta da contrapartida':'777'}],r['token'])
    wb=load_workbook(io.BytesIO(out));assert wb.active['D2'].value==508 and wb.active['E2'].value==777
    assert wb.active['D3'].value=='=1+1'

def test_account_history_catalog_is_scoped_and_bounded(monkeypatch):
    records=[record(),record(bank='bb'),record(nature='recebido',account='777')]
    setup(monkeypatch,records)
    report=agent.investigate(242,book(),use_ai=False)
    payment=next(item for item in report['catalogo_contas'] if item['conta']=='412')
    assert payment['natureza']=='pago'
    assert payment['bancos']==['bb','itau_508']
    assert len(payment['exemplos'])==2
    assert payment['exemplos'][0]['historico']=='Pago: Empresa: Fornecedor ABC'
    assert payment['exemplos'][0]['periodos']==['2026-07','2026-08']
    assert payment['exemplos'][0]['ocorrencias']==20
    other=agent.investigate(1408,book([[-100,'',512,'Pago: Empresa: Fornecedor ABC']]),use_ai=False)
    assert other['catalogo_contas']==[]
